from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SOURCE = Path(
    r"C:\Users\梁嘎嘎\Desktop\权利要求与附图优化版_建筑内局部证据与跨事件关系记忆建筑损毁评估.docx"
)
OUTPUT = Path(
    r"D:\software\dowmload\BRIGHT-master\deliverables\全面审核修订版_建筑内局部证据与跨事件关系记忆建筑损毁评估方法.docx"
)

TITLE = "一种基于建筑内局部证据聚合 跨事件关系记忆与效用验证的多模态建筑损毁评估方法"

ABSTRACT = (
    "本发明公开一种多模态建筑损毁评估方法。利用灾前光学遥感影像和灾后合成孔径雷达影像获得基础预测；"
    "在预测建筑实例内生成局部区域，以同建筑硬约束排除跨建筑关联，并在建筑实例与局部邻域的交集中搜索雷达响应，"
    "以零初始化、幅值受限的变化残差修正光学结构关联。聚合局部软峰值、匹配质量和关键区域表示，形成建筑内损毁证据结构。"
    "仅复核基础类别为完整或处于完整与损毁混淆状态的建筑；从排除同灾害事件样本的关系记忆库获得历史证据，"
    "根据当前证据与历史证据的互斥关系和检索可靠度估计修正效用。仅当基础类别、疑似损毁分数和修正效用满足门控条件时，"
    "将目标掩膜内原完整类别像素修正为损毁类别，其他类别及基础网络中间特征保持不变。"
)

CLAIM_1 = [
    "1.一种基于建筑内局部证据聚合、跨事件关系记忆与效用验证的多模态建筑损毁评估方法，其特征在于，包括：",
    "获取目标区域的灾前光学遥感影像和灾后合成孔径雷达影像，将所述灾前光学遥感影像和灾后合成孔径雷达影像输入基础建筑损毁评估网络，获得建筑位置概率、基础建筑损毁类别概率、光学结构特征和雷达响应特征，所述基础建筑损毁类别包括背景类别、完整类别、损毁类别和毁坏类别；",
    "对所述建筑位置概率进行阈值化和连通域分解，得到多个预测建筑实例；针对第i个预测建筑实例，根据其建筑面积确定局部区域中心数量并将各局部区域中心吸附至该预测建筑实例内的有效建筑像素；构造像素与局部区域中心之间的同建筑硬约束掩膜，使不属于同一预测建筑实例的像素与局部区域中心的掩膜值为零，并在软关联归一化之前排除对应的跨建筑候选；",
    "针对第i个预测建筑实例内的建筑像素，根据该建筑像素的光学结构特征、候选局部区域中心的光学结构特征及二者的相对位置计算结构关联得分；在所述第i个预测建筑实例与以该建筑像素为中心的预设搜索邻域的交集中，对雷达响应特征进行相似度加权搜索，得到匹配后的雷达响应特征、匹配置信度和匹配不确定度；根据所述光学结构特征与匹配后的雷达响应特征生成变化残差，利用初始值为零且绝对值受预设上限约束的门控系数将所述变化残差叠加至结构关联得分，并在所述同建筑硬约束掩膜内归一化，得到像素与局部区域中心之间的软关联；",
    "按照所述软关联聚合各局部区域的光学结构特征、匹配后的雷达响应特征、跨模态关系统计量、匹配置信度、匹配不确定度和区域面积占比，并计算区域内像素证据的可微软峰值相对于加权均值的增量，得到局部区域证据表示；根据关键性分数选取关键局部区域，将同一预测建筑实例内全部局部区域证据表示的总体均值、内部离散度和所述关键局部区域的证据表示组合为建筑内损毁证据结构；",
    "当预测建筑实例的基础建筑损毁类别为完整类别，或者完整类别概率与损毁类别概率之差的绝对值不大于预设混淆阈值时，将该预测建筑实例确定为疑难建筑候选；将所述疑难建筑候选的基础建筑损毁类别概率、基础预测不确定度和建筑内损毁证据结构输入条件损毁专家，得到疑似损毁分数和当前证据裕量；",
    "将历史建筑的建筑级检索表示、关键局部区域证据表示集合、建筑损毁类别、历史修正效用和灾害事件标识存入跨灾害事件关系记忆库；检索所述疑难建筑候选时，先排除灾害事件标识与所述疑难建筑候选相同的历史建筑，再依次执行建筑级粗检索和关键局部区域双向匹配重排序，得到历史损毁支持度、历史完整支持度以及由二者之差形成的历史证据裕量；根据当前证据裕量和历史证据裕量的符号及其是否达到强度阈值，将二者的关系编码为强证据方向一致、仅当前证据达到强度阈值、仅历史证据达到强度阈值或强证据方向冲突四种互斥关系；",
    "将所述疑似损毁分数、历史证据裕量、所述互斥关系、基础预测不确定度以及由检索相似度、历史邻居类别一致性、历史邻居事件多样性和检索得分离散度形成的检索可靠度输入残差效用验证模块，估计候选修正相对于保持基础预测的效用值；仅在基础建筑损毁类别为完整类别、疑似损毁分数达到专家阈值且效用值达到效用阈值时开启修正门，在保持所述基础建筑损毁评估网络中间特征不变的情况下，将目标预测建筑实例或其关键局部区域内原预测为完整类别的像素修正为损毁类别，并保持背景类别、原损毁类别和原毁坏类别不变；未开启所述修正门时保持基础建筑损毁预测结果不变。",
]

DEPENDENT_CLAIMS = [
    "2.根据权利要求1所述的方法，其特征在于，所述第i个预测建筑实例的局部区域中心数量为该预测建筑实例的面积与目标局部区域面积之比的向上取整值，并限制在1与每栋建筑的最大局部区域数量之间；各局部区域中心经紧凑网格初始化后，分别吸附至对应预测建筑实例内距离最近的有效建筑像素。",
    "3.根据权利要求1所述的方法，其特征在于，所述相似度加权搜索根据光学结构特征映射向量与雷达响应特征映射向量的归一化内积生成搜索权重，以搜索权重的归一化熵作为匹配不确定度、以一减所述归一化熵作为匹配置信度；所述变化残差由光学结构特征、匹配后的雷达响应特征、二者的绝对差和逐元素乘积生成；所述门控系数由零初始化的可学习增益经双曲正切函数限幅后与Sigmoid门值相乘得到，其绝对值不超过预设上限。",
    "4.根据权利要求1所述的方法，其特征在于，所述可微软峰值为局部区域内像素证据得分按照温度控制指数权重形成的加权和；所述关键性分数由可微软峰值与加权均值之差、跨模态关系统计量、匹配置信度和区域面积占比确定，并据此选取前M个局部区域；不存在局部区域损毁标注时，仅在同一预测建筑实例内生成并归一化各局部区域的注意力权重，以局部区域证据表示的加权和获得建筑级预测，并利用建筑级损毁标签训练局部区域证据表示和条件损毁专家。",
    "5.根据权利要求1所述的方法，其特征在于，所述建筑级粗检索利用建筑内损毁证据结构的检索表示计算当前建筑与历史建筑的相似度；所述关键局部区域双向匹配重排序分别计算当前建筑关键局部区域在历史建筑中的最大局部支持和历史建筑关键局部区域在当前建筑中的最大局部支持，并将两个方向的平均支持与所述相似度组合为重排序得分；所述检索可靠度由重排序得分、历史邻居类别一致性、历史邻居事件多样性和检索得分方差形成，所述历史邻居事件多样性为历史邻居所覆盖的灾害事件数量与历史邻居数量之比；所述残差效用验证模块以候选修正前后的分类损失差是否大于预设效用增量生成二值监督信号。",
]


def _replace_text(paragraph, text: str) -> None:
    rpr = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        rpr = deepcopy(paragraph.runs[0]._r.rPr)
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)


def _remove_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _insert_claims_section_break(claim_last_paragraph, description_break_paragraph) -> None:
    source_ppr = description_break_paragraph._p.pPr
    if source_ppr is None:
        raise RuntimeError("Description section properties are missing")
    source_sectpr = source_ppr.find(qn("w:sectPr"))
    if source_sectpr is None:
        raise RuntimeError("Description section break is missing")

    claim_ppr = claim_last_paragraph._p.get_or_add_pPr()
    existing = claim_ppr.find(qn("w:sectPr"))
    if existing is not None:
        claim_ppr.remove(existing)
    claim_sectpr = deepcopy(source_sectpr)
    for header_ref in list(claim_sectpr.findall(qn("w:headerReference"))):
        claim_sectpr.remove(header_ref)
    claim_ppr.append(claim_sectpr)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    doc = Document(SOURCE)
    paragraphs = doc.paragraphs

    # 摘要与方法权利要求。
    _replace_text(paragraphs[0], ABSTRACT)
    for paragraph, text in zip(paragraphs[4:12], CLAIM_1, strict=True):
        _replace_text(paragraph, text)
    for paragraph, text in zip(paragraphs[12:16], DEPENDENT_CLAIMS, strict=True):
        _replace_text(paragraph, text)

    # 权利要求书采用略收紧的行距，避免权利要求5末行孤立成页；不改变字体、字号或技术内容。
    for paragraph in paragraphs[4:16]:
        paragraph.paragraph_format.line_spacing = 1.45

    # 删除系统权利要求；桌面源文件已不存在原存储介质权利要求。
    _remove_paragraph(paragraphs[16])

    # 说明书题名、技术领域与末尾实施方式同步收敛为方法主题。
    _replace_text(paragraphs[17], TITLE)
    _replace_text(
        paragraphs[19],
        "本发明涉及遥感图像智能解译、灾害监测、合成孔径雷达影像处理以及多模态机器学习技术领域，尤其涉及一种利用灾前光学遥感影像和灾后合成孔径雷达影像，在预测建筑实例内部构造局部损毁证据，并经跨灾害事件历史证据验证后选择性修正建筑损毁预测结果的方法。",
    )

    # 桌面源文件把权利要求书和说明书合并在同一分节中；在权利要求5后恢复分节。
    _insert_claims_section_break(paragraphs[15], paragraphs[191])

    # 删除通用参数校准段、系统实现段和无明确技术边界的等同替换兜底段。
    _remove_paragraph(paragraphs[187])
    _remove_paragraph(paragraphs[189])
    _remove_paragraph(paragraphs[190])

    doc.core_properties.title = TITLE
    doc.core_properties.subject = "中国发明专利申请文件"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)

    # 为新分节建立独立的“权利要求书”页眉，同时保留说明书页眉和页码重新起算。
    doc = Document(OUTPUT)
    if len(doc.sections) != 5:
        raise RuntimeError(f"Expected five document sections, got {len(doc.sections)}")
    claim_header = doc.sections[2].header
    claim_header.is_linked_to_previous = False
    description_header_paragraph = doc.sections[3].header.paragraphs[0]
    if claim_header.paragraphs:
        claim_header._element.remove(claim_header.paragraphs[0]._p)
    claim_header._element.append(deepcopy(description_header_paragraph._p))
    _replace_text(claim_header.paragraphs[0], "权利要求书")
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
