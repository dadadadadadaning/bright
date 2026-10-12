from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document


SOURCE = Path(
    r"C:\Users\梁嘎嘎\Desktop\全面审核修订版_建筑内局部证据与跨事件关系记忆建筑损毁评估方法.docx"
)
OUTPUT = Path(
    r"D:\software\dowmload\BRIGHT-master\deliverables\按参考模板重构版_建筑内局部证据与跨事件关系记忆建筑损毁评估方法.docx"
)
FIGURE = Path(
    r"D:\software\dowmload\BRIGHT-master\tmp\patent_reference\figure1_s1_s6.png"
)


CLAIM_1 = [
    "1.一种基于建筑内局部证据聚合、跨事件关系记忆与效用验证的多模态建筑损毁评估方法，其特征在于，包括：",
    "S1：获取目标区域的灾前光学遥感影像和灾后合成孔径雷达影像，输入基础建筑损毁评估网络，获得建筑位置概率、基础建筑损毁类别概率、光学结构特征和雷达响应特征；",
    "S2：根据所述建筑位置概率形成预测建筑实例，在每一预测建筑实例内设置局部区域中心，并以同建筑硬约束排除跨建筑候选；",
    "S3：在同一预测建筑实例与局部搜索邻域的交集中搜索雷达响应，以幅值受限的变化残差修正光学结构关联，得到建筑内软关联；",
    "S4：按照所述建筑内软关联聚合局部光学结构、匹配后的雷达响应、跨模态关系统计量和局部强响应，形成局部区域证据表示，并构建建筑内损毁证据结构；",
    "S5：筛选完整类疑难建筑，由条件损毁专家生成疑似损毁分数和当前证据裕量；从排除同灾害事件病例的关系记忆库中检索历史损毁支持度与历史完整支持度，并由二者之差形成历史证据裕量，确定当前证据与历史证据的关系；",
    "S6：根据所述疑似损毁分数、历史证据裕量、所述关系和基础预测不确定度估计候选修正效用；仅当所述候选修正效用满足接受条件时，在保持基础建筑损毁评估网络中间特征不变的情况下，将目标范围内原预测为完整类别的像素修正为损毁类别，否则保持基础建筑损毁预测结果不变。",
]


DEPENDENT_CLAIMS = [
    "2.根据权利要求1所述的方法，其特征在于，S2中，第i个预测建筑实例的局部区域中心数量为该预测建筑实例的面积与目标局部区域面积之比的向上取整值，并限制在1与每栋建筑的最大局部区域数量之间；各局部区域中心经紧凑网格初始化后吸附至对应预测建筑实例内距离最近的有效建筑像素；在建筑内软关联归一化之前，将不属于同一预测建筑实例的像素与局部区域中心之间的候选关系置为无效。",
    "3.根据权利要求1所述的方法，其特征在于，S3中，根据光学结构特征映射向量与雷达响应特征映射向量的归一化内积生成搜索权重，以所述搜索权重的归一化熵作为匹配不确定度，并以一减所述归一化熵作为匹配置信度；由光学结构特征、匹配后的雷达响应特征、二者的绝对差和逐元素乘积生成所述变化残差；由零初始化的可学习增益经双曲正切函数限幅后与Sigmoid门值相乘得到门控系数，使所述门控系数的绝对值不超过预设上限。",
    "4.根据权利要求1所述的方法，其特征在于，S4中，所述局部强响应为局部区域内像素证据得分按照温度控制指数权重形成的可微软峰值；根据所述可微软峰值与加权均值之差、跨模态关系统计量、匹配置信度和区域面积占比确定关键性分数并选取前M个关键局部区域；在不存在局部区域损毁标注时，仅在同一预测建筑实例内生成并归一化各局部区域的注意力权重，利用建筑级损毁标签训练局部区域证据表示和条件损毁专家。",
    "5.根据权利要求1所述的方法，其特征在于，S5中，先依据建筑级检索表示对跨事件历史建筑进行粗检索，再依据当前建筑与历史建筑的关键局部区域之间的双向最大局部支持进行重排序，获得历史损毁支持度和历史完整支持度；根据当前证据裕量与由所述历史损毁支持度和历史完整支持度之差形成的历史证据裕量，将二者的关系编码为强证据方向一致、仅当前证据达到强度阈值、仅历史证据达到强度阈值或强证据方向冲突四种互斥关系；S6中，根据候选修正前后的分类损失差是否大于预设效用增量生成残差效用验证的二值监督信号。",
]


TECHNICAL_SOLUTION = [
    "为实现上述目的，本发明采用如下技术方案。",
    "S1：获取目标区域的灾前光学遥感影像和灾后合成孔径雷达影像，输入基础建筑损毁评估网络，获得建筑位置概率、基础建筑损毁类别概率、光学结构特征和雷达响应特征；",
    "S2：根据建筑位置概率形成预测建筑实例，在每一预测建筑实例内设置局部区域中心，并以同建筑硬约束排除跨建筑候选；",
    "S3：在同一预测建筑实例与局部搜索邻域的交集中搜索雷达响应，以幅值受限且初始影响为零的变化残差修正光学结构关联，得到建筑内软关联；",
    "S4：按照建筑内软关联聚合局部光学结构、匹配后的雷达响应、跨模态关系统计量和局部强响应，形成局部区域证据表示及建筑内损毁证据结构；",
    "S5：筛选基础预测中的完整类疑难建筑，由条件损毁专家生成疑似损毁分数和当前证据裕量，并从排除同灾害事件建筑的关系记忆库检索历史损毁支持度和历史完整支持度，确定当前证据与历史证据的关系；",
    "S6：根据疑似损毁分数、历史损毁支持度与历史完整支持度之差、两类证据的关系和基础预测不确定度估计候选修正效用，仅在候选修正效用满足接受条件时，对原预测为完整类别的目标像素实施末端损毁修正。",
    "进一步地，S2中，根据预测建筑实例的面积确定局部区域中心数量，将局部区域中心吸附至对应建筑内部的有效位置，并在软关联归一化之前通过同建筑硬约束掩膜排除跨建筑候选。",
    "进一步地，S3中，利用光学结构特征与候选雷达响应之间的相似度进行建筑内局部搜索，获得匹配后的雷达响应、匹配置信度和匹配不确定度；将光学结构特征与匹配后的雷达响应形成的变化残差经零初始化、幅值受限的门控系数叠加至光学结构关联。",
    "进一步地，S4中，在区域均值之外保留跨模态关系离散度和可微软峰值，并结合匹配质量与区域面积占比形成局部区域证据表示；根据同一建筑内全部局部区域证据表示的总体统计量、内部离散度和关键局部区域表示形成建筑内损毁证据结构。",
]


OVERVIEW_STEPS = [
    "如图1所示，本发明实施例提供的多模态建筑损毁评估方法包括：",
    CLAIM_1[1],
    CLAIM_1[2],
    CLAIM_1[3],
    CLAIM_1[4],
    CLAIM_1[5],
    CLAIM_1[6],
]


def replace_text(paragraph, text: str) -> None:
    rpr = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        rpr = deepcopy(paragraph.runs[0]._r.rPr)
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)


def remove_paragraph(paragraph) -> None:
    parent = paragraph._element.getparent()
    if parent is not None:
        parent.remove(paragraph._element)


def insert_before(anchor, text: str, format_source) -> None:
    paragraph = anchor.insert_paragraph_before()
    source_ppr = format_source._p.pPr
    if source_ppr is not None:
        current_ppr = paragraph._p.pPr
        if current_ppr is not None:
            paragraph._p.remove(current_ppr)
        paragraph._p.insert(0, deepcopy(source_ppr))
    paragraph.add_run(text)
    if format_source.runs and format_source.runs[0]._r.rPr is not None:
        paragraph.runs[0]._r.insert(0, deepcopy(format_source.runs[0]._r.rPr))


def draw_centered_multiline(draw, box, text, font, spacing=12):
    left, top, right, bottom = box
    lines = text.split("\n")
    heights = []
    widths = []
    for line in lines:
        bounds = draw.textbbox((0, 0), line, font=font)
        widths.append(bounds[2] - bounds[0])
        heights.append(bounds[3] - bounds[1])
    total = sum(heights) + spacing * (len(lines) - 1)
    y = top + (bottom - top - total) / 2
    for line, width, height in zip(lines, widths, heights):
        x = left + (right - left - width) / 2
        draw.text((x, y), line, font=font, fill="black")
        y += height + spacing


def create_flow_figure() -> bytes:
    FIGURE.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (2000, 2300), "white")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(r"C:\Windows\Fonts\simhei.ttf", 70)
    body_font = ImageFont.truetype(r"C:\Windows\Fonts\simsun.ttc", 49)
    label_font = ImageFont.truetype(r"C:\Windows\Fonts\simhei.ttf", 58)
    title = "多模态建筑损毁评估方法流程"
    title_box = draw.textbbox((0, 0), title, font=title_font)
    draw.text(((2000 - (title_box[2] - title_box[0])) / 2, 75), title, font=title_font, fill="black")

    items = [
        "获取灾前光学影像与灾后SAR影像\n输出基础预测和双模态特征",
        "生成预测建筑实例与局部区域\n以同建筑硬约束排除跨建筑候选",
        "在建筑内部搜索雷达局部响应\n以限幅变化残差修正光学结构关联",
        "聚合局部证据并选取关键区域\n构建建筑内损毁证据结构",
        "筛选完整类疑难建筑\n结合条件专家与跨事件关系记忆证据",
        "验证候选修正效用\n选择性实施完整类别至损毁类别修正",
    ]
    left, right = 210, 1580
    top, height, gap = 260, 250, 95
    for index, item in enumerate(items, 1):
        y1 = top + (index - 1) * (height + gap)
        y2 = y1 + height
        draw.rounded_rectangle((left, y1, right, y2), radius=18, outline="black", width=5)
        draw_centered_multiline(draw, (left + 35, y1 + 15, right - 35, y2 - 15), item, body_font)
        mid = (y1 + y2) // 2
        draw.line((right, mid, 1700, mid), fill="black", width=4)
        draw.text((1740, mid - 34), f"S{index}", font=label_font, fill="black")
        if index < len(items):
            x = (left + right) // 2
            arrow_top = y2
            arrow_bottom = y2 + gap - 18
            draw.line((x, arrow_top, x, arrow_bottom), fill="black", width=5)
            draw.polygon([(x, y2 + gap), (x - 16, arrow_bottom - 2), (x + 16, arrow_bottom - 2)], fill="black")
    image.save(FIGURE, format="PNG", optimize=True)
    return FIGURE.read_bytes()


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    doc = Document(SOURCE)
    p = doc.paragraphs
    if len(p) < 196:
        raise RuntimeError(f"Unexpected source structure: {len(p)} paragraphs")

    # 权利要求1按参考专利的写法收敛为S1至S6主链，细节下沉至四项从属权利要求。
    for paragraph, text in zip(p[4:11], CLAIM_1, strict=True):
        replace_text(paragraph, text)
    remove_paragraph(p[11])
    for paragraph, text in zip(p[12:16], DEPENDENT_CLAIMS, strict=True):
        replace_text(paragraph, text)
    for paragraph in p[4:16]:
        paragraph.paragraph_format.line_spacing = 1.5

    replace_text(
        p[18],
        "本发明涉及遥感影像智能解译与多模态机器学习技术领域，具体涉及一种利用灾前光学遥感影像和灾后合成孔径雷达影像进行建筑损毁评估的方法。",
    )
    replace_text(
        p[31],
        "本发明要解决的技术问题是：在光学与雷达影像存在成像差异和局部错位的条件下，避免不同建筑之间发生证据串联，保留建筑内部的小面积强损毁响应，并防止跨事件历史病例在证据不可靠时误改基础预测。",
    )

    for paragraph, text in zip(p[33:43], TECHNICAL_SOLUTION, strict=True):
        replace_text(paragraph, text)
    replace_text(
        p[43],
        "进一步地，S5中，仅对基础类别为完整或完整类别与损毁类别相互混淆的预测建筑启动条件损毁专家；跨事件关系记忆检索先排除同灾害事件病例，再依次进行建筑级粗检索和关键局部区域双向匹配重排序。S6中，将当前证据与历史证据的互斥关系、检索可靠度和基础预测不确定度用于候选修正效用判断。建筑内局部区域、跨模态残差、软关联、局部强响应、历史证据裕量和修正门满足下列关系：",
    )

    # 具体实施方式先重复S1至S6总流程，再逐步展开，形成与参考专利一致的“总述—分步详述”结构。
    overview_anchor = p[68]
    for text in OVERVIEW_STEPS:
        insert_before(overview_anchor, text, p[68])
    replace_text(
        p[68],
        "在一种可实施方式中，S1具体包括以下处理过程。获取灾前光学遥感影像Iopt和灾后合成孔径雷达影像Isar，对二者进行地理参考统一、尺度归一化和必要的辐射预处理。全局配准只需使同一建筑落入相邻区域，不要求光学像素与雷达像素严格同坐标对应。",
    )
    replace_text(
        p[74],
        "在一种可实施方式中，S2具体包括以下处理过程。对Pbuilding进行阈值化和四连通或八连通分析，删除面积小于预设阈值的伪建筑，获得预测建筑实例集合B={B1，B2，…，BN}。所述建筑实例来自模型预测而非推理阶段的真实标注，因此能够直接用于部署。",
    )
    replace_text(
        p[89],
        "在一种可实施方式中，S3具体包括以下处理过程。对于光学位置p，在Bi与以p为中心的预设搜索邻域的交集中构造雷达候选集合Ni(p)。利用分别作用于光学特征和雷达特征的映射函数计算候选关系分数，并对候选关系分数进行归一化加权，得到匹配后的雷达表示。该步骤允许雷达响应相对于光学结构发生小范围位移。",
    )
    replace_text(
        p[106],
        "在一种可实施方式中，S4具体包括以下处理过程。针对局部区域Rik，利用Q计算光学结构均值、雷达响应均值和跨模态关系均值。进一步计算跨模态关系标准差，以描述区域内部异质性；匹配置信度均值和匹配熵均值用于描述该区域跨模态证据的可靠性；区域面积与建筑面积之比用于描述区域尺度。",
    )
    replace_text(
        p[129],
        "在一种可实施方式中，S5具体包括以下处理过程。根据Pbase计算每个预测建筑的类别概率和不确定度。若基础建筑类别为完整，或者完整与损毁概率之差小于预设阈值，则把该建筑加入疑难候选集合；其他建筑仍直接采用基础预测。",
    )
    replace_text(
        p[164],
        "在一种可实施方式中，S6具体包括以下处理过程。证据残差效用验证模块的输入包括sDI、基础预测不确定度、历史支持差、原型支持差、前若干邻居的相似度、类别一致性、事件多样性以及检索得分离散度，输出0至1之间的候选修正效用值u。",
    )
    replace_text(
        p[181],
        "在一种训练方式中，按照S1至S6对应的模块分阶段训练。第一阶段训练或复现基础建筑损毁评估网络；第二阶段固定基础网络并训练光学结构关联；第三阶段打开受限变化残差，利用建筑级弱监督训练跨模态分区和区域证据；第四阶段训练条件损毁专家；第五阶段建立跨事件关系记忆库并训练或校准残差效用验证模块。",
    )
    remove_paragraph(p[186])
    replace_text(p[189], "图1 本发明多模态建筑损毁评估方法S1至S6总体技术流程示意图")

    # 用S1至S6流程图替换原S1至S8流程图；该图片在摘要附图和说明书附图中共用。
    figure_blob = create_flow_figure()
    replaced = False
    for image_part in doc.part.package.image_parts:
        if str(image_part.partname) == "/word/media/image1.png":
            image_part._blob = figure_blob
            replaced = True
    if not replaced:
        raise RuntimeError("Flow figure image part was not found")

    doc.core_properties.subject = "中国发明专利申请文件"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
