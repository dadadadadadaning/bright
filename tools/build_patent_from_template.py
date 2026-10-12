from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(r"D:\software\dowmload\BRIGHT-master")
TEMPLATE = Path(
    r"D:\software\微信文件\xwechat_files\wxid_g66h01ctvlv222_99ac\temp\RWTemp\2025-12\22bc873dd63a0677ecec0644b2e355f9\专利撰写模板.docx"
)
SOURCE_PATENT = ROOT / "deliverables" / "建筑内局部证据与跨事件效用验证的建筑损毁评估发明专利初稿.docx"
FIG_DIR = ROOT / ".patent_build" / "figures"
OUTPUT = ROOT / "deliverables" / "按专利模板修订版_建筑内局部证据与跨事件关系记忆建筑损毁评估.docx"

TITLE = "一种基于建筑内局部证据聚合 跨事件关系记忆与效用验证的多模态建筑损毁评估方法 系统及存储介质"

ABSTRACT = (
    "本发明公开一种多模态建筑损毁评估方法、系统及存储介质。"
    "利用灾前光学影像和灾后合成孔径雷达影像获得基础预测；在每个预测建筑内设置局部区域，以光学结构形成基础分区，以建筑内搜索的雷达关系作有界残差修边，并聚合局部统计量、匹配质量和面积形成建筑内损毁证据结构。"
    "仅对被判为完整或与损毁混淆的建筑启动条件损毁专家；从排除同事件的关系记忆库检索历史证据，按照当前建议与历史建议的一致、单边或冲突关系验证修正效用。"
    "效用满足条件时，在不改变基础网络中间特征的情况下，将目标区域内原完整类别修正为损毁类别；否则保持原结果。"
    "本发明可降低跨模态错位、局部证据稀释、类别长尾及历史证据误用造成的误判。"
)


CLAIMS = [
    (
        "1.一种基于建筑内局部证据聚合、跨事件关系记忆与效用验证的多模态建筑损毁评估方法，其特征在于，包括：",
        [
            "获取目标区域的灾前光学遥感影像和灾后合成孔径雷达影像，将所述灾前光学遥感影像和灾后合成孔径雷达影像输入基础建筑损毁评估网络，获得建筑位置概率、基础建筑损毁类别概率、光学结构特征和雷达响应特征，所述基础建筑损毁类别至少包括背景类别、完整类别、损毁类别和毁坏类别；",
            "根据所述建筑位置概率获得多个预测建筑实例，并针对每一预测建筑实例设置多个局部区域中心，将任一建筑像素的候选局部区域中心限制为同一预测建筑实例内的局部区域中心，获得不跨越不同预测建筑实例的建筑约束候选集合；",
            "根据所述光学结构特征形成所述建筑约束候选集合内的结构关联得分，在同一预测建筑实例内从所述雷达响应特征中搜索与光学结构位置相对应的局部响应，获得跨模态关系特征、匹配置信度和匹配不确定度，并以幅值受限的变化残差修正所述结构关联得分，得到像素与局部区域中心之间的软关联；",
            "按照所述软关联聚合各局部区域的光学结构特征、匹配后的雷达响应特征、跨模态关系统计量、局部强响应、匹配质量和区域面积占比，形成区域证据表示，并根据同一预测建筑实例内多个区域证据表示的总体统计量、内部异质性和关键局部区域表示形成建筑内损毁证据结构；",
            "从基础建筑损毁预测结果中选取被预测为完整类别或者完整类别与损毁类别的预测分数满足预设混淆条件的预测建筑实例作为疑难建筑候选，利用以所述建筑内损毁证据结构为输入的条件损毁专家获得当前修正建议和疑似损毁分数；",
            "将所述疑难建筑候选与跨灾害事件关系记忆库进行比较，排除与所述疑难建筑候选属于同一灾害事件的历史建筑，获得历史修正建议和历史损毁支持信息，并根据所述当前修正建议与历史修正建议形成一致关系、仅当前证据支持关系、仅历史证据支持关系或冲突关系；",
            "根据所述疑似损毁分数、历史损毁支持信息、所述关系以及基础建筑损毁预测的不确定度，估计候选修正相对于保持基础预测的残差效用；仅在所述残差效用满足预设接受条件时，在保持所述基础建筑损毁评估网络的中间特征不变的情况下，将目标预测建筑实例或其关键局部区域内原预测为完整类别的像素选择性修正为损毁类别；在不满足所述预设接受条件时保持基础建筑损毁预测结果不变。",
        ],
    ),
    (
        "2.根据权利要求1所述的方法，其特征在于，所述多个预测建筑实例通过对建筑位置概率阈值化并进行连通域分析获得；每一预测建筑实例的局部区域中心数量根据该预测建筑实例的面积、目标区域面积、最小区域面积和每栋建筑的最大区域数量确定，且候选局部区域中心被吸附至该预测建筑实例内部的有效位置。",
        [],
    ),
    (
        "3.根据权利要求1所述的方法，其特征在于，所述跨模态局部响应搜索在对应预测建筑实例与预设邻域的交集中进行；所述变化残差由光学特征、候选雷达特征、二者的绝对差、二者的乘积、匹配置信度、匹配不确定度和基础建筑损毁预测的不确定度中的至少两项生成，并通过初始值为零且绝对值不超过预设上限的门控系数作用于所述结构关联得分。",
        [],
    ),
    (
        "4.根据权利要求1所述的方法，其特征在于，所述局部强响应为根据区域内像素证据得分进行温度控制加权得到的可微软峰值；在不存在局部区域损毁标注时，将同一预测建筑实例的多个区域证据表示通过最大聚合、前若干项聚合、注意力聚合或对数和指数聚合中的至少一种方式得到建筑级结果，并以建筑级损毁标签对所述区域证据表示和条件损毁专家进行弱监督训练。",
        [],
    ),
    (
        "5.根据权利要求1所述的方法，其特征在于，所述跨灾害事件关系记忆库至少存储历史建筑的建筑级检索表示、局部区域证据表示集合、建筑损毁类别、历史修正效用和灾害事件标识；先依据建筑级检索表示进行粗检索，再依据关键局部区域表示集合进行细粒度重排序，并根据历史损毁支持度与历史完整支持度之差、检索相似度、邻居类别一致性、邻居事件多样性、检索得分离散度以及所述当前修正建议与历史修正建议的关系中的至少两项确定所述残差效用。",
        [],
    ),
    (
        "6.一种多模态建筑损毁评估系统，其特征在于，包括基础预测模块、建筑约束局部区域生成模块、建筑内跨模态搜索模块、建筑内损毁证据结构构建模块、条件损毁专家模块、跨事件关系记忆检索模块、残差效用验证模块以及选择性末端修正模块；各模块被配置为协同执行权利要求1至5任一项所述的方法。",
        [],
    ),
    (
        "7.一种计算机可读存储介质，其上存储有计算机程序，其特征在于，所述计算机程序被处理器执行时实现权利要求1至5任一项所述的方法。",
        [],
    ),
]


FIGURES = [
    ("figure_1_overall.png", "图1 本发明多模态建筑损毁评估方法的总体技术流程示意图"),
    ("figure_2_building_regions.png", "图2 预测建筑约束的局部区域生成示意图"),
    ("figure_3_cross_modal.png", "图3 建筑内跨模态搜索与有界残差修边示意图"),
    ("figure_4_evidence_structure.png", "图4 区域证据表示与建筑内损毁证据结构示意图"),
    ("figure_5_specialist_memory.png", "图5 条件损毁专家与跨事件关系记忆验证示意图"),
    ("figure_6_late_correction.png", "图6 残差效用验证与选择性末端修正示意图"),
]


def set_run_font(run, size=12, bold=False, east_asia="宋体", latin="Times New Roman"):
    run.font.name = latin
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), east_asia)


def style_paragraph(paragraph, role="body"):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.left_indent = Pt(0)
    pf.right_indent = Pt(0)
    if role in {"body", "claim"}:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.first_line_indent = Pt(24)
        pf.line_spacing = 1.5
    elif role == "heading":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Pt(0)
        pf.space_before = Pt(3)
        pf.line_spacing = 1.3
    elif role == "title":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Pt(0)
        pf.line_spacing = 1.5
    elif role in {"formula", "caption", "figure"}:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Pt(0)
        pf.line_spacing = 1.5
        if role == "caption":
            pf.space_before = Pt(15.6)
    elif role == "page_break":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Pt(0)


def move_before(paragraph, marker):
    paragraph._p.getparent().remove(paragraph._p)
    marker._p.addprevious(paragraph._p)


def add_text(doc, marker, text, role="body", bold=False):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    style_paragraph(p, role)
    run = p.add_run(text)
    size = 15 if role == "title" else 12
    set_run_font(run, size=size, bold=(bold or role in {"title", "heading"}))
    if marker is not None:
        move_before(p, marker)
    return p


def add_figure(doc, marker, image_path, caption, width=6.35, add_page_break=False):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    style_paragraph(p, "figure")
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width))
    if marker is not None:
        move_before(p, marker)
    cap = add_text(doc, marker, caption, role="caption")
    if add_page_break:
        br = doc.add_paragraph()
        style_paragraph(br, "page_break")
        br.add_run().add_break(WD_BREAK.PAGE)
        if marker is not None:
            move_before(br, marker)
    return cap


def clear_template_body(doc):
    break_paras = [
        p
        for p in doc.paragraphs
        if p._p.pPr is not None and p._p.pPr.sectPr is not None
    ]
    if len(break_paras) != 4:
        raise RuntimeError(f"Expected four section-break paragraphs, got {len(break_paras)}")
    keep = {p._p for p in break_paras}
    body = doc._element.body
    for child in list(body):
        if child.tag == qn("w:p") and child not in keep:
            body.remove(child)
        elif child.tag == qn("w:tbl"):
            body.remove(child)
    for p in break_paras:
        for child in list(p._p):
            if child.tag != qn("w:pPr"):
                p._p.remove(child)
    return break_paras


def normalize_spec_text(text):
    text = re.sub(r"^\[\d{4}\]\s*", "", text.strip())
    replacements = {
        "跨灾害事件历史证据库": "跨灾害事件关系记忆库",
        "历史证据库": "关系记忆库",
        "跨事件历史证据验证": "跨事件关系记忆验证",
        "跨事件历史证据": "跨事件关系记忆证据",
        "证据效用判定模块": "残差效用验证模块",
        "效用判定模块": "残差效用验证模块",
        "效用模块": "残差效用验证模块",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def extract_spec_items():
    src = Document(SOURCE_PATENT)
    items = []
    start = next(i for i, p in enumerate(src.paragraphs) if p.text.strip() == "技术领域")
    end = next(i for i, p in enumerate(src.paragraphs) if p.text.strip() == "说明书附图")
    for p in src.paragraphs[start:end]:
        text = p.text.strip()
        if not text:
            continue
        role = "body"
        if p.style.name in {"Heading 2", "Heading 3"}:
            role = "heading"
        elif p.style.name == "Equation":
            role = "formula"
        original = text
        text = normalize_spec_text(text)
        items.append((role, text))
        if original.startswith("[0018]"):
            items.append(
                (
                    "body",
                    "条件损毁专家亦可称为当前证据分支，其输出表示仅依据当前建筑影像和基础特征形成的修正建议；关系记忆库输出表示依据跨事件历史病例形成的修正建议。二者的关系被显式编码为建议一致、仅当前证据支持、仅关系记忆支持和建议冲突四类，用作残差效用验证的输入，而不是对两路概率作固定平均。",
                )
            )
        if original.startswith("[0051]"):
            items.append(
                (
                    "body",
                    "根据当前证据分支和关系记忆库是否建议改变基础预测，将疑难建筑划分为两者一致、仅当前证据建议修正、仅关系记忆建议修正以及两者冲突四种关系。所述关系与各自置信度、邻居质量和基础模型不确定度共同输入残差效用验证模块，使关系记忆仅在历史建议能够提高基础结果时参与最终修正。",
                )
            )
    return items


def build():
    if len(re.sub(r"\s", "", ABSTRACT)) > 300:
        raise ValueError("Abstract exceeds 300 Chinese characters")
    for path in [TEMPLATE, SOURCE_PATENT, *[FIG_DIR / name for name, _ in FIGURES]]:
        if not path.exists():
            raise FileNotFoundError(path)

    doc = Document(TEMPLATE)
    markers = clear_template_body(doc)

    # Section 1: abstract.
    add_text(doc, markers[0], ABSTRACT, role="body")

    # Section 2: abstract figure.
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    style_paragraph(p, "figure")
    p.paragraph_format.space_before = Pt(20)
    p.add_run().add_picture(str(FIG_DIR / FIGURES[0][0]), width=Inches(6.55))
    move_before(p, markers[1])

    # Section 3: claims.
    for head, steps in CLAIMS:
        add_text(doc, markers[2], head, role="claim")
        for step in steps:
            add_text(doc, markers[2], step, role="claim")

    # Section 4: specification.
    add_text(doc, markers[3], TITLE, role="title")
    for role, text in extract_spec_items():
        if text == "技术领域":
            add_text(doc, markers[3], text, role="heading")
            continue
        add_text(doc, markers[3], text, role=role)

    # Section 5: drawings. One principal figure per page for reliable filing layout.
    for index, (filename, caption) in enumerate(FIGURES):
        add_figure(
            doc,
            None,
            FIG_DIR / filename,
            caption,
            width=6.45,
            add_page_break=index < len(FIGURES) - 1,
        )

    doc.core_properties.title = TITLE
    doc.core_properties.subject = "中国发明专利申请文件"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)
    print(f"abstract_chars={len(re.sub(r'\\s', '', ABSTRACT))}")
    print(f"claims={len(CLAIMS)}")


if __name__ == "__main__":
    build()
