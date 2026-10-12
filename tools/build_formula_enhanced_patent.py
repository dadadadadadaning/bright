from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches
from PIL import Image, ImageDraw


ROOT = Path(r"D:\software\dowmload\BRIGHT-master")
TEMPLATE = Path(
    r"D:\software\微信文件\xwechat_files\wxid_g66h01ctvlv222_99ac\msg\file\2025-12\专利撰写模板.docx"
)
OUTPUT = ROOT / "deliverables" / "权利要求与附图优化版_建筑内局部证据与跨事件关系记忆建筑损毁评估.docx"
WORK_DIR = ROOT / ".patent_formula_work"
FIG_DIR = WORK_DIR / "figures"
EQUATION_JSON = WORK_DIR / "equations.json"

sys.path.insert(0, str(ROOT / "tools"))
import build_patent_from_template as base  # noqa: E402
import redraw_patent_flowcharts as flow  # noqa: E402


FIGURES = [
    ("figure_1_overall.png", "图1 本发明多模态建筑损毁评估方法的总体技术流程示意图"),
    ("figure_2_local_evidence.png", "图2 建筑内局部证据构建的空间约束和信息聚合结构示意图"),
    ("figure_3_memory_utility.png", "图3 跨事件关系记忆与效用门控的双分支决策结构示意图"),
]


REVISED_CLAIMS = [
    (
        "1.一种基于建筑内局部证据聚合、跨事件关系记忆与效用验证的多模态建筑损毁评估方法，其特征在于，包括：",
        [
            "获取目标区域的灾前光学遥感影像和灾后合成孔径雷达影像，将所述灾前光学遥感影像和灾后合成孔径雷达影像输入基础建筑损毁评估网络，获得建筑位置概率、基础建筑损毁类别概率、光学结构特征和雷达响应特征，所述基础建筑损毁类别包括背景类别、完整类别、损毁类别和毁坏类别；",
            "对所述建筑位置概率进行阈值化和连通域分解，得到多个预测建筑实例；针对第i个预测建筑实例，根据其建筑面积确定局部区域中心数量并将各局部区域中心限制在该预测建筑实例内部；构造像素与局部区域中心之间的同建筑硬约束掩膜，使不属于同一预测建筑实例的像素与局部区域中心的掩膜值为零，并在软关联归一化之前排除对应的跨建筑候选；",
            "针对第i个预测建筑实例内的建筑像素，根据该建筑像素的光学结构特征、候选局部区域中心的光学结构特征及二者的相对位置计算结构关联得分；在所述第i个预测建筑实例与以该建筑像素为中心的预设搜索邻域的交集中，对雷达响应特征进行相似度加权搜索，得到匹配后的雷达响应特征、匹配置信度和匹配不确定度；根据所述匹配后的雷达响应特征与光学结构特征生成变化残差，利用初始值为零且绝对值受预设上限约束的门控系数将所述变化残差叠加至结构关联得分，并在所述同建筑硬约束掩膜内归一化，得到像素与局部区域中心之间的软关联；",
            "按照所述软关联分别聚合各局部区域的光学结构特征、匹配后的雷达响应特征、跨模态关系统计量、匹配置信度、匹配不确定度和区域面积占比，并计算区域内像素证据的可微软峰值相对于加权均值的增量，得到局部区域证据表示；根据关键性分数选择多个关键局部区域，将同一预测建筑实例内全部局部区域证据表示的总体均值、内部离散度和所述关键局部区域的证据表示组合为建筑内损毁证据结构；",
            "当预测建筑实例的基础建筑损毁类别为完整类别，或者完整类别概率与损毁类别概率之差的绝对值不大于预设混淆阈值时，将该预测建筑实例确定为疑难建筑候选；将所述疑难建筑候选的基础类别概率、基础预测不确定度和建筑内损毁证据结构输入条件损毁专家，得到疑似损毁分数和当前证据裕量；",
            "将历史建筑的建筑级检索表示、关键局部区域证据表示集合、建筑损毁类别、历史修正效用和灾害事件标识存入跨灾害事件关系记忆库；在对所述疑难建筑候选进行检索时，先排除灾害事件标识与所述疑难建筑候选相同的历史建筑，再依次执行建筑级粗检索和关键局部区域双向匹配重排序，得到历史损毁支持度、历史完整支持度及二者之差所形成的历史证据裕量；根据当前证据裕量和历史证据裕量的符号及强度，将二者的关系编码为强证据方向一致、仅当前证据达到强度阈值、仅历史证据达到强度阈值或强证据方向冲突四种互斥关系；",
            "将所述疑似损毁分数、历史损毁支持度与历史完整支持度之差、所述互斥关系、基础预测不确定度以及由检索相似度、邻居类别一致性、邻居事件多样性和检索得分离散度形成的检索可靠度共同输入残差效用验证模块，估计候选修正相对于保持基础预测的效用值；仅在基础建筑损毁类别为完整类别、疑似损毁分数达到专家阈值且效用值达到效用阈值时开启修正门，在保持所述基础建筑损毁评估网络的中间特征不变的情况下，将目标预测建筑实例或其关键局部区域内原预测为完整类别的像素修正为损毁类别，并保持背景类别、原损毁类别和原毁坏类别不变；未开启所述修正门时保持基础建筑损毁预测结果不变。",
        ],
    ),
    (
        "2.根据权利要求1所述的方法，其特征在于，所述第i个预测建筑实例的局部区域中心数量通过对该预测建筑实例的面积与目标局部区域面积之比向上取整后，限制在1与每栋建筑的最大局部区域数量之间得到；各局部区域中心经紧凑网格初始化后吸附至对应预测建筑实例内部的有效建筑像素，并将面积小于预设最小区域面积的局部区域并入同一预测建筑实例内的相邻局部区域。",
        [],
    ),
    (
        "3.根据权利要求1所述的方法，其特征在于，所述相似度加权搜索以光学结构特征的映射向量与雷达响应特征的映射向量之间的归一化相似度生成搜索权重，并根据所述搜索权重的归一化熵确定匹配不确定度、根据一减所述归一化熵确定匹配置信度；所述变化残差由光学结构特征、匹配后的雷达响应特征、二者的绝对差和二者的逐元素乘积生成，所述门控系数由零初始化的可学习增益、双曲正切函数和Sigmoid函数共同限定。",
        [],
    ),
    (
        "4.根据权利要求1所述的方法，其特征在于，所述可微软峰值通过对局部区域内的像素证据得分进行温度控制的指数加权得到；所述关键性分数由可微软峰值与加权均值之差、跨模态关系统计量、匹配置信度和区域面积占比共同确定；在不存在局部区域损毁标注时，对同一预测建筑实例的多个局部区域证据表示进行注意力聚合得到建筑级预测，并利用建筑级损毁标签训练局部区域证据表示和条件损毁专家。",
        [],
    ),
    (
        "5.根据权利要求1所述的方法，其特征在于，所述建筑级粗检索采用建筑内损毁证据结构的检索表示计算相似度，所述关键局部区域双向匹配重排序分别计算当前建筑到历史建筑以及历史建筑到当前建筑的最大局部支持并进行组合；所述检索可靠度由重排序后的检索相似度、历史邻居类别一致性、历史邻居所覆盖的灾害事件数量与历史邻居数量之比以及检索得分方差共同形成；所述残差效用验证模块以候选修正前后的分类损失差是否大于预设效用增量作为监督信号进行训练。",
        [],
    ),
    (
        "6.一种多模态建筑损毁评估系统，其特征在于，包括基础预测模块、同建筑硬约束局部区域生成模块、建筑内跨模态搜索和有界残差关联模块、建筑内损毁证据结构构建模块、条件损毁专家模块、同事件排除的跨事件关系记忆检索模块、残差效用验证模块以及选择性末端修正模块；各模块被配置为协同执行权利要求1至5任一项所述的方法。",
        [],
    ),
    (
        "7.一种计算机可读存储介质，其上存储有计算机程序，其特征在于，所述计算机程序被处理器执行时实现权利要求1至5任一项所述的方法。",
        [],
    ),
]


class EquationCollector:
    SENTINEL = "¤" * 128

    def __init__(self) -> None:
        self.items: list[dict[str, str]] = []

    def add(self, doc: Document, marker, linear: str) -> None:
        placeholder = f"[[EQ{len(self.items) + 1:04d}]]"
        # Word's professional equation builder may temporarily expand a math range
        # while converting linear delimiters.  A same-paragraph sentinel buffer
        # prevents that expansion from touching the following equation paragraph.
        p = base.add_text(doc, marker, placeholder + self.SENTINEL, role="formula")
        p.paragraph_format.keep_together = True
        p.paragraph_format.keep_with_next = False
        self.items.append({"placeholder": placeholder, "linear": linear})


def _center_text(draw: ImageDraw.ImageDraw, box, text: str, size: int = 38, *, bold=False) -> None:
    font = flow.img_font(size, bold=bold)
    wrapped = flow.split_text(draw, text, font, box[2] - box[0] - 28)
    tw, th = flow.text_size(draw, wrapped, font, spacing=8)
    draw.multiline_text(
        ((box[0] + box[2] - tw) / 2, (box[1] + box[3] - th) / 2),
        wrapped,
        font=font,
        fill=flow.INK,
        spacing=8,
        align="center",
    )


def _panel(draw: ImageDraw.ImageDraw, box, title: str) -> None:
    draw.rounded_rectangle(box, radius=24, outline=flow.INK, width=5)
    title_box = (box[0], box[1], box[2], box[1] + 105)
    draw.line([(box[0], title_box[3]), (box[2], title_box[3])], fill=flow.INK, width=4)
    _center_text(draw, title_box, title, size=38, bold=True)


def _dashed_line(draw: ImageDraw.ImageDraw, start, end, *, width=4, dash=22, gap=14) -> None:
    x1, y1 = start
    x2, y2 = end
    length = max(((x2 - x1) ** 2 + (y2 - y1) ** 2) ** 0.5, 1)
    ux, uy = (x2 - x1) / length, (y2 - y1) / length
    pos = 0.0
    while pos < length:
        stop = min(pos + dash, length)
        draw.line(
            [(x1 + ux * pos, y1 + uy * pos), (x1 + ux * stop, y1 + uy * stop)],
            fill=flow.INK,
            width=width,
        )
        pos += dash + gap


def _draw_diamond(draw: ImageDraw.ImageDraw, center, size, text: str, font_size=34) -> tuple[int, int, int, int]:
    cx, cy = center
    w, h = size
    points = [(cx, cy - h // 2), (cx + w // 2, cy), (cx, cy + h // 2), (cx - w // 2, cy)]
    draw.polygon(points, fill=flow.PAPER, outline=flow.INK)
    draw.line(points + [points[0]], fill=flow.INK, width=5)
    box = (cx - w // 3, cy - h // 4, cx + w // 3, cy + h // 4)
    _center_text(draw, box, text, size=font_size, bold=True)
    return (cx - w // 2, cy - h // 2, cx + w // 2, cy + h // 2)


def _draw_cylinder(draw: ImageDraw.ImageDraw, box, text: str) -> None:
    x1, y1, x2, y2 = box
    cap_h = 55
    draw.rectangle((x1, y1 + cap_h // 2, x2, y2 - cap_h // 2), fill=flow.PAPER, outline=flow.INK, width=5)
    draw.ellipse((x1, y1, x2, y1 + cap_h), fill=flow.PAPER, outline=flow.INK, width=5)
    draw.arc((x1, y2 - cap_h, x2, y2), 0, 180, fill=flow.INK, width=5)
    _center_text(draw, (x1 + 18, y1 + 48, x2 - 18, y2 - 25), text, size=34, bold=True)


def draw_figure_2_architecture() -> Path:
    image = Image.new("RGB", (flow.CANVAS_W, flow.CANVAS_H), flow.PAPER)
    draw = ImageDraw.Draw(image)

    structure_panel = (70, 100, 790, 1040)
    search_panel = (70, 1160, 790, 2200)
    association_panel = (850, 100, 1445, 2200)
    evidence_panel = (1500, 100, 1930, 2200)
    _panel(draw, structure_panel, "同建筑硬约束")
    _panel(draw, search_panel, "建筑内跨模态搜索")
    _panel(draw, association_panel, "结构关联与有界残差")
    _panel(draw, evidence_panel, "局部证据聚合")

    # Building instance geometry and forbidden cross-building association.
    bi = [(145, 325), (280, 235), (530, 270), (670, 430), (590, 760), (330, 850), (150, 675)]
    bj = [(520, 735), (685, 700), (730, 900), (565, 955), (455, 865)]
    draw.polygon(bi, fill=flow.PAPER, outline=flow.INK)
    draw.line(bi + [bi[0]], fill=flow.INK, width=6)
    draw.polygon(bj, fill=flow.PAPER, outline=flow.INK)
    draw.line(bj + [bj[0]], fill=flow.INK, width=4)
    draw.text((175, 350), "建筑实例 B_i", font=flow.img_font(36, bold=True), fill=flow.INK)
    draw.text((535, 865), "B_j", font=flow.img_font(34, bold=True), fill=flow.INK)
    p = (360, 520)
    centers = [(245, 420), (510, 430), (400, 720)]
    draw.ellipse((p[0] - 12, p[1] - 12, p[0] + 12, p[1] + 12), fill=flow.INK)
    draw.text((p[0] + 18, p[1] - 24), "p", font=flow.img_font(32), fill=flow.INK)
    for idx, c in enumerate(centers, 1):
        draw.ellipse((c[0] - 19, c[1] - 19, c[0] + 19, c[1] + 19), fill=flow.PAPER, outline=flow.INK, width=5)
        draw.text((c[0] + 24, c[1] - 23), f"c_i{idx}", font=flow.img_font(29), fill=flow.INK)
        flow.draw_arrow(draw, p, c, width=3, head=16)
    forbidden = (585, 835)
    _dashed_line(draw, p, forbidden, width=3)
    draw.line([(555, 790), (610, 845)], fill=flow.INK, width=7)
    draw.line([(610, 790), (555, 845)], fill=flow.INK, width=7)
    draw.text((175, 925), "同建筑候选 M_pik=1    跨建筑候选 M_pjk=0", font=flow.img_font(28), fill=flow.INK)

    # Optical/SAR grids and local search window.
    grid_y = 1380
    for gx, label in [(135, "光学结构 F_rgb"), (470, "雷达响应 F_sar")]:
        draw.text((gx + 40, 1280), label, font=flow.img_font(31, bold=True), fill=flow.INK)
        for row in range(5):
            for col in range(5):
                x1 = gx + col * 54
                y1 = grid_y + row * 54
                draw.rectangle((x1, y1, x1 + 54, y1 + 54), outline=flow.INK, width=2)
    draw.ellipse((135 + 2 * 54 + 17, grid_y + 2 * 54 + 17, 135 + 2 * 54 + 37, grid_y + 2 * 54 + 37), fill=flow.INK)
    draw.text((260, 1615), "p", font=flow.img_font(31), fill=flow.INK)
    draw.rectangle((470 + 54, grid_y + 54, 470 + 4 * 54, grid_y + 4 * 54), outline=flow.INK, width=6)
    for qx, qy in [(470 + 2 * 54 + 27, grid_y + 2 * 54 + 27), (470 + 3 * 54 + 27, grid_y + 54 + 27)]:
        draw.ellipse((qx - 10, qy - 10, qx + 10, qy + 10), fill=flow.INK)
    flow.draw_arrow(draw, (405, 1515), (470, 1515), width=5, head=22)
    draw.text((235, 1745), "搜索域 N_i(p)=B_i∩邻域", font=flow.img_font(31, bold=True), fill=flow.INK)
    _center_text(draw, (145, 1835, 715, 2085), "相似度加权得到匹配雷达响应\n同时输出匹配置信度 C_p 与不确定度 H_p", size=32)

    # Parallel structure and change branches followed by a masked association.
    boxes = {
        "struct": (915, 300, 1380, 560),
        "search": (915, 720, 1380, 980),
        "res": (915, 1140, 1380, 1400),
        "mask": (915, 1600, 1380, 1880),
    }
    flow.draw_box(draw, boxes["struct"], "光学结构关联\nLstruct", max_font=39, min_font=30)
    flow.draw_box(draw, boxes["search"], "匹配关系特征\nR_p、C_p、H_p", max_font=38, min_font=29)
    flow.draw_box(draw, boxes["res"], "零初始化有界残差\nα·ΔLchange", max_font=38, min_font=29)
    flow.draw_box(draw, boxes["mask"], "同建筑掩膜归一化\n软关联 Q_pik", max_font=38, min_font=29)
    flow.draw_arrow(draw, (1148, 560), (1148, 720), width=5)
    flow.draw_arrow(draw, (1148, 980), (1148, 1140), width=5)
    flow.draw_arrow(draw, (1148, 1400), (1148, 1600), width=5)
    draw.text((1240, 1480), "Lfinal", font=flow.img_font(30, bold=True), fill=flow.INK)
    flow.draw_arrow(draw, (790, 520), (915, 430), width=5)
    flow.draw_arrow(draw, (790, 1660), (915, 850), width=5)

    # Region tokens converge into a building evidence structure.
    token_boxes = [(1560, 300, 1870, 470), (1560, 560, 1870, 730), (1560, 820, 1870, 990)]
    for idx, box in enumerate(token_boxes, 1):
        label = f"局部区域证据 z_i{idx}" if idx < 3 else "关键局部区域 z_iM"
        flow.draw_box(draw, box, label, max_font=33, min_font=27)
        flow.draw_arrow(draw, (1380, 1740), (box[0], (box[1] + box[3]) / 2), width=3, head=17)
    merge = _draw_diamond(draw, (1715, 1260), (330, 300), "均值\n离散度\n软峰值", font_size=31)
    for box in token_boxes:
        flow.draw_arrow(draw, ((box[0] + box[2]) / 2, box[3]), (1715, merge[1]), width=3, head=16)
    output_box = (1545, 1580, 1885, 1950)
    flow.draw_box(draw, output_box, "建筑内损毁\n证据结构 Z_i", max_font=39, min_font=30, bold=True)
    flow.draw_arrow(draw, (1715, merge[3]), (1715, output_box[1]), width=5)
    _center_text(
        draw,
        (1535, 2020, 1900, 2170),
        "避免关键证据被平均稀释",
        size=25,
        bold=True,
    )

    path = FIG_DIR / "figure_2_local_evidence.png"
    image.save(path, dpi=(300, 300), optimize=True)
    return path


def draw_figure_3_architecture() -> Path:
    image = Image.new("RGB", (flow.CANVAS_W, flow.CANVAS_H), flow.PAPER)
    draw = ImageDraw.Draw(image)

    # Two evidence branches.
    _panel(draw, (70, 100, 770, 980), "当前建筑证据分支")
    _panel(draw, (70, 1160, 770, 2200), "跨事件关系记忆分支")
    flow.draw_box(draw, (135, 285, 705, 500), "疑难建筑候选\n基础预测为完整或完整/损毁混淆", max_font=36, min_font=28)
    flow.draw_box(draw, (135, 655, 705, 875), "条件损毁专家\n疑似损毁分数 sDI 与当前证据裕量 dcur", max_font=36, min_font=28)
    flow.draw_arrow(draw, (420, 500), (420, 655), width=5)

    _draw_cylinder(draw, (120, 1325, 400, 1605), "跨事件\n关系记忆库")
    filter_box = _draw_diamond(draw, (585, 1470), (300, 250), "排除同事件\ne_j≠e_i", font_size=31)
    flow.draw_arrow(draw, (400, 1470), (filter_box[0], 1470), width=5)
    flow.draw_box(draw, (135, 1760, 705, 1950), "建筑级粗检索 + 关键区域双向重排序", max_font=35, min_font=28)
    flow.draw_box(draw, (135, 2015, 705, 2155), "历史证据裕量 dmem 与检索可靠度", max_font=33, min_font=27)
    flow.draw_poly_arrow(draw, [(585, filter_box[3]), (585, 1695), (420, 1695), (420, 1760)], width=5)
    flow.draw_arrow(draw, (420, 1950), (420, 2015), width=5)

    # Four mutually exclusive relation states.
    relation_outer = (835, 360, 1350, 1530)
    _panel(draw, relation_outer, "互斥关系编码")
    relation_boxes = [
        ((900, 560, 1285, 735), "强证据方向一致"),
        ((900, 800, 1285, 975), "仅当前证据达到阈值"),
        ((900, 1040, 1285, 1215), "仅历史证据达到阈值"),
        ((900, 1280, 1285, 1455), "强证据方向冲突"),
    ]
    for box, label in relation_boxes:
        flow.draw_box(draw, box, label, max_font=33, min_font=27)
    flow.draw_poly_arrow(draw, [(770, 770), (815, 770), (815, 675), (900, 675)], width=4)
    flow.draw_poly_arrow(draw, [(770, 2085), (810, 2085), (810, 1370), (900, 1370)], width=4)
    draw.text((860, 1645), "关系向量 r_i", font=flow.img_font(34, bold=True), fill=flow.INK)

    # Utility gate and two terminal branches.
    utility_box = (1450, 280, 1910, 650)
    flow.draw_box(
        draw,
        utility_box,
        "残差效用验证\n关系 + 基础不确定度\n历史支持差 + 检索可靠度\n输出效用值 u_i",
        max_font=34,
        min_font=27,
        bold=True,
    )
    flow.draw_arrow(draw, (1350, 900), (1450, 500), width=5)
    flow.draw_poly_arrow(draw, [(705, 2050), (1400, 2050), (1400, 570), (1450, 570)], width=4)
    decision = _draw_diamond(draw, (1680, 980), (500, 400), "Base=I\nsDI≥τs\nu_i≥τu", font_size=34)
    flow.draw_arrow(draw, (1680, utility_box[3]), (1680, decision[1]), width=5)
    correct_box = (1425, 1350, 1745, 1615)
    keep_box = (1755, 1350, 1950, 1615)
    flow.draw_box(draw, correct_box, "仅在目标掩膜内\n执行 I→D", max_font=35, min_font=28, bold=True)
    flow.draw_box(draw, keep_box, "保持基础预测", max_font=33, min_font=27, bold=True)
    flow.draw_poly_arrow(draw, [(1570, decision[3] - 15), (1570, 1270), (1585, correct_box[1])], width=5)
    flow.draw_poly_arrow(draw, [(1790, decision[3] - 15), (1790, 1270), (1850, keep_box[1])], width=5)
    draw.text((1515, 1260), "是", font=flow.img_font(30, bold=True), fill=flow.INK)
    draw.text((1825, 1260), "否", font=flow.img_font(30, bold=True), fill=flow.INK)
    final_box = (1460, 1810, 1920, 2070)
    flow.draw_box(draw, final_box, "最终建筑损毁结果", max_font=40, min_font=31, bold=True)
    flow.draw_arrow(draw, ((correct_box[0] + correct_box[2]) / 2, correct_box[3]), (1600, final_box[1]), width=5)
    flow.draw_arrow(draw, ((keep_box[0] + keep_box[2]) / 2, keep_box[3]), (1800, final_box[1]), width=5)
    lock_box = (880, 1775, 1320, 2110)
    flow.draw_box(draw, lock_box, "基础网络中间特征\n全程保持不变\n其他类别禁止改写", max_font=35, min_font=28)
    _dashed_line(draw, (1320, 1940), (1460, 1940), width=4)

    path = FIG_DIR / "figure_3_memory_utility.png"
    image.save(path, dpi=(300, 300), optimize=True)
    return path


def draw_figures() -> None:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    flow.FIG_DIR = FIG_DIR
    flow.draw_vertical_steps(
        "figure_1_overall.png",
        [
            ("S1", "获取灾前光学影像与灾后合成孔径雷达影像，并进行预处理"),
            ("S2", "输入基础评估网络，获得建筑位置、基础损毁类别及多模态特征"),
            ("S3", "在每个预测建筑内生成自适应局部区域及硬约束候选集合"),
            ("S4", "执行建筑内跨模态搜索，以零初始化有界残差形成软关联"),
            ("S5", "聚合局部统计、软峰值与关键区域，构建建筑内损毁证据结构"),
            ("S6", "条件损毁专家产生当前建议，关系记忆排除同事件后产生历史建议"),
            ("S7", "编码当前建议与历史建议的四类关系并估计残差效用"),
            ("S8", "仅在效用满足条件时将完整类别选择性修正为损毁类别"),
        ],
        max_font=41,
        min_font=28,
    )
    draw_figure_2_architecture()
    draw_figure_3_architecture()


CORE_FORMULAS = [
    "K_i=clip(⌈A_i/A_0⌉,1,K_(max))",
    "R_(ik)⊆B_i,   R_(ik)∩B_j=∅  (i≠j)",
    "N_i^(sar)(p)=B_i∩{q:‖q-p‖_∞≤r_s}",
    "F̃^(sar)(p)=∑_(q∈N_i^(sar)(p)) a_(pq)F^(sar)(q)",
    "L_(pik)^(final)=L_(pik)^(struct)+α_(pik)ΔL_(pik)^(change)",
    "α_(pik)=α_(max)tanh(g)σ(v_(pik)),   |α_(pik)|≤α_(max),   g=0",
    "Q_(pik)=(M_(pik) e^(L_(pik)^(final)/T_r))/(∑_(j=1)^(K_i) M_(pij) e^(L_(pij)^(final)/T_r)+ε)",
    "ΔE_(ik)=E_(ik)^(peak)-E_(ik)^(mean)",
    "Δ_i^(mem)=S_i^D-S_i^I",
    "G_i=𝕀[ŷ_i^(base)=I ∧ s_i^(DI)≥τ_s ∧ u_i≥τ_u]",
]


DETAIL_FORMULAS: dict[int, list[str]] = {
    43: [
        "I_(rgb)^(pre)∈ℝ^(H×W×C_r),   I_(sar)^(post)∈ℝ^(H×W×C_s)",
        "(F^(rgb),F^(sar),L^(loc),L^(base))=F_θ(I_(rgb)^(pre),I_(sar)^(post))",
        "P^(building)(p)=Softmax(L^(loc)(p))_(building)",
        "P_c^(base)(p)=e^(L_c^(base)(p))/(∑_(k=0)^3 e^(L_k^(base)(p)))",
    ],
    44: [
        "M^(building)(p)=𝕀[P^(building)(p)≥τ_b]",
        "{B_1,…,B_N}=CC(M^(building)),   |B_i|≥A_(min)",
    ],
    45: [
        "A_i=|B_i|",
        "K_i=clip(⌈A_i/A_0⌉,1,K_(max))",
    ],
    46: [
        "C_i={c_(i1),…,c_(iK_i)}",
        "C_i(p)={c_(ik)∈C_i:‖p-c_(ik)‖_2≤r_c}",
        "M_(pik)=𝕀[p∈B_i ∧ c_(ik)∈B_i ∧ ‖p-c_(ik)‖_2≤r_c]",
        "M_(pik)=0 ⇒ Q_(pik)=0",
    ],
    47: [
        "Δx_(pik)=(x_p-x_(c_ik))/W,   Δy_(pik)=(y_p-y_(c_ik))/H",
        "L_(pik)^(struct)=φ_s(F^(rgb)(p),F^(rgb)(c_(ik)),Δx_(pik),Δy_(pik),E^(rgb)(p))",
        "L_(pik)^(struct)=-λ_f‖W_rF^(rgb)(p)-W_rF^(rgb)(c_(ik))‖_2^2-λ_s‖p-c_(ik)‖_2^2",
    ],
    48: [
        "N_i^(sar)(p)=B_i∩{q:‖q-p‖_∞≤r_s}",
        "s_(pq)=((W_rF^(rgb)(p))^T(W_sF^(sar)(q)))/(‖W_rF^(rgb)(p)‖_2‖W_sF^(sar)(q)‖_2)",
        "a_(pq)=e^(s_(pq)/T_a)/(∑_(q'∈N_i^(sar)(p)) e^(s_(pq')/T_a))",
        "F̃^(sar)(p)=∑_(q∈N_i^(sar)(p)) a_(pq) F^(sar)(q)",
    ],
    49: [
        "H_p=-(∑_q a_(pq)log(a_(pq)+ε))/log|N_i^(sar)(p)|",
        "C_p=1-H_p,   U_p^(base)=1-max_c P_c^(base)(p)",
        "R_p=Concat[F^(rgb)(p),F̃^(sar)(p),|F^(rgb)(p)-F̃^(sar)(p)|,C_p,H_p]",
    ],
    50: [
        "ΔL_(pik)^(change)=φ_Δ(R_p,F^(rgb)(c_(ik)))",
        "α_(pik)=α_(max)tanh(g)σ(v_(pik)),   |α_(pik)|≤α_(max)",
        "g|_(t=0)=0",
        "L_(pik)^(final)=L_(pik)^(struct)+α_(pik)ΔL_(pik)^(change)",
    ],
    51: [
        "Q_(pik)=(M_(pik) e^(L_(pik)^(final)/T_r))/(∑_(j=1)^(K_i) M_(pij) e^(L_(pij)^(final)/T_r)+ε)",
        "∑_(k=1)^(K_i) Q_(pik)=1,   p∈B_i",
    ],
    52: [
        "m_(ik)=∑_(p∈B_i) Q_(pik),   ρ_(ik)=m_(ik)/(|B_i|+ε)",
        "z_(ik)^(rgb)=(∑_(p∈B_i) Q_(pik) F^(rgb)(p))/(m_(ik)+ε)",
        "z_(ik)^(sar)=(∑_(p∈B_i) Q_(pik) F̃^(sar)(p))/(m_(ik)+ε)",
        "μ_(ik)^(rel)=(∑_(p∈B_i) Q_(pik) R_p)/(m_(ik)+ε)",
        "C̄_(ik)=(∑_(p∈B_i) Q_(pik) C_p)/(m_(ik)+ε)",
    ],
    53: [
        "e_p=ψ(R_p)",
        "ω_(pik)=(Q_(pik) e^(e_p/T_p))/(∑_(q∈B_i) Q_(qik) e^(e_q/T_p)+ε)",
        "E_(ik)^(peak)=∑_(p∈B_i) ω_(pik) e_p",
        "E_(ik)^(mean)=(∑_(p∈B_i) Q_(pik) e_p)/(m_(ik)+ε)",
        "ΔE_(ik)=E_(ik)^(peak)-E_(ik)^(mean)",
    ],
    54: [
        "z_(ik)=Φ(z_(ik)^(rgb),z_(ik)^(sar),μ_(ik)^(rel),C̄_(ik),ρ_(ik),ΔE_(ik))",
        "κ_(ik)=C̄_(ik)(β_1|ΔE_(ik)|+β_2‖μ_(ik)^(rel)‖_1+β_3ρ_(ik))",
        "μ_i^R=(1/K_i) ∑_(k=1)^(K_i) z_(ik)",
        "σ_i^R=√((1/K_i) ∑_(k=1)^(K_i) (z_(ik)-μ_i^R)^2)",
        "Z_i=Concat[μ_i^R,σ_i^R,z_(i,t_1),…,z_(i,t_M)]",
    ],
    55: [
        "a_(ik)^(bag)=e^(φ_b(z_(ik)))/(∑_(j=1)^(K_i) e^(φ_b(z_(ij))))",
        "o_i^(bag)=∑_(k=1)^(K_i) a_(ik)^(bag) o_(ik)",
        "p_i^(bag)=Softmax(o_i^(bag))",
        "L_(bag)=-∑_i log p_(i,y_i)^(bag)",
    ],
    56: [
        "p_(ic)^(base)=(1/|B_i|) ∑_(p∈B_i) P_c^(base)(p)",
        "ŷ_i^(base)=argmax_c p_(ic)^(base),   U_i=1-max_c p_(ic)^(base)",
        "H={B_i:ŷ_i^(base)=I}∪{B_i:|p_(iI)^(base)-p_(iD)^(base)|≤τ_(ID)}",
    ],
    57: [
        "x_i^(expert)=Concat[p_i^(base),U_i,Z_i]",
        "s_i^(DI)=σ(f_(θ_e)(x_i^(expert))),   r_i^(cur)=𝕀[s_i^(DI)≥τ_s]",
    ],
    59: [
        "w_i=1/((N_(y_i)+ε)^(γ_c)(N_(e_i,y_i)+ε)^(γ_e))",
        "w̃_i=w_i/((1/N) ∑_(j=1)^N w_j)",
        "L_(expert)=-∑_i w̃_i[y_i^(DI)log s_i^(DI)+(1-y_i^(DI))log(1-s_i^(DI))]",
    ],
    60: [
        "M_j=(k_j,T_j,y_j,e_j,u_j^(hist))",
        "k_j=φ_k(Z_j),   T_j={z_(j,t_1),…,z_(j,t_M)}",
    ],
    61: [
        "S_(ij)^(global)=cos(k_i,k_j)",
        "S̄_(ij)^(global)=S_(ij)^(global)𝕀[e_j≠e_i]",
        "J_i=TopR_j(S̄_(ij)^(global))",
    ],
    62: [
        "S_(ij)^(q→m)=(1/M)∑_(a=1)^M max_b cos(z_(i,t_a),z_(j,t_b))",
        "S_(ij)^(m→q)=(1/M)∑_(b=1)^M max_a cos(z_(j,t_b),z_(i,t_a))",
        "S_(ij)=λ_gS_(ij)^(global)+λ_l(S_(ij)^(q→m)+S_(ij)^(m→q))/2",
    ],
    63: [
        "S_i^D=(∑_(j∈J_i) S_(ij) 𝕀[y_j=D])/(∑_(j∈J_i) |S_(ij)|+ε)",
        "S_i^I=(∑_(j∈J_i) S_(ij) 𝕀[y_j=I])/(∑_(j∈J_i) |S_(ij)|+ε)",
        "Δ_i^(mem)=S_i^D-S_i^I,   r_i^(mem)=𝕀[Δ_i^(mem)≥τ_m]",
    ],
    64: [
        "d_i^(cur)=2s_i^(DI)-1,   d_i^(mem)=Δ_i^(mem)",
        "R_i^(agree)=𝕀[d_i^(cur)d_i^(mem)>0 ∧ |d_i^(cur)|≥τ_c ∧ |d_i^(mem)|≥τ_m]",
        "R_i^(cur)=𝕀[|d_i^(cur)|≥τ_c ∧ |d_i^(mem)|<τ_m]",
        "R_i^(mem)=𝕀[|d_i^(cur)|<τ_c ∧ |d_i^(mem)|≥τ_m]",
        "R_i^(conflict)=𝕀[d_i^(cur)d_i^(mem)<0 ∧ |d_i^(cur)|≥τ_c ∧ |d_i^(mem)|≥τ_m]",
        "r_i=[R_i^(agree),R_i^(cur),R_i^(mem),R_i^(conflict)]",
    ],
    66: [
        "A_i^(cls)=max(S_i^D,S_i^I),   D_i^(event)=|{e_j:j∈J_i}|/|J_i|",
        "V_i^(retr)=Var({S_(ij):j∈J_i})",
        "x_i^(util)=Concat[s_i^(DI),Δ_i^(mem),r_i,U_i,A_i^(cls),D_i^(event),V_i^(retr)]",
        "u_i=σ(f_(θ_u)(x_i^(util)))",
    ],
    67: [
        "ΔU_i=ℓ(ŷ_i^(base),y_i)-ℓ(ŷ_i^(corr),y_i)",
        "y_i^(util)=𝕀[ΔU_i>δ_u]",
        "L_(util)=-∑_i [y_i^(util) log u_i+(1-y_i^(util)) log(1-u_i)]",
    ],
    68: [
        "G_i=𝕀[ŷ_i^(base)=I ∧ s_i^(DI)≥τ_s ∧ u_i≥τ_u]",
    ],
    69: [
        "Ω_i^R=⋃_(k:κ_(ik)≥τ_κ){p∈B_i:Q_(pik)≥τ_Q}",
        "Ω_i∈{B_i,Ω_i^R}",
    ],
    70: [
        "G_i=1 ∧ ŷ_p^(base)=I ⇒ ŷ_p^(final)=D,   p∈Ω_i",
        "G_i=0 ∨ ŷ_p^(base)≠I ⇒ ŷ_p^(final)=ŷ_p^(base)",
    ],
    71: [
        "L_(compact)=∑_(i,k,p) Q_(pik) ‖p-c_(ik)‖_2^2",
        "L_(balance)=∑_i ∑_k (m_(ik)/|B_i|-1/K_i)^2",
        "L_(res)=∑_(i,k,p) |α_(pik)ΔL_(pik)^(change)|",
        "L_(total)=L_(base)+λ_b L_(bag)+λ_e L_(expert)+λ_u L_(util)+λ_c L_(compact)+λ_q L_(balance)+λ_r L_(res)",
    ],
}


def add_equation_intro(doc: Document, marker, text: str) -> None:
    p = base.add_text(doc, marker, text, role="body")
    p.paragraph_format.keep_with_next = True


def build_document() -> tuple[Path, int]:
    if len(re.sub(r"\s", "", base.ABSTRACT)) > 300:
        raise ValueError("Abstract exceeds 300 Chinese characters")
    for path in (TEMPLATE, base.SOURCE_PATENT):
        if not path.exists():
            raise FileNotFoundError(path)
    for filename, _ in FIGURES:
        if not (FIG_DIR / filename).exists():
            raise FileNotFoundError(FIG_DIR / filename)

    doc = Document(TEMPLATE)
    markers = base.clear_template_body(doc)
    equations = EquationCollector()

    # 第一节：说明书摘要。
    base.add_text(doc, markers[0], base.ABSTRACT, role="body")

    # 第二节：摘要附图。
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    base.style_paragraph(p, "figure")
    p.paragraph_format.space_before = base.Pt(20)
    p.add_run().add_picture(str(FIG_DIR / FIGURES[0][0]), width=Inches(6.55))
    base.move_before(p, markers[1])

    # 第三节：权利要求书。保留七项精简权利要求，不在权利要求中堆叠推导公式。
    for head, steps in REVISED_CLAIMS:
        base.add_text(doc, markers[2], head, role="claim")
        for step in steps:
            base.add_text(doc, markers[2], step, role="claim")

    # 第四节：说明书。
    base.add_text(doc, markers[3], base.TITLE, role="title")
    items = base.extract_spec_items()
    for index, (role, text) in enumerate(items):
        if 26 <= index <= 29:
            continue
        if index == 39:
            text = (
                "图1为本发明多模态建筑损毁评估方法的总体技术流程示意图；"
                "图2为建筑内局部证据构建的空间约束和信息聚合结构示意图；"
                "图3为跨事件关系记忆与效用门控的双分支决策结构示意图。"
            )
            role = "body"
        base.add_text(doc, markers[3], text, role=role)

        if index == 25:
            add_equation_intro(
                doc,
                markers[3],
                "为锁定上述技术方案的必要边界，建筑内局部区域、跨模态残差、软关联、局部强响应、关系记忆及效用门控满足下列核心关系：",
            )
            for formula in CORE_FORMULAS:
                equations.add(doc, markers[3], formula)
            base.add_text(
                doc,
                markers[3],
                "其中，B_i为第i个建筑实例，A_i为其面积，K_i为局部区域数量，M为同建筑硬约束，Q为区域软关联，ΔE为局部强响应相对平均响应的增量，Δ_i^(mem)表示跨事件历史损毁支持相对完整支持的差值，G_i为末端修正门。由此，跨建筑关联被硬排除，雷达信息仅以有界残差参与区域生成，关系记忆仅在效用验证通过时改变最终决策。",
                role="body",
            )

        if index == 41:
            add_equation_intro(
                doc,
                markers[3],
                "下列实施例对符号、约束和训练目标作进一步数学定义。除特别说明外，p、q表示像素位置，i、j表示建筑实例，k表示建筑内局部区域，I、D分别表示完整类别和损毁类别，ε为防止分母为零的正常数。",
            )

        if index in DETAIL_FORMULAS:
            for formula in DETAIL_FORMULAS[index]:
                equations.add(doc, markers[3], formula)

    # 第五节：说明书附图。严格控制为三幅主图，每页一幅。
    for index, (filename, caption) in enumerate(FIGURES):
        base.add_figure(
            doc,
            None,
            FIG_DIR / filename,
            caption,
            width=6.45,
            add_page_break=index < len(FIGURES) - 1,
        )

    doc.core_properties.title = base.TITLE
    doc.core_properties.subject = "中国发明专利申请文件"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    EQUATION_JSON.write_text(
        json.dumps(equations.items, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return OUTPUT, len(equations.items)


def main() -> None:
    draw_figures()
    output, equation_count = build_document()
    print(output)
    print(EQUATION_JSON)
    print(f"equations={equation_count}")
    print(f"claims={len(REVISED_CLAIMS)}")
    print(f"figures={len(FIGURES)}")


if __name__ == "__main__":
    main()
