from __future__ import annotations

from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BUILD_DIR = ROOT / ".patent_build"
FIG_DIR = BUILD_DIR / "figures"
OUT_DIR = ROOT / "deliverables"
OUT_PATH = OUT_DIR / "建筑内局部证据与跨事件效用验证的建筑损毁评估发明专利初稿.docx"


def font_path() -> str:
    candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\simsun.ttc"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return str(candidate)
    raise FileNotFoundError("No CJK font found")


FONT_FILE = font_path()


def img_font(size: int, bold: bool = False):
    bold_candidates = [Path(r"C:\Windows\Fonts\msyhbd.ttc"), Path(r"C:\Windows\Fonts\simhei.ttf")]
    if bold:
        for p in bold_candidates:
            if p.exists():
                return ImageFont.truetype(str(p), size)
    return ImageFont.truetype(FONT_FILE, size)


def text_bbox(draw: ImageDraw.ImageDraw, text: str, font):
    box = draw.multiline_textbbox((0, 0), text, font=font, spacing=8, align="center")
    return box[2] - box[0], box[3] - box[1]


def fit_text(draw, text, box, max_size=34, min_size=20, bold=False):
    x1, y1, x2, y2 = box
    width, height = x2 - x1 - 30, y2 - y1 - 24
    for size in range(max_size, min_size - 1, -1):
        font = img_font(size, bold=bold)
        tw, th = text_bbox(draw, text, font)
        if tw <= width and th <= height:
            return font
    return img_font(min_size, bold=bold)


def box(draw, xy, text, fill="#FFFFFF", width=3, radius=14, bold=False):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline="#111111", width=width)
    font = fit_text(draw, text, xy, bold=bold)
    tw, th = text_bbox(draw, text, font)
    x1, y1, x2, y2 = xy
    draw.multiline_text(((x1 + x2 - tw) / 2, (y1 + y2 - th) / 2 - 2), text,
                        font=font, fill="#111111", spacing=8, align="center")


def arrow(draw, start, end, width=4):
    draw.line([start, end], fill="#111111", width=width)
    x1, y1 = start
    x2, y2 = end
    dx, dy = x2 - x1, y2 - y1
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    size = 15
    tip = (x2, y2)
    left = (x2 - ux * size + px * size * 0.55, y2 - uy * size + py * size * 0.55)
    right = (x2 - ux * size - px * size * 0.55, y2 - uy * size - py * size * 0.55)
    draw.polygon([tip, left, right], fill="#111111")


def label(draw, xy, text, size=28, bold=False, anchor="mm"):
    draw.text(xy, text, fill="#111111", font=img_font(size, bold=bold), anchor=anchor,
              align="center", spacing=7)


def make_canvas():
    return Image.new("RGB", (1600, 900), "white")


def save_canvas(img, name):
    path = FIG_DIR / name
    img.save(path, dpi=(220, 220))
    return path


def draw_figures():
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    paths = []

    # Figure 1: overall pipeline.
    img = make_canvas(); d = ImageDraw.Draw(img)
    label(d, (800, 60), "整体技术流程", 42, True)
    coords = [(70, 150, 450, 290), (610, 150, 990, 290), (1150, 150, 1530, 290),
              (1150, 530, 1530, 670), (610, 530, 990, 670), (70, 530, 450, 670)]
    texts = ["灾前光学影像\n灾后 SAR 影像", "基础网络\n建筑与损毁预测", "同楼硬约束\n局部区域生成",
             "建筑内证据结构\n条件损毁专家", "跨事件病例验证\n效用门控", "选择性末端修正\n仅 Intact → Damaged"]
    fills = ["#FFFFFF", "#F2F2F2", "#FFFFFF", "#F2F2F2", "#FFFFFF", "#F2F2F2"]
    for c, t, f in zip(coords, texts, fills): box(d, c, t, fill=f, bold=True)
    arrow(d, (450, 220), (610, 220)); arrow(d, (990, 220), (1150, 220))
    arrow(d, (1340, 290), (1340, 530)); arrow(d, (1150, 600), (990, 600)); arrow(d, (610, 600), (450, 600))
    label(d, (800, 790), "基础网络中间特征保持不变   附加证据只决定是否修改最终决策", 30, True)
    paths.append(save_canvas(img, "figure_1_overall.png"))

    # Figure 2: building-first regions.
    img = make_canvas(); d = ImageDraw.Draw(img)
    label(d, (800, 55), "预测建筑约束的局部区域生成", 40, True)
    box(d, (90, 145, 460, 690), "预测建筑掩膜\n\nB1      B2\n\n各建筑独立连通", "#FFFFFF", bold=True)
    box(d, (615, 145, 985, 690), "按建筑面积\n自适应设置区域中心\n\n每栋建筑独立估计 K\n\n过小区域邻接合并", "#F2F2F2", bold=True)
    box(d, (1140, 145, 1510, 690), "像素候选集合\n\n仅包含同一建筑\n附近的区域中心\n\n跨建筑关联概率恒为 0", "#FFFFFF", bold=True)
    arrow(d, (460, 420), (615, 420)); arrow(d, (985, 420), (1140, 420))
    label(d, (800, 805), "硬约束来自候选集合，而非依赖损失函数事后抑制", 30, True)
    paths.append(save_canvas(img, "figure_2_building_regions.png"))

    # Figure 3: local cross-modal search and bounded residual.
    img = make_canvas(); d = ImageDraw.Draw(img)
    label(d, (800, 55), "建筑内跨模态搜索与有界修边", 40, True)
    box(d, (55, 150, 375, 305), "光学结构特征\n形成基础关联 Lstruct", "#FFFFFF", bold=True)
    box(d, (55, 520, 375, 675), "同一预测建筑内\nSAR 局部候选搜索", "#FFFFFF", bold=True)
    box(d, (515, 520, 855, 675), "关系特征\n匹配置信度与熵", "#F2F2F2", bold=True)
    box(d, (980, 330, 1310, 500), "有界残差\nα × ΔLchange", "#F2F2F2", bold=True)
    box(d, (1380, 330, 1560, 500), "软关联\nQ", "#FFFFFF", bold=True)
    arrow(d, (375, 598), (515, 598)); arrow(d, (855, 598), (1060, 500))
    arrow(d, (375, 228), (1060, 330)); arrow(d, (1310, 415), (1380, 415))
    label(d, (800, 775), "Lfinal = Lstruct + αΔLchange      |α| 受上限约束且初始为 0", 30, True)
    paths.append(save_canvas(img, "figure_3_cross_modal.png"))

    # Figure 4: tokens and building evidence.
    img = make_canvas(); d = ImageDraw.Draw(img)
    label(d, (800, 55), "区域证据 Token 与建筑内损毁结构", 40, True)
    box(d, (60, 180, 390, 670), "同一建筑内\n多个局部区域\n\nR1  R2  …  RK", "#FFFFFF", bold=True)
    box(d, (520, 125, 905, 330), "区域统计\nRGB 均值  SAR 均值\n关系均值与离散度", "#F2F2F2", bold=True)
    box(d, (520, 500, 905, 705), "强证据保留\n软峰值  匹配质量\n面积占比", "#F2F2F2", bold=True)
    box(d, (1070, 180, 1510, 670), "建筑证据结构\n\n整体均值\n内部异质性\nTop M 关键区域\n区域邻接关系", "#FFFFFF", bold=True)
    arrow(d, (390, 360), (520, 230)); arrow(d, (390, 490), (520, 600))
    arrow(d, (905, 230), (1070, 360)); arrow(d, (905, 600), (1070, 490))
    label(d, (800, 810), "避免小面积损毁在整栋建筑平均池化中被稀释", 30, True)
    paths.append(save_canvas(img, "figure_4_evidence_structure.png"))

    # Figure 5: specialist and memory verifier.
    img = make_canvas(); d = ImageDraw.Draw(img)
    label(d, (800, 55), "条件损毁专家与跨事件历史验证", 40, True)
    box(d, (70, 290, 370, 500), "基础预测为完整\n或完整与损毁混淆", "#FFFFFF", bold=True)
    box(d, (490, 290, 790, 500), "条件损毁专家\n输出疑似损毁分数", "#F2F2F2", bold=True)
    box(d, (910, 120, 1240, 330), "建筑级粗检索\n排除同一灾害事件", "#FFFFFF", bold=True)
    box(d, (910, 500, 1240, 710), "关键区域集合\n细粒度重排序", "#FFFFFF", bold=True)
    box(d, (1360, 290, 1550, 500), "效用\n验证器", "#F2F2F2", bold=True)
    arrow(d, (370, 395), (490, 395)); arrow(d, (790, 350), (910, 225)); arrow(d, (790, 440), (910, 605))
    arrow(d, (1240, 225), (1400, 290)); arrow(d, (1240, 605), (1400, 500))
    label(d, (800, 805), "历史病例不直接改写主干特征  仅验证当前修正是否值得采用", 29, True)
    paths.append(save_canvas(img, "figure_5_specialist_memory.png"))

    # Figure 6: selective late correction.
    img = make_canvas(); d = ImageDraw.Draw(img)
    label(d, (800, 55), "选择性末端决策修正", 40, True)
    box(d, (90, 170, 450, 335), "Base = Intact", "#FFFFFF", bold=True)
    box(d, (90, 545, 450, 710), "Base ≠ Intact", "#FFFFFF", bold=True)
    box(d, (610, 170, 1010, 335), "专家分数与历史证据\n通过效用门限", "#F2F2F2", bold=True)
    box(d, (610, 545, 1010, 710), "未通过门限\n或不属于目标类别", "#F2F2F2", bold=True)
    box(d, (1180, 170, 1510, 335), "仅原 Intact 像素\n修正为 Damaged", "#FFFFFF", bold=True)
    box(d, (1180, 545, 1510, 710), "保持基础预测", "#FFFFFF", bold=True)
    arrow(d, (450, 252), (610, 252)); arrow(d, (1010, 252), (1180, 252))
    arrow(d, (450, 628), (610, 628)); arrow(d, (1010, 628), (1180, 628))
    label(d, (800, 815), "背景  已损毁  已毁坏像素均不被该修正步骤覆盖", 30, True)
    paths.append(save_canvas(img, "figure_6_late_correction.png"))
    return paths


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="8"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag); borders.append(element)
        element.set(qn("w:val"), "single"); element.set(qn("w:sz"), size); element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def suppress_paragraph_borders(paragraph_or_style):
    element = paragraph_or_style._element
    if element.tag == qn("w:style"):
        p_pr = element.find(qn("w:pPr"))
        if p_pr is None:
            p_pr = OxmlElement("w:pPr")
            element.append(p_pr)
    else:
        p_pr = element.get_or_add_pPr()
    for old in list(p_pr.findall(qn("w:pBdr"))):
        p_pr.remove(old)
    p_bdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "nil")
        p_bdr.append(border)
    p_pr.append(p_bdr)


def set_east_asia_font(run, name="宋体"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar"); fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar"); fld_sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t"); txt.text = "1"
    fld_end = OxmlElement("w:fldChar"); fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, txt, fld_end])
    set_east_asia_font(run); run.font.size = Pt(9)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink"); hyperlink.set(qn("r:id"), rid)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "000000")
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline]); new_run.append(r_pr)
    t = OxmlElement("w:t"); t.text = text; new_run.append(t); hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_equation(doc, text):
    p = doc.add_paragraph(style="Equation")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.name = "Cambria Math"; r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Cambria Math")
    r.font.size = Pt(10.5)
    return p


def add_body(doc, text, *, bold_lead=None, indent=True, keep=False):
    p = doc.add_paragraph(style="Body Text")
    if indent:
        p.paragraph_format.first_line_indent = Pt(22)
    if keep:
        p.paragraph_format.keep_with_next = True
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead); r1.bold = True; set_east_asia_font(r1, "黑体")
        r2 = p.add_run(text[len(bold_lead):]); set_east_asia_font(r2)
    else:
        r = p.add_run(text); set_east_asia_font(r)
    return p


def add_claim(doc, number, paragraphs):
    for idx, text in enumerate(paragraphs):
        p = doc.add_paragraph(style="Claim")
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.left_indent = Pt(0 if idx == 0 else 22)
        p.paragraph_format.keep_together = False
        if idx == 0:
            lead = f"{number}. "
            r = p.add_run(lead); r.bold = True; set_east_asia_font(r, "黑体")
        r = p.add_run(text); set_east_asia_font(r)


def add_figure(doc, path, caption, width=6.55):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = False
    r = cap.add_run(caption); r.bold = True; set_east_asia_font(r, "宋体")


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Inches(0.78); section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.9); section.right_margin = Inches(0.9)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(4); normal.paragraph_format.line_spacing = 1.35

    title = styles["Title"]
    title.font.name = "黑体"; title._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    title.font.size = Pt(22); title.font.bold = True; title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(20)
    suppress_paragraph_borders(title)

    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = styles[name]
        style.font.name = "黑体"; style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12 if name != "Heading 1" else 16)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True

    if "Body Text" not in styles:
        styles.add_style("Body Text", WD_STYLE_TYPE.PARAGRAPH)
    body = styles["Body Text"]
    body.font.name = "宋体"; body._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    body.font.size = Pt(10.5); body.font.color.rgb = RGBColor(0, 0, 0)
    body.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    body.paragraph_format.space_after = Pt(4)
    body.paragraph_format.widow_control = True

    if "Claim" not in styles:
        styles.add_style("Claim", WD_STYLE_TYPE.PARAGRAPH)
    claim = styles["Claim"]
    claim.font.name = "宋体"; claim._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    claim.font.size = Pt(11); claim.font.color.rgb = RGBColor(0, 0, 0)
    claim.paragraph_format.line_spacing = 1.45; claim.paragraph_format.space_after = Pt(5)

    if "Equation" not in styles:
        styles.add_style("Equation", WD_STYLE_TYPE.PARAGRAPH)
    eq = styles["Equation"]
    eq.font.name = "Cambria Math"; eq.font.size = Pt(10.5); eq.font.color.rgb = RGBColor(0, 0, 0)
    eq.paragraph_format.space_before = Pt(4); eq.paragraph_format.space_after = Pt(7)

    caption = styles["Caption"]
    caption.font.name = "宋体"; caption._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    caption.font.size = Pt(10); caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.space_after = Pt(9)

    for s in doc.sections:
        add_page_number(s.footer.paragraphs[0])
    return doc


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if level > 1 else WD_ALIGN_PARAGRAPH.CENTER
    return p


def add_numbered_paragraphs(doc, items, start):
    n = start
    for item in items:
        add_body(doc, f"[{n:04d}] {item}")
        n += 1
    return n


def build_document(figs):
    doc = setup_document()

    # Cover.
    p = doc.add_paragraph(style="Title")
    suppress_paragraph_borders(p)
    p.paragraph_format.space_before = Pt(90)
    r = p.add_run("一种基于建筑内局部证据与跨事件效用验证的\n多模态建筑损毁评估方法 系统及存储介质")
    r.bold = True; set_east_asia_font(r, "黑体"); r.font.size = Pt(22)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_before = Pt(22)
    rs = sub.add_run("中国发明专利申请文件初稿")
    rs.bold = True; rs.font.size = Pt(14); set_east_asia_font(rs, "黑体")
    desc = doc.add_paragraph(); desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    desc.paragraph_format.space_before = Pt(140)
    rd = desc.add_run("包含说明书摘要  权利要求书  说明书  说明书附图\n附录为代理人沟通用创新性定位说明")
    rd.font.size = Pt(11); set_east_asia_font(rd)
    doc.add_page_break()

    # Abstract.
    add_heading(doc, "说明书摘要", 1)
    abstract = (
        "本发明公开一种基于建筑内局部证据与跨事件效用验证的多模态建筑损毁评估方法、系统及存储介质。"
        "利用灾前光学影像和灾后合成孔径雷达影像获得基础预测，并在预测建筑内设置局部区域，限制像素仅关联同一建筑的区域中心；"
        "以光学结构形成基础分区，以建筑内雷达关系作有界残差修边，聚合局部统计量、匹配质量和面积，形成建筑内损毁证据结构。"
        "仅对被判为完整或完整与损毁混淆的建筑启动条件损毁专家，并以排除同灾害事件的历史病例验证修正效用。"
        "效用满足条件时，在不改变基础模型中间特征的情况下，仅将目标区域内原完整类别修正为损毁类别；否则保持原结果。"
        "本发明可降低跨模态错位、局部证据稀释、类别长尾及历史证据误用造成的误判。"
    )
    add_body(doc, abstract)
    add_heading(doc, "摘要附图", 2)
    add_figure(doc, figs[0], "图1", width=6.4)
    doc.add_page_break()

    # Claims.
    add_heading(doc, "权利要求书", 1)
    add_claim(doc, 1, [
        "一种基于建筑内局部证据与跨事件效用验证的多模态建筑损毁评估方法，其特征在于，包括：",
        "获取目标区域的灾前光学遥感影像和灾后合成孔径雷达影像，将所述灾前光学遥感影像和灾后合成孔径雷达影像输入基础建筑损毁评估网络，获得建筑位置概率、基础建筑损毁类别概率、光学结构特征和雷达响应特征，所述基础建筑损毁类别至少包括背景类别、完整类别、损毁类别和毁坏类别；",
        "根据所述建筑位置概率获得多个预测建筑实例，并针对每一预测建筑实例设置与该预测建筑实例面积相适应的多个局部区域中心，将任一建筑像素的候选局部区域中心限制为同一预测建筑实例内的局部区域中心，以获得不跨越不同预测建筑实例的建筑约束候选集合；",
        "根据所述光学结构特征形成所述建筑约束候选集合内的结构关联得分，在同一预测建筑实例内从所述雷达响应特征中搜索与光学结构位置相对应的局部响应，获得跨模态关系特征、匹配置信度和匹配不确定度，并以幅值受限的变化残差修正所述结构关联得分，得到像素与局部区域中心之间的软关联；",
        "按照所述软关联聚合各局部区域的光学结构特征、匹配后的雷达响应特征、跨模态关系均值、跨模态关系离散度、局部强响应、匹配质量和区域面积占比，形成多个区域证据表示，并利用所述区域证据表示的总体统计量、内部异质性及关键局部区域表示形成建筑内损毁证据结构；",
        "从基础建筑损毁预测结果中选取被预测为完整类别或者完整类别与损毁类别的预测分数满足预设混淆条件的预测建筑实例，将其作为疑难建筑候选，并利用以建筑内损毁证据结构为输入的条件损毁专家获得疑似损毁分数；",
        "将所述疑难建筑候选的建筑级表示和关键局部区域表示与跨灾害事件历史证据库进行比较，在排除与所述疑难建筑候选属于同一灾害事件的历史建筑后，获得历史损毁支持信息；",
        "根据所述疑似损毁分数、历史损毁支持信息和基础建筑损毁预测的不确定度确定附加证据效用，并仅在所述附加证据效用满足预设接受条件时，在保持所述基础建筑损毁评估网络的中间特征不变的情况下，将目标预测建筑实例或其关键局部区域内原预测为完整类别的像素选择性修正为损毁类别；在不满足所述预设接受条件时保持基础建筑损毁预测结果不变。"
    ])
    add_claim(doc, 2, [
        "根据权利要求1所述的方法，其特征在于，所述多个预测建筑实例通过对建筑位置概率阈值化并进行连通域分析获得；每一预测建筑实例的局部区域中心数量根据该预测建筑实例的面积、目标区域面积、最小区域面积和每栋建筑的最大区域数量确定，且候选局部区域中心被吸附至该预测建筑实例内部的有效位置。"
    ])
    add_claim(doc, 3, [
        "根据权利要求1所述的方法，其特征在于，所述跨模态局部响应搜索在对应预测建筑实例与预设邻域的交集中进行；所述变化残差由光学特征、候选雷达特征、二者的绝对差、二者的乘积、匹配置信度、匹配不确定度和基础建筑损毁预测的不确定度中的至少两项生成，并通过初始值为零且绝对值不超过预设上限的门控系数作用于所述结构关联得分。"
    ])
    add_claim(doc, 4, [
        "根据权利要求1所述的方法，其特征在于，所述局部强响应为根据区域内像素证据得分进行温度控制加权得到的可微软峰值；在不存在局部区域损毁标注时，将同一预测建筑实例的多个区域证据表示通过最大聚合、前若干项聚合、注意力聚合或对数和指数聚合中的至少一种方式得到建筑级结果，并以建筑级损毁标签对所述区域证据表示进行弱监督训练。"
    ])
    add_claim(doc, 5, [
        "根据权利要求1所述的方法，其特征在于，所述条件损毁专家在保持基础建筑损毁评估网络的主体表示固定或以低于所述条件损毁专家的学习率更新的条件下训练，并采用类别平衡采样、灾害事件平衡采样、少数类别加权、动态正负样本重加权和完整类别与损毁类别的间隔约束中的至少一种，以降低完整类别样本对损毁类别的梯度支配。"
    ])
    add_claim(doc, 6, [
        "根据权利要求1所述的方法，其特征在于，所述跨灾害事件历史证据库至少存储历史建筑的建筑级检索表示、局部区域证据表示集合、建筑损毁类别和灾害事件标识；先依据建筑级检索表示进行粗检索，再依据当前建筑与候选历史建筑的关键局部区域表示集合进行细粒度重排序，并根据历史损毁支持度与历史完整支持度之差、检索相似度、邻居类别一致性、邻居事件多样性和检索得分离散度中的至少两项确定所述附加证据效用；所述选择性修正的作用范围为整栋预测建筑实例、一个或多个关键局部区域或者由空间邻接的关键局部区域形成的区域簇，且不改变背景类别、已预测损毁类别和已预测毁坏类别的像素。"
    ])
    add_claim(doc, 7, [
        "一种多模态建筑损毁评估系统，其特征在于，包括基础预测模块、建筑约束局部区域生成模块、建筑内跨模态搜索模块、区域证据结构构建模块、条件损毁专家模块、跨事件历史证据验证模块、证据效用判定模块以及选择性末端修正模块；各模块被配置为协同执行权利要求1至6任一项所述的方法。"
    ])
    add_claim(doc, 8, [
        "一种计算机可读存储介质，其上存储有计算机程序，其特征在于，所述计算机程序被处理器执行时实现权利要求1至6任一项所述的方法。"
    ])
    doc.add_page_break()

    # Specification.
    add_heading(doc, "说明书", 1)
    add_heading(doc, "发明名称", 2)
    add_body(doc, "一种基于建筑内局部证据与跨事件效用验证的多模态建筑损毁评估方法、系统及存储介质", indent=False)

    add_heading(doc, "技术领域", 2)
    n = 1
    n = add_numbered_paragraphs(doc, [
        "本发明涉及遥感图像智能解译、灾害监测、合成孔径雷达影像处理以及多模态机器学习技术领域，尤其涉及一种利用灾前光学遥感影像和灾后合成孔径雷达影像，在建筑实例内部构造局部损毁证据，并经跨灾害事件历史证据验证后选择性修正建筑损毁预测结果的方法、系统及计算机可读存储介质。"
    ], n)

    add_heading(doc, "背景技术", 2)
    n = add_numbered_paragraphs(doc, [
        "地震、洪水、火灾、爆炸、火山活动及武装冲突等事件可能在短时间内造成大量建筑物损毁。快速获得建筑位置和损毁等级，对于应急救援、人员疏散、物资调度以及灾情统计具有重要作用。",
        "灾后光学遥感影像容易受到云层、烟尘、降雨、夜间成像及重访周期的影响。合成孔径雷达具有全天时、全天候观测能力，因此，利用灾前光学遥感影像提供的清晰建筑结构与灾后合成孔径雷达影像提供的散射响应进行协同评估，能够提高灾后快速响应能力。",
        "但是，光学影像与合成孔径雷达影像的成像机理差异显著。前者主要表达颜色、纹理和边界，后者主要表达微波散射响应，并伴随散斑噪声和侧视几何效应。即使完成全局配准，相同坐标位置仍可能不对应相同的建筑结构，直接逐像素差分或者在同一坐标处拼接特征容易把模态差异和局部错位误认为真实损毁。",
        "现有的超像素融合方法通常在参考影像上完成超像素分割，将相同分割结果应用于多幅损毁检测图像，并对相互对应的超像素置信度进行融合。该类方法能够减轻不同空间分辨率造成的插值误差，但仍以跨图像的对应超像素作为融合单元，未解决预测建筑之间的区域串联、建筑内部的非同坐标局部匹配以及少量强损毁证据被区域平均稀释的问题。",
        "现有的光学与雷达多特征方法还可按照建筑矢量单元提取极化、相干性、指数或位移等多尺度特征，并训练分类网络。该类方法侧重特征库构建和整体分类，未显式区分基础模型已经能够稳定处理的普通样本与需要二次复核的疑难样本，也未设置仅在修正有益时才改变最终结果的约束。",
        "建筑损毁通常具有明显的局部性。中度损毁可能只出现在屋顶局部、边缘、连接部位或与周边散射体交互的位置。若直接对整栋建筑进行平均池化，小面积损毁所对应的强异常会被大量正常区域稀释；若直接采用最大值，又容易受到雷达散斑和配准误差产生的孤立强响应干扰。",
        "建筑损毁数据还存在类别长尾。完整建筑通常数量较多，中度损毁建筑数量较少，基础模型易将少数损毁建筑预测为完整建筑。若为了改善少数类别而重新改变整个基础特征空间，可能同时破坏大量已经正确分类的完整建筑和毁坏建筑。",
        "历史灾害事件中的建筑病例包含可供参考的局部损毁结构，但不同事件在灾种、地表环境、建筑类型、传感器参数和成像方向方面存在差异。把历史病例特征直接写入当前模型可能形成事件捷径或引入误导，检索到相似病例本身也不能说明采用该病例一定能够改善当前预测。",
        "因此，需要一种以基础模型为稳定主预测器、在预测建筑边界内形成位置容错的跨模态局部证据、仅针对完整与损毁混淆的长尾疑难建筑进行复核、利用跨事件历史病例验证修正效用，并在附加证据可靠时才实施受控末端修正的建筑损毁评估方案。"
    ], n)

    add_heading(doc, "发明内容", 2)
    add_heading(doc, "要解决的技术问题", 3)
    n = add_numbered_paragraphs(doc, [
        "本发明旨在至少解决以下技术问题：异源光学与雷达特征难以逐像素可靠比较；不同建筑之间的局部区域可能发生错误连接；小面积损毁证据容易被建筑级平均特征稀释；长尾损毁类别容易被完整类别压制；历史病例可能因事件差异而误导当前判断；以及附加模块缺少对修正收益的显式判断，容易在救回少量损毁建筑的同时误伤大量正确完整建筑。"
    ], n)
    add_heading(doc, "技术方案", 3)
    n = add_numbered_paragraphs(doc, [
        "为解决上述技术问题，本发明提供一种多模态建筑损毁评估方法。该方法不把超像素或局部区域本身作为最终分类结果，而将其作为建筑内部损毁证据的组织单元；不要求光学特征与雷达特征在同一坐标处直接对应，而在同一预测建筑内部进行局部候选搜索；不让雷达变化证据从训练初期任意改变光学结构分区，而以幅值受限且初始为零的残差对结构关联进行修正。",
        "所述方法首先由基础建筑损毁评估网络输出建筑位置概率、基础损毁类别概率、光学结构特征和雷达响应特征。基础网络可以采用卷积网络、Transformer网络、状态空间网络或其组合，本发明不限定基础网络的具体类型。",
        "依据建筑位置概率形成预测建筑实例，并按照每栋预测建筑的面积设置局部区域中心。像素只与同一预测建筑内的候选中心建立关联，从候选集合的结构上排除跨建筑连接。",
        "在候选集合内，以光学结构特征、光学边缘和建筑局部坐标形成结构关联；同时，在同一预测建筑范围内为光学位置搜索多个雷达候选位置，获得匹配后的雷达表示、关系特征、匹配置信度和匹配不确定度。利用受限门控系数将变化残差叠加至结构关联，形成最终软关联。",
        "按照软关联聚合区域证据。除区域均值外，还保留关系离散度和可微软峰值，并结合匹配质量与区域面积占比形成区域证据表示。对同一建筑的多个区域证据表示计算总体均值、内部离散程度、关键区域表示和区域关系，形成建筑内损毁证据结构。",
        "基础网络继续负责全部建筑的常规预测。仅当基础预测为完整类别，或者完整类别与损毁类别的分数接近、预测不确定度较高时，才启动条件损毁专家。条件损毁专家专门判断所述疑难建筑是否具有被基础模型遗漏的损毁证据。",
        "历史证据库由训练阶段已经标注的历史建筑建立。查询时排除与当前建筑属于同一灾害事件的历史建筑，先在建筑级表示上粗检索，再根据关键局部区域集合细粒度重排序，分别获得支持完整解释和支持损毁解释的历史证据。",
        "证据效用判定模块综合条件损毁专家分数、基础模型不确定度、历史支持差、检索相似度、邻居类别一致性、邻居事件多样性和检索得分离散度，判断采用附加证据是否值得。仅当预设接受条件满足时执行完整类别至损毁类别的末端修正。",
        "末端修正不将区域证据或历史证据重新写入基础网络的中间特征。在一种优选方式中，仅修改目标建筑范围内原预测为完整类别的像素，背景、已预测损毁及已预测毁坏像素均保持不变；修正范围还可以限定为关键局部区域或相邻关键局部区域形成的区域簇。"
    ], n)
    add_equation(doc, "Lfinal = Lstruct + α × ΔLchange")
    add_equation(doc, "α = αmax × tanh(g) × sigmoid(v)，且 |α| ≤ αmax")
    add_equation(doc, "Δmemory = Supportdamaged − Supportintact")
    add_equation(doc, "Final = Damaged，仅当 Base = Intact 且 Utility ≥ τ；否则 Final = Base")

    add_heading(doc, "有益效果", 3)
    n = add_numbered_paragraphs(doc, [
        "与把全图超像素作为跨影像融合单元的方案相比，本发明首先形成预测建筑实例，再从候选集合结构上禁止像素跨建筑关联，使局部证据与独立建筑实体保持一致。",
        "与同坐标特征差分或拼接相比，本发明在同一建筑内部搜索雷达候选响应，并将匹配置信度和不确定度作为独立证据，能够降低跨模态成像差异和局部配准误差的影响。",
        "与直接由雷达变化特征决定区域边界相比，本发明以光学结构关联作为基础，仅允许有界变化残差进行修边，并采用零初始门控，使新增模块在训练开始时不破坏基础结构。",
        "通过同时保存区域均值、离散度和软峰值，本发明在抑制孤立噪声的同时保留局部强异常，降低小面积损毁被整体平均稀释的风险。",
        "通过仅对基础模型中的完整类疑难建筑启用条件专家，本发明把长尾处理集中在主要混淆方向，而不强制重新塑造全部类别的基础表示。",
        "通过排除同事件病例、采用建筑级粗检索和局部区域细排，并由效用判定模块决定是否采用历史证据，本发明将历史病例从直接分类器转变为修正验证器，减少事件捷径和错误邻居造成的干扰。",
        "通过选择性末端修正，本发明保留基础网络已经正确给出的背景、损毁和毁坏结果，仅对高可信的完整至损毁混淆进行有限修正，便于审计每次修改的收益与风险。"
    ], n)

    add_heading(doc, "附图说明", 2)
    n = add_numbered_paragraphs(doc, [
        "图1为本发明总体技术流程示意图；图2为预测建筑约束的局部区域生成示意图；图3为建筑内跨模态搜索与有界修边示意图；图4为区域证据表示与建筑内损毁结构示意图；图5为条件损毁专家与跨事件历史验证示意图；图6为选择性末端决策修正示意图。"
    ], n)

    add_heading(doc, "具体实施方式", 2)
    n = add_numbered_paragraphs(doc, [
        "下面结合附图和实施例对本发明作进一步说明。应当理解，所述实施例用于解释本发明，而不用于限定本发明的保护范围。在不冲突的情况下，不同实施例中的技术特征可以相互组合。",
        "实施例一为完整的训练与推理流程。获取灾前光学遥感影像 Iopt 和灾后合成孔径雷达影像 Isar，对二者进行地理参考统一、尺度归一化和必要的辐射预处理。全局配准只需使同一建筑落入相邻区域，不要求光学像素与雷达像素严格同坐标对应。",
        "将两种影像输入基础建筑损毁评估网络。网络输出建筑位置概率图 Pbuilding、四类别概率图 Pbase、光学结构特征 Frgb 和雷达响应特征 Fsar。四类别依次为背景、完整、损毁和毁坏；也可以根据应用需要把损毁划分为轻度、中度和重度。",
        "对 Pbuilding 进行阈值化和四连通或八连通分析，删除面积小于预设阈值的伪建筑，获得预测建筑实例集合 B={B1，B2，…，BN}。所述建筑实例来自模型预测而非推理阶段的真实标注，因此能够直接用于部署。",
        "针对建筑 Bi，按照建筑面积 Ai 估计局部区域数量 Ki。一个示例为 Ki 取 Ai 与目标区域面积之比的向上取整值，并同时受最小区域面积和每栋建筑最大区域数量限制。局部区域中心按照紧凑网格初始化，再吸附至 Bi 内最近的有效建筑像素。",
        "对于 Bi 内像素 p，仅建立与 Bi 内邻近区域中心的候选关系。若中心 c 不属于 Bi，则 p 与 c 的候选关系被置为无效，软关联概率恒为零。由此，不同建筑之间即使在空间上相邻，也不会共享同一个局部区域。",
        "结构关联分支以 Frgb、光学边缘和建筑内部归一化坐标为输入，输出像素与候选中心之间的结构关联 Lstruct。还可以加入固定的空间紧凑度偏置，以限制像素关联至过远区域中心。",
        "对于光学位置 p，在 Bi 与以 p 为中心的预设搜索邻域的交集中构造雷达候选集合 Ni(p)。利用分别作用于光学特征和雷达特征的映射函数计算候选关系分数，并对候选关系分数进行归一化加权，得到匹配后的雷达表示。该步骤允许雷达响应相对于光学结构发生小范围位移。",
        "变化分支的输入可以包括映射后的光学值、匹配后的雷达值、绝对差、逐元素乘积、匹配置信度、匹配熵、局部关系变化和基础模型不确定度。变化分支输出候选中心级变化残差 ΔLchange 和门控值 v。",
        "最终关联采用结构主导的残差形式。变化增益 g 初始设为零，门控系数 α 的绝对值受 αmax 限制。因而在训练开始时 Lfinal 与 Lstruct 完全一致；随着训练进行，只有经过证据支持的雷达变化才能在有限幅度内修正光学结构分区。",
        "对 Lfinal 在同一建筑的有效候选集合内进行掩膜归一化，获得软关联 Q。训练阶段使用 Q 完成可微池化和重投影；可视化或部署时可以把最大关联中心作为硬区域标签，并合并面积过小且与相邻区域相似的区域。",
        "针对局部区域 Rik，利用 Q 计算光学结构均值、雷达响应均值和跨模态关系均值。进一步计算跨模态关系标准差，以描述区域内部异质性；匹配置信度均值和匹配熵均值用于描述该区域跨模态证据的可靠性；区域面积与建筑面积之比用于描述区域尺度。",
        "局部强响应采用软峰值而非硬最大值。具体地，根据每个像素的证据得分和温度参数计算指数权重，再与 Q 相乘并归一化。温度较低时更关注强证据像素，温度较高时更接近均值。该方式保留可微性，并降低单个噪声像素直接主导结果的风险。",
        "将上述统计量编码为区域证据表示 zik。对同一建筑的全部 zik 计算均值和标准差，依据软峰值与区域均值之差、区域内部关系变化以及可学习异常分数确定关键性分数，选择前 M 个关键区域。建筑证据结构 Zi 由总体均值、内部离散度、前 M 个关键区域和可选的区域邻接关系组成。",
        "当训练数据仅具有建筑级损毁标签时，不把局部区域称为真实局部损毁标注，而把其作为损毁相关证据区域。将各区域的证据预测通过最大聚合、前若干项聚合、注意力聚合或对数和指数聚合得到建筑级预测，并利用建筑级标签训练区域证据分支。对标签纯度不足的预测建筑可以降低权重或跳过。",
        "实施例二说明条件损毁专家。根据 Pbase 计算每个预测建筑的类别概率和不确定度。若基础建筑类别为完整，或者完整与损毁概率之差小于预设阈值，则把该建筑加入疑难候选集合；其他建筑仍直接采用基础预测。",
        "条件损毁专家的输入包括基础类别概率、基础不确定度、建筑证据结构 Zi、关键区域统计量以及局部强响应与均值的差。条件损毁专家输出疑似损毁分数 sDI，用于回答基础模型判定为完整的建筑是否包含被遗漏的损毁证据。",
        "条件损毁专家可以在冻结基础网络的情况下单独训练，或者以显著较低的学习率联合微调。训练样本在完整与损毁类别之间平衡，并可进一步在灾害事件之间平衡。对于少数损毁类别增加权重，或者根据当前误分类动态调整正负样本梯度。",
        "一个类别与事件联合采样示例为：属于类别 c 和事件 e 的样本，其采样权重与该类别样本数和该事件内该类别样本数的乘积成反比。该策略使数据量较大的完整类别或单一灾害事件不易主导条件专家。",
        "实施例三说明跨事件历史证据验证。利用训练集中已标注建筑建立历史证据库 M。每条病例至少包含建筑级检索表示 kj、局部区域证据表示集合 Tj、损毁类别 yj 和灾害事件标识 ej；还可以包含建筑面积、方向、成像质量和基础模型预测信息。",
        "对于当前疑难建筑，首先计算其建筑级表示与历史建筑级表示的相似度，得到粗检索候选。删除事件标识与当前建筑相同的病例，并限制每个来源事件可进入结果的邻居数量，以提高历史邻居的事件多样性。",
        "在粗检索候选内部，利用当前建筑关键区域集合与历史建筑关键区域集合进行细粒度匹配。细粒度匹配可以采用单向最大支持、双向集合匹配、最优传输距离或者重构距离。匹配结果用于重新排列历史候选，而不是直接替代基础分类。",
        "将历史完整建筑和历史损毁建筑分别形成支持集合，计算完整支持度 SI 和损毁支持度 SD，历史支持差为 SD 与 SI 之差。还可以先计算当前建筑与完整原型、损毁原型的相似度差，作为进入实例级检索的前置筛选。",
        "为避免检索模块仅依赖基础分类答案，可以分别采用打乱历史标签、打乱事件标识、使用错误区域集合和随机邻居作为训练或验证对照。只有真实的跨事件、区域感知历史病例在相同修正覆盖率下优于上述对照时，才启用历史证据。",
        "实施例四说明证据效用判定。证据效用模块的输入包括 sDI、基础不确定度、历史支持差、原型支持差、前若干邻居的相似度、类别一致性、事件多样性以及检索得分离散度。模块输出 0 至 1 之间的效用值 u。",
        "效用模块的训练目标不是再次预测建筑类别，而是判断采用候选修正相对于保留基础结果是否产生净收益。可以根据训练集或内部校准集上候选修正前后的正确像素数变化、建筑分类正确性变化或者预设代价函数生成效用监督。",
        "推理时只有同时满足基础类别为完整、条件损毁专家分数达到阈值以及效用值达到阈值时，才设置修正门控 Gi 为真。若任一条件不满足，则 Gi 为假，当前建筑保持基础预测。",
        "当 Gi 为真时，可以采用建筑级、区域簇级或单区域级修正。建筑级修正仅把 Bi 内原预测为完整类别的像素改为损毁类别；区域簇级修正先选择关键性分数高于阈值且空间邻接的区域，再只修改所述区域簇内的完整类别像素；单区域级修正仅作用于一个或多个关键区域。",
        "在上述任一修正粒度下，背景像素、已预测为损毁的像素和已预测为毁坏的像素均不改变。区域证据、条件专家输出和历史证据只参与最终门控，不写回基础网络的中间特征，从而保持基础模型行为可复现。",
        "实施例五说明训练阶段。第一阶段训练或复现基础建筑损毁评估网络；第二阶段在基础网络固定的情况下训练光学结构关联，使局部区域满足建筑内紧凑性、结构重构和区域规模平衡；第三阶段打开受限变化残差，并利用建筑级弱监督训练跨模态分区和区域证据；第四阶段训练条件损毁专家；第五阶段建立跨事件历史证据库并训练或校准效用模块。",
        "推理阶段不需要局部损毁像素真值，也不需要当前灾害事件的已标注病例。历史证据库仅使用先前已标注建筑构建，当前查询建筑只读取灾前光学影像、灾后雷达影像、基础模型输出及历史证据库。",
        "本发明中的阈值、局部搜索半径、目标区域面积、最大区域数量、软峰值温度、粗检索数量、细检索数量和效用阈值均可依据内部校准集确定。参数选择应在不访问最终测试标签的条件下完成。",
        "本发明可以由服务器、工作站、边缘计算设备或云端计算平台执行。系统中的各功能模块可以由软件、硬件或软硬件组合实现；各模块可以集成于同一处理器，也可以分布于不同处理节点。",
        "本领域技术人员在不脱离本发明实质的情况下，可以对局部区域生成方式、关系计算方式、聚合函数、检索距离和效用判定模型作等同替换。只要仍然采用建筑实例内候选硬约束、结构主导的受限跨模态修边、建筑内证据结构、疑难完整建筑条件复核、跨事件效用验证和选择性末端修正的组合，均可落入本发明构思。"
    ], n)

    doc.add_page_break()
    add_heading(doc, "说明书附图", 1)
    for idx, path in enumerate(figs, 1):
        add_figure(doc, path, f"图{idx}", width=6.35)
        if idx not in (2, 4, 6):
            # Two figures per page where layout permits.
            pass
        if idx in (2, 4):
            doc.add_page_break()

    # Internal appendix for novelty positioning.
    doc.add_page_break()
    add_heading(doc, "附录  创新性定位与代理人沟通要点", 1)
    add_body(doc, "本附录用于提交前与专利代理人沟通保护边界，不属于正式申请文件。正式递交时可单独保留或删除。", indent=False)
    add_heading(doc, "建议坚持的独立权利要求组合", 2)
    add_body(doc, "独立权利要求不以“光学影像与雷达影像使用超像素融合”为中心，而把以下技术特征绑定为一条连续链：预测建筑实例内候选硬约束；光学结构基础分区与建筑内雷达局部搜索；幅值受限且零初始的变化残差修边；均值、离散度和软峰值共同构成建筑内证据结构；只复核基础模型中的完整类疑难建筑；排除同事件的历史病例验证；效用门控；不回写主干特征的完整至损毁末端修正。任何删减都可能扩大现有技术重合风险。", indent=False)

    add_heading(doc, "快速检索对比", 2)
    rows = [
        ["公开方案", "公开方案的核心", "本稿的差异化边界"],
        ["CN109360190B\n基于图像超像素融合的建筑损毁检测", "对雷达和光学损毁检测图像进行对应超像素处理；逐超像素形成BBA；融合对应超像素BBA得到结果。", "先形成预测建筑实例并硬性禁止跨建筑关联；在建筑内进行位置容错的雷达候选搜索；光学结构为基础、雷达变化仅作有界残差修边；局部区域用于构造证据结构而非直接BBA融合。"],
        ["CN120164117B\n耦合SAR图像和光学图像多特征的建筑物损毁检测", "围绕建筑矢量单元构造极化、相干性、指数和位移等多尺度特征库，并使用变分自编码器和改进分类器。", "不依赖预制多尺度物理特征库；基础模型先完成常规预测，仅对完整与损毁混淆建筑启动条件专家；历史证据只验证修正效用，最终不改写主干表示。"],
        ["CN105701481B\n一种倒塌建筑物提取方法", "依据极化SAR图像检索内容建立样本库，并通过距离度量学习提取倒塌建筑。", "历史库不是直接提取或替代分类器；查询排除同事件病例，采用建筑级粗检索与关键区域细排，并由效用模块决定是否保持基础结果。"]
    ]
    table = doc.add_table(rows=len(rows), cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(1.55), Inches(2.15), Inches(2.75)]
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j); cell.width = widths[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if i == 0:
                set_cell_shading(cell, "404040")
            elif i % 2 == 0:
                set_cell_shading(cell, "F2F2F2")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 or i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.2
            r = p.add_run(value); set_east_asia_font(r)
            r.font.size = Pt(9.2); r.bold = i == 0
            r.font.color.rgb = RGBColor(255, 255, 255) if i == 0 else RGBColor(0, 0, 0)
    set_repeat_table_header(table.rows[0]); set_table_borders(table)

    add_heading(doc, "撰写与审查意见答复要点", 2)
    points = [
        "若审查意见把本稿概括为超像素融合，应强调局部区域只是建筑内证据载体，核心限定还包括同楼候选硬约束、建筑内位置容错搜索、结构主导的受限修边、条件专家、跨事件效用验证以及不回写主干的末端修正。",
        "若审查意见认为多特征分类属于常规手段，应强调本稿没有用附加特征替换基础分类器，而是仅在基础模型的特定混淆方向上判断修正是否具有净效用，并明确限定未通过门控时保持基础结果。",
        "若审查意见质疑历史病例库的创造性，应强调同事件排除、建筑级粗检索与关键区域集合细排、事件多样性约束以及历史证据只用于修正验证的组合关系。",
        "若需要进一步收窄独立权利要求，优先保留“光学结构基础关联＋有界变化残差”“仅基础完整类候选”“跨事件病例排除”“不回写主干且只修改原完整像素”四组限制。"
    ]
    for idx, point in enumerate(points, 1):
        p = doc.add_paragraph(style="Body Text")
        p.paragraph_format.left_indent = Pt(18); p.paragraph_format.first_line_indent = Pt(-18)
        r = p.add_run(f"{idx}. "); r.bold = True; set_east_asia_font(r, "黑体")
        r2 = p.add_run(point); set_east_asia_font(r2)

    add_heading(doc, "检索依据", 2)
    p = doc.add_paragraph(style="Body Text")
    p.add_run("1. "); add_hyperlink(p, "CN109360190B 基于图像超像素融合的建筑损毁检测方法及装置", "https://patents.google.com/patent/CN109360190B/zh")
    p = doc.add_paragraph(style="Body Text")
    p.add_run("2. "); add_hyperlink(p, "CN120164117B 耦合SAR图像和光学图像多特征的建筑物损毁检测方法及装置", "https://patents.google.com/patent/CN120164117B/zh")
    p = doc.add_paragraph(style="Body Text")
    p.add_run("3. "); add_hyperlink(p, "CN105701481B 一种倒塌建筑物提取方法", "https://patents.google.com/patent/CN105701481B/zh")
    add_body(doc, "上述对比用于确定本稿的起草边界，不构成完整的新颖性检索结论。正式申请前应由专利代理人结合申请日、同族文献、权利要求解释和审查实践继续检索。", indent=False)

    # Core properties and metadata.
    props = doc.core_properties
    props.title = "建筑内局部证据与跨事件效用验证的建筑损毁评估发明专利初稿"
    props.subject = "中国发明专利申请文件初稿"
    props.keywords = "建筑损毁 光学遥感 SAR 局部区域 跨事件历史证据 选择性修正"
    props.comments = "Patent draft generated from the supplied technical disclosure"
    return doc


def main():
    BUILD_DIR.mkdir(parents=True, exist_ok=True)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    figs = draw_figures()
    doc = build_document(figs)
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
