from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(r"D:\software\dowmload\BRIGHT-master\论文写作\光学SAR建筑损毁小论文写作框架.docx")

BLUE = "1F4E78"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "F5F6F7"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), BORDER)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, east_asia="宋体", western="Times New Roman", size=10.5, bold=None, color=BLACK):
    run.font.name = western
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), western)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), western)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = color


def style_paragraph(paragraph, first_line=True, line=1.5, before=0, after=0):
    fmt = paragraph.paragraph_format
    fmt.line_spacing = line
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    if first_line:
        fmt.first_line_indent = Cm(0.74)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    style_paragraph(p)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    style_paragraph(p, first_line=False, line=1.35, after=2)
    set_run_font(p.add_run(text))
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    style_paragraph(p, first_line=False, line=1.35, after=2)
    set_run_font(p.add_run(text))
    return p


def add_equation(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, east_asia="Cambria Math", western="Cambria Math", size=11)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        set_run_font(r, east_asia="黑体", size=9.5, bold=True, color=RGBColor(255, 255, 255))
    for ridx, row_data in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cell = cells[i]
            set_cell_shading(cell, PALE_BLUE if ridx % 2 else "FFFFFF")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.2
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_run_font(r, size=9)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def remove_paragraph_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.4)
sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.6)
sec.right_margin = Cm(2.6)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.font.size = Pt(10.5)

title_style = styles["Title"]
title_style.font.name = "Times New Roman"
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
title_style.font.size = Pt(22)
title_style.font.bold = True
title_style.font.color.rgb = BLACK
title_p_pr = title_style._element.get_or_add_pPr()
title_p_bdr = title_p_pr.find(qn("w:pBdr"))
if title_p_bdr is not None:
    title_p_pr.remove(title_p_bdr)

for style_name, size, font in (("Heading 1", 16, "黑体"), ("Heading 2", 13, "黑体"), ("Heading 3", 11, "黑体")):
    s = styles[style_name]
    s.font.name = "Times New Roman"
    s._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = BLACK
    s.paragraph_format.space_before = Pt(10 if style_name != "Heading 1" else 14)
    s.paragraph_format.space_after = Pt(5)

# Title page
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(85)
p.paragraph_format.space_after = Pt(22)
p.add_run("光学SAR建筑损毁小论文写作框架")
remove_paragraph_border(p)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(18)
r = p.add_run("跨事件局部证据互补与修正效用验证")
set_run_font(r, east_asia="黑体", size=15, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(25)
r = p.add_run("用于论文结构冻结  方法撰写  实验规划与结果归因")
set_run_font(r, east_asia="宋体", size=11)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(170)
r = p.add_run("2026年9月")
set_run_font(r, size=11)

doc.add_page_break()

add_heading(doc, "文档用途与研究边界", 1)
add_body(doc, "本文档用于组织光学–SAR建筑损毁语义分割小论文的整体逻辑、章节内容、方法表达和实验闭环。论文保持像素级语义分割任务不变，以P3或DamageFormer为基础模型，以SLIC区域作为候选修改单元，以Direct Head和跨事件Memory提供互补局部证据，并利用Residual Utility Verifier决定保留基础预测、采用Direct建议或采用Memory建议。")
add_body(doc, "当前已获得的正面证据是跨事件Memory相对Direct存在额外Oracle候选空间；该结果仅说明Memory值得继续验证，不能视为最终无真实标签推理性能。论文中的最终性能结论必须来自M2至M4在独立验证集上的真实推理结果。")

add_table(doc, ["内容", "当前定位", "写作边界"], [
    ["研究任务", "光学–SAR建筑损毁语义分割", "不转为实例分割，不重新制作实例标签"],
    ["基础模型", "P3或DamageFormer兼容基线", "基础模型冻结后再训练局部Verifier"],
    ["Direct Head", "主要局部建议器", "以真实无GT推理结果评价"],
    ["跨事件Memory", "互补候选与历史效用证据", "默认排除同图、排除同事件"],
    ["Oracle分析", "候选上限与互补性证明", "不得将3.9401个百分点写成最终模型增益"],
    ["实例Token和大型新模块", "暂不纳入主线", "仅在已有实验支持时作为补充"],
], widths=[3.2, 5.2, 7.0])

add_heading(doc, "候选题目", 1)
add_table(doc, ["适用条件", "推荐题目"], [
    ["当前阶段", "基于跨事件记忆与效用验证的光学–SAR建筑损毁语义分割方法"],
    ["M4显著超过M3", "RUM-BDA  面向光学–SAR建筑损毁评估的残差效用记忆与可靠局部细化方法"],
    ["Memory未超过Direct Verifier", "基于效用验证与选择性回退的光学–SAR建筑损毁语义分割方法"],
], widths=[4.2, 11.2])

add_heading(doc, "论文逻辑总览", 1)
add_body(doc, "论文沿用“关键问题—对应方法—实验验证”的组织方式。核心逻辑不是罗列SLIC、Direct和Memory，而是回答以下问题：如何利用当前区域证据与跨事件历史证据识别基础模型中的可修正错误，并在修改可能有害时可靠回退。")
add_equation(doc, "基础预测  →  局部候选  →  Direct与Memory互补证据  →  Utility验证  →  选择性写回或回退")
add_table(doc, ["关键问题", "对应方法", "必要证据"], [
    ["基础模型存在局部漏检和类别混淆", "建筑概率引导的SLIC候选区域", "候选覆盖率与漏检恢复漏斗"],
    ["单一预测头无法覆盖跨事件损毁模式", "Direct当前证据与跨事件Memory证据", "Memory-only候选及事件级互补性"],
    ["局部修改可能同时有益和有害", "Residual Utility Verifier与选择性回退", "M2至M5真实推理和打乱标签对照"],
], widths=[5.1, 5.1, 5.2])

add_heading(doc, "摘要与关键词", 1)
add_heading(doc, "摘要组织", 2)
add_body(doc, "摘要按照背景、问题、方法、实验和结论五部分组织。首先说明建筑损毁快速评估对应急救援的重要性，以及灾后SAR全天候观测的优势；随后指出灾前RGB与灾后SAR存在模态差异和局部空间错位，基础模型容易产生局部漏检与类别混淆，而无条件局部修正又可能引入新的错误。")
add_body(doc, "方法部分概述基础语义分割、建筑引导候选构建、Direct与跨事件Memory互补证据以及Residual Utility Verifier。实验部分报告BRIGHT上的总体mIoU、类别IoU和F1、事件级性能、消融结果、漏检恢复率及修改可靠性。最终摘要仅填写真实无GT推理结果，不使用Oracle上限替代模型性能。")
add_heading(doc, "关键词", 2)
add_body(doc, "建筑损毁评估；光学–SAR遥感；语义分割；跨事件记忆；效用验证；选择性细化")

doc.add_page_break()

add_heading(doc, "I 引言", 1)
add_heading(doc, "第一段 研究背景与应用需求", 2)
add_body(doc, "介绍地震、洪水、火灾等灾害发生后建筑损毁信息对应急救援、资源配置和灾后重建的重要性。说明高分辨率遥感能够快速覆盖大范围区域，灾后SAR影像不受云雾和昼夜条件限制，灾前RGB影像则能够提供较清晰的建筑位置、轮廓和纹理信息。")

add_heading(doc, "第二段 现有研究发展", 2)
add_body(doc, "概述建筑损毁评估由灾前灾后光学影像逐渐发展到多任务网络、Transformer以及光学–SAR多模态方法。说明DamageFormer等方法将建筑定位与损毁分类进行解耦，BRIGHT进一步推动全天候、跨事件建筑损毁评估研究。")

add_heading(doc, "第三段 现有方法的主要不足", 2)
add_number(doc, "RGB与SAR之间存在局部空间错位和显著模态差异，对应位置的特征不一定表达同一地物。")
add_number(doc, "损毁证据通常集中在局部区域，单一预测头难以覆盖不同灾害事件中的复杂损毁模式。")
add_number(doc, "局部后处理或残差修正可能同时产生有效修改与有害修改，缺少修改前的效用判断和无效修改回退机制。")

add_heading(doc, "第四段 问题本质与研究目标", 2)
add_body(doc, "将研究问题概括为：如何在保持基础语义分割结果稳定性的前提下，利用当前区域证据与跨事件历史证据识别可修正错误，并只接受具有正效用的局部修改。图1建议展示RGB–SAR局部错位、P3局部漏检、Direct与Memory建议一致或冲突，以及有效修改和有害修改的差异。")

add_heading(doc, "第五段 方法概述", 2)
add_number(doc, "以P3或DamageFormer作为基础模型，生成初始损毁预测和多模态特征。")
add_number(doc, "基于建筑概率和SLIC超像素构造局部候选区域。")
add_number(doc, "利用Direct Head生成当前区域建议，并从排除同事件的Memory中检索跨事件相似证据。")
add_number(doc, "构建Residual Utility Verifier，预测保持P3、采用Direct或采用Memory的相对效用。")
add_number(doc, "只将可靠修改写回基础预测，形成最终建筑损毁语义图。")

add_heading(doc, "第六段 主要贡献", 2)
add_number(doc, "提出面向光学–SAR建筑损毁语义分割的可靠局部细化框架，在不改变原有语义分割任务的情况下，将当前区域判断与跨事件历史证据用于基础预测修正。")
add_number(doc, "构建Direct–Memory互补局部证据机制。Direct Head根据当前区域特征生成主要建议，跨事件Memory通过排除同事件检索提供Direct未覆盖的补充候选。该贡献需以M3超过M2为实验前提。")
add_number(doc, "提出残差效用验证与选择性回退方法，根据基础预测不确定性、Direct证据、Memory证据和区域质量，在保持基础结果、采用Direct和采用Memory之间进行选择。")
add_body(doc, "若M4没有超过M2，应将第二项贡献调整为跨事件Memory候选互补性分析，不能宣称Memory提高了最终模型性能。")

add_heading(doc, "第七段 文章结构", 2)
add_body(doc, "第II节介绍建筑损毁语义分割、光学–SAR多模态学习以及记忆增强与可靠细化相关研究；第III节介绍总体框架、局部候选构建、Direct–Memory证据推理和效用验证方法；第IV节介绍数据集、实验设置、定量比较、定性分析、消融研究和过程分析；第V节总结全文。")

doc.add_page_break()

add_heading(doc, "II 相关工作", 1)
add_heading(doc, "A 建筑损毁遥感评估", 2)
add_body(doc, "介绍xBD、xView2、BRIGHT等数据集及建筑损毁任务的发展，梳理U-Net、孪生网络、ChangeOS、DamageFormer和ChangeMamba等代表性方法。重点指出现有工作主要关注整体语义分割精度，对局部错误的可控修正研究相对不足。")

add_heading(doc, "B 光学–SAR多模态建筑损毁识别", 2)
add_body(doc, "介绍灾前光学影像提供建筑结构先验、灾后SAR提供全天候观测的基本思路。总结多模态编码、特征对齐、早期融合和后期融合方法，并指出跨模态表征差异和局部配准误差仍会影响损毁等级判断。")

add_heading(doc, "C 区域级证据聚合与预测细化", 2)
add_body(doc, "介绍超像素、区域特征聚合、局部Token、注意力池化及多实例学习等方法。说明区域建模能够增强局部空间一致性，但区域级修改如果缺少可靠性控制，也可能放大基础模型错误。")

add_heading(doc, "D 记忆增强学习与可靠性估计", 2)
add_body(doc, "介绍类别原型、Memory Bank、近邻检索和跨事件知识迁移方法，以及置信度、预测熵、类别margin和选择性预测等可靠性估计方法。进一步指出现有Memory通常保存类别信息，较少显式记录某次修改是否真正改善基础预测的历史效用。")

doc.add_page_break()

add_heading(doc, "III 所提方法", 1)
add_heading(doc, "A 问题定义与框架概述", 2)
add_body(doc, "给定灾前RGB影像X_rgb和灾后SAR影像X_sar，基础模型输出背景、完好、受损和毁坏类别的像素级预测。模型最终仍执行语义分割，不要求输出独立建筑实例ID。")
add_equation(doc, "Y_base = f_base(X_rgb, X_sar)")
add_body(doc, "图2建议按照“输入与基础模型—建筑引导候选区域—Direct与Memory证据—Utility Verifier—局部回写与最终损毁图”的顺序绘制。")

add_heading(doc, "B 建筑引导的局部候选与特征表示", 2)
add_heading(doc, "1 建筑候选区域生成", 3)
add_body(doc, "由基础模型获得建筑概率图，通过阈值筛选预测建筑区域，并结合RGB影像的颜色和空间连续性生成SLIC超像素。超像素只作为局部证据和预测修改单元，不改变语义分割任务。")
add_heading(doc, "2 局部区域特征构建", 3)
add_body(doc, "对每个候选区域聚合RGB与SAR融合特征、基础模型类别概率、预测熵、类别margin、平均建筑概率、超像素面积、边界强度及可用的对齐可靠性特征，形成区域描述向量。")
add_equation(doc, "r_j = [ f_j^(rgb-sar),  p_j^base,  q_j^region ]")

add_heading(doc, "C Direct与跨事件Memory互补证据推理", 2)
add_heading(doc, "1 Direct Head当前证据", 3)
add_body(doc, "Direct Head直接读取当前区域描述，输出区域类别概率或修改建议。它是主要建议器，反映当前影像内部的局部损毁证据。")
add_heading(doc, "2 跨事件Memory检索", 3)
add_body(doc, "以区域描述为查询，在训练集Memory Bank中检索Top-k相似区域。训练和检索时排除同一图像，并默认排除同一灾害事件。Memory输出类别posterior、最大相似度、Top-k一致性、邻居来源事件数量和事件熵等信息。")
add_heading(doc, "3 Direct–Memory关系建模", 3)
add_body(doc, "将候选关系划分为二者一致、仅Direct修改、仅Memory修改和二者冲突。关系类型、置信度差异及相似度共同作为Verifier输入。")

add_heading(doc, "D Residual Utility Memory与效用验证器", 2)
add_heading(doc, "1 修改效用定义", 3)
add_body(doc, "对每个候选区域形成保持基础预测、采用Direct建议和采用Memory建议三种假设，并使用训练集真实标签计算Direct和Memory相对基础预测的效用。")
add_equation(doc, "H_0 = Keep Base,     H_D = Use Direct,     H_M = Use Memory")
add_equation(doc, "U_D = Q(H_D) - Q(H_0),     U_M = Q(H_M) - Q(H_0)")
add_heading(doc, "2 Residual Utility Memory", 3)
add_body(doc, "除区域特征和类别外，进一步保存建议来源、建议方向、采用建议后的效用、Accept或Reject状态以及来源事件。推理时判断当前建议更接近历史上的有效修改还是有害修改。")
add_heading(doc, "3 Utility Verifier", 3)
add_body(doc, "Verifier读取基础预测、Direct证据、Memory证据、二者关系和区域质量，预测不同选择的相对效用。第一版采用Logistic Regression或低容量线性分类器，避免复杂模型造成的过拟合与归因混淆。")
add_equation(doc, "a* = argmax { 0,  U_D_hat,  U_M_hat }")

add_heading(doc, "E 选择性局部修正", 2)
add_body(doc, "只有Verifier预测为正效用的建议才写回基础预测；当Direct和Memory均不可靠时，模型回退至基础预测。最终输出仍为像素级建筑损毁语义图。")
add_equation(doc, "Y_final = Y_base + G_D ⊙ R_D + G_M ⊙ R_M")

add_heading(doc, "F 训练与推理策略", 2)
add_table(doc, ["阶段", "处理流程", "数据约束"], [
    ["训练", "冻结基础模型；生成候选；计算Utility；构建Residual Memory；训练Verifier", "只使用Train；检索排除同图和同事件"],
    ["验证", "生成候选；查询冻结Memory；Verifier选择；计算指标", "Val只评估，不写入Memory，不使用GT选择"],
    ["测试", "冻结模型、阈值和Memory后进行一次性推理", "在方法和参数确定前不访问Test"],
], widths=[2.4, 8.0, 5.0])
add_body(doc, "如果局部对齐模块尚未取得独立增益，应将其作为基础特征预处理或补充消融，不直接列为主要创新模块。")

doc.add_page_break()

add_heading(doc, "IV 实验与讨论", 1)
add_heading(doc, "A 数据集与实验设置", 2)
add_heading(doc, "1 数据集与划分", 3)
add_body(doc, "介绍BRIGHT的灾前RGB、灾后SAR、建筑损毁标签和灾害事件构成，明确Train、Val和Test的使用原则。Memory构建、Utility标签生成和Verifier训练均只使用Train。")
add_heading(doc, "2 评价指标", 3)
for item in ("总体mIoU和各类别IoU", "各类别F1，重点关注受损与毁坏类别", "事件级mIoU及提高事件数量", "有效修改率和有害修改率", "Affected漏检恢复率", "推理时间、Memory检索开销和峰值显存"):
    add_bullet(doc, item)
add_heading(doc, "3 实现细节", 3)
add_body(doc, "说明基础模型、输入尺寸、SLIC参数、候选筛选规则、Memory容量、Top-k、相似度度量、Verifier模型、训练策略和随机种子。所有对照实验保持数据划分、基础权重和评价脚本一致。")

add_heading(doc, "B 与代表性方法的定量比较", 2)
add_body(doc, "在统一数据划分与评价协议下，将最终方法与基础模型及代表性光学–SAR建筑损毁方法进行比较。主表报告总体指标、类别指标和必要的资源开销。")

add_heading(doc, "C 主要方法结果", 2)
add_table(doc, ["配置", "方法内容", "回答的问题"], [
    ["M0", "P3基础模型", "基础性能是多少"],
    ["M1", "Direct固定规则", "不训练Verifier时直接修改是否有效"],
    ["M2", "Direct Logistic Verifier", "没有Memory时可训练Verifier的真实基线"],
    ["M3", "Direct加普通Memory证据", "现有Memory信息是否帮助Verifier"],
    ["M4", "Residual Utility Memory", "历史修改效用是否优于普通类别Memory"],
    ["M5", "打乱Utility标签", "收益是否来自正确历史效用而非参数增加"],
], widths=[2.0, 6.4, 7.0])
add_body(doc, "核心判断为M3是否超过M2，以及M4是否进一步超过M3。若Memory未形成真实增益，应保留Direct Verifier作为最终方法，并将Memory结果降级为互补性分析。")

add_heading(doc, "D 定性结果分析", 2)
add_body(doc, "展示P3基础预测、Direct建议、Memory建议、Verifier选择、最终预测和真实标签。案例应覆盖Memory有效、Memory有害、Direct与Memory一致以及二者冲突四类情况。")

add_heading(doc, "E 消融研究", 2)
add_table(doc, ["消融方向", "建议比较内容", "目标"], [
    ["候选区域", "无SLIC、当前SLIC、不同尺度和建筑概率阈值", "验证候选覆盖和空间组织作用"],
    ["Memory证据", "移除posterior、相似度、Top-k一致性、事件熵和关系特征", "识别有效Memory信息"],
    ["Verifier特征", "Base、Base加Direct、加入Memory、加入区域质量和可靠性", "验证决策信息来源"],
    ["Utility监督", "类别Memory、Residual Utility Memory、打乱Utility标签", "验证效用记忆真实性"],
    ["数据隔离", "普通检索与排除同事件检索", "验证跨事件迁移而非同场景记忆"],
], widths=[3.0, 7.5, 4.9])

add_heading(doc, "F 机制与过程分析", 2)
add_heading(doc, "1 Oracle上限与互补性", 3)
add_body(doc, "报告P3、Direct、Memory和联合候选的Oracle上限，说明Oracle只用于衡量候选可修正空间。重点报告联合Oracle相对Direct Oracle的额外增益，而不是将联合Oracle作为最终性能。")
add_heading(doc, "2 候选重叠与冲突", 3)
add_body(doc, "统计Direct与Memory一致、Direct-only、Memory-only和冲突候选，并计算不同类型候选的正效用比例。")
add_heading(doc, "3 漏检恢复漏斗", 3)
add_equation(doc, "P3受损漏检 → 预测建筑覆盖 → SLIC候选覆盖 → 建议Affected → Verifier接受 → 正确恢复")
add_heading(doc, "4 事件级泛化", 3)
add_body(doc, "报告不同事件的性能变化，分析Memory在哪些灾害类型和事件中提供Direct没有的补充证据，并检查是否存在少数事件明显下降。")
add_heading(doc, "5 效率分析", 3)
add_body(doc, "报告参数量、训练时间、推理时间、Memory容量、检索开销和峰值显存，说明可靠局部细化的额外成本。")

add_heading(doc, "G 局限性讨论", 2)
add_body(doc, "讨论基础模型建筑覆盖范围、SLIC候选质量、灾害类型差异和Memory邻居质量对方法的影响。说明Oracle互补性并不必然转化为真实推理增益，最终效果取决于Utility标签质量和Verifier可分性。")

doc.add_page_break()

add_heading(doc, "V 结论与未来研究方向", 1)
add_body(doc, "总结本文围绕建筑损毁语义分割中的局部漏检、跨事件模式差异和修正可靠性问题，构建Direct–Memory互补证据与效用验证机制。结论只报告真实无GT推理结果；联合Oracle结果只用于说明Memory候选具有额外互补空间。")
add_body(doc, "未来工作可考虑扩展建筑外围候选、多尺度区域划分和更高效的跨事件检索，但当前论文不继续扩展到实例分割、重新标注或大型新主干网络。")

add_heading(doc, "建议图表清单", 1)
add_table(doc, ["编号", "建议内容", "所在章节"], [
    ["图1", "典型问题与研究动机：漏检、互补候选、有益和有害修改", "引言"],
    ["图2", "整体模型框架", "方法概述"],
    ["图3", "SLIC局部候选与Direct–Memory证据构建", "方法"],
    ["图4", "Residual Utility Memory与Verifier", "方法"],
    ["图5", "BRIGHT数据与不同灾害样例", "实验设置"],
    ["表1", "与代表性方法的总体比较", "定量比较"],
    ["表2", "M0至M5主要实验", "主要结果"],
    ["表3", "模块与Verifier特征消融", "消融研究"],
    ["表4", "各事件逐事件结果", "事件级分析"],
    ["图6", "典型场景定性结果", "定性分析"],
    ["图7", "Direct与Memory候选关系和效用比例", "机制分析"],
    ["图8", "Affected漏检恢复漏斗", "机制分析"],
    ["图9", "归一化混淆矩阵", "错误分析"],
    ["表5", "参数量、推理时间和显存开销", "效率分析"],
], widths=[2.0, 9.5, 4.2])

doc.add_page_break()
add_heading(doc, "实验完成后的题目与贡献决策", 1)
add_table(doc, ["实验结果", "论文处理方式"], [
    ["M4稳定超过M3和M2", "以Residual Utility Memory作为核心创新，采用RUM-BDA题目"],
    ["M3超过M2但M4无额外提升", "保留普通Memory证据，弱化Residual Utility Memory表述"],
    ["M3和M4均未超过M2", "以Direct Utility Verifier为最终方法，Memory仅作为候选互补性分析"],
    ["M2也未超过M0", "停止扩大结构，重新检查Utility标签、数据泄漏、候选定义和评价协议"],
], widths=[5.0, 10.7])

add_heading(doc, "建议执行顺序", 1)
add_number(doc, "冻结P3、数据划分、评价脚本和当前SLIC参数，建立不可变的M0基线。")
add_number(doc, "依次完成M2、M3、M4和M5，先验证Memory能否在无GT推理中超过Direct Verifier。")
add_number(doc, "确定最终方法后，再补充特征消融、事件级分析、定性结果和效率实验。")
add_number(doc, "实验进行的同时撰写引言、相关工作和方法章节；最终指标确定后完成摘要、实验结论和题目。")
add_number(doc, "全文检查结果归因，明确区分Oracle上限、候选互补性和真实推理性能。")

# Footer with page number
footer = sec.footer
p = footer.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
fld_char1 = OxmlElement("w:fldChar")
fld_char1.set(qn("w:fldCharType"), "begin")
instr_text = OxmlElement("w:instrText")
instr_text.set(qn("xml:space"), "preserve")
instr_text.text = " PAGE "
fld_char2 = OxmlElement("w:fldChar")
fld_char2.set(qn("w:fldCharType"), "end")
r = p.add_run()
r._r.append(fld_char1)
r._r.append(instr_text)
r._r.append(fld_char2)
set_run_font(r, size=9)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(str(OUTPUT))
