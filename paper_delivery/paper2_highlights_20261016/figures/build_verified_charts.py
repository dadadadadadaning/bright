#!/usr/bin/env python3
"""Plot source-reported manuscript tables; no estimated experiment scores.

Comparison bars and means are arithmetically derived from printed scores.
These plots establish numerical associations, not module-level causality.
The optional historical curve belongs to a different experimental version.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document


DEFAULT_DOCX = Path("/workspace/attachments/15490848-5f32-40a1-adfd-65fc7054e14c/论文2.docx")
DEFAULT_HISTORY = Path("/workspace/outputs/bright_paper_data_analysis_20261016/detailed_results_analysis.md")
OUT = Path(__file__).resolve().parent


def number(text):
    return float(re.fullmatch(r"\s*([+-]?\d+(?:\.\d+)?)[*†]?\s*", text).group(1))


def write_csv(name, rows):
    with (OUT / name).open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def save(fig, name):
    for ext in ("png", "pdf", "svg"):
        fig.savefig(OUT / f"{name}.{ext}", dpi=360, facecolor="white")
    plt.close(fig)


def clean(ax, grid="y"):
    ax.spines[["top", "right"]].set_visible(False)
    ax.tick_params(width=.6, length=3)
    ax.grid(axis=grid, color="#dce1e5", linewidth=.5, zorder=0)
    ax.set_axisbelow(True)


def main(docx, history):
    plt.rcParams.update({
        "font.family": "DejaVu Sans", "font.size": 8,
        "axes.titlesize": 9, "axes.labelsize": 8,
        "xtick.labelsize": 7.2, "ytick.labelsize": 7.2,
        "legend.fontsize": 7.1, "axes.linewidth": .6,
        "pdf.fonttype": 42, "ps.fonttype": 42,
        "svg.fonttype": "none",
    })
    document = Document(docx)
    metric_keys = ["localization_f1_percent", "classification_f1_percent", "oa_percent",
                   "fourclass_miou_percent", "background_iou_percent", "intact_iou_percent",
                   "damaged_iou_percent", "destroyed_iou_percent"]
    main_rows = []
    for row in document.tables[0].rows[1:]:
        cells = [c.text.strip() for c in row.cells]
        main_rows.append({"method": cells[0].replace("†", ""),
                          **dict(zip(metric_keys, map(number, cells[1:])))})
    by_method = {r["method"]: r for r in main_rows}
    assert by_method["Ours"]["fourclass_miou_percent"] == 69.64
    assert by_method["Ours"]["intact_iou_percent"] == 76.25
    write_csv("source_main_table.csv", main_rows)

    event_names = ["UNet", "DeepLabV3+", "SiamAttnUNet", "SiamCRNN", "DamageFormer", "ChangeMamba", "Ours"]
    event_rows = []
    printed_event_average = None
    for row in document.tables[1].rows[1:]:
        cells = [c.text.strip() for c in row.cells]
        values = {name: number(value) for name, value in zip(event_names, cells[1:])}
        if cells[0] == "Average":
            printed_event_average = values
            continue
        event_rows.append({"event": cells[0], **values,
                           "ours_minus_damageformer_pp": round(values["Ours"] - values["DamageFormer"], 8),
                           "ours_minus_changemamba_pp": round(values["Ours"] - values["ChangeMamba"], 8)})
    assert len(event_rows) == 14
    event_average = {k: float(np.mean([r[k] for r in event_rows])) for k in event_names}
    for k in event_names:
        assert round(event_average[k], 2) == printed_event_average[k], (k, event_average[k], printed_event_average[k])
    write_csv("source_event_table_and_derived_deltas.csv", event_rows)

    # Figure 9: foreground per-class scores and paired four-class comparisons.
    fig, axes = plt.subplots(1, 2, figsize=(17 / 2.54, 8.8 / 2.54), gridspec_kw={"width_ratios": [1.22, 1]})
    fig.subplots_adjust(left=.075, right=.99, bottom=.19, top=.76, wspace=.38)
    classes = ["Intact", "Damaged", "Destroyed"]
    colors = ["#b5bdc7", "#73839d", "#56a7aa", "#b64649"]
    methods = ["DamageFormer", "ChangeMamba", "MPF", "Ours"]
    x = np.arange(3)
    width = .19
    for j, (method, color) in enumerate(zip(methods, colors)):
        values = [by_method[method][f"{c.lower()}_iou_percent"] for c in classes]
        bars = axes[0].bar(x + (j - 1.5) * width, values, width=width, color=color, label=method, zorder=3)
        for bar, value in zip(bars, values):
            axes[0].text(bar.get_x() + bar.get_width() / 2, value + .7, f"{value:.2f}",
                         ha="center", va="bottom", fontsize=5.5, rotation=90)
    axes[0].set_xticks(x, classes)
    axes[0].set_ylim(0, 86)
    axes[0].set_ylabel("IoU (%)")
    axes[0].set_title("(a) Foreground class performance", loc="left", pad=34)
    axes[0].legend(loc="lower left", bbox_to_anchor=(-.02, 1.04), ncol=2, frameon=False, handlelength=1.2, columnspacing=1.2)
    clean(axes[0])

    all_classes = ["Background", "Intact", "Damaged", "Destroyed"]
    delta_rows = []
    for j, (baseline, color) in enumerate([("DamageFormer", "#73839d"), ("MPF", "#56a7aa")]):
        deltas = [round(by_method["Ours"][f"{c.lower()}_iou_percent"] - by_method[baseline][f"{c.lower()}_iou_percent"], 8) for c in all_classes]
        positions = np.arange(4) + (j - .5) * .33
        axes[1].bar(positions, deltas, width=.31, color=color, label=f"Ours − {baseline}", zorder=3)
        for xpos, delta, cls in zip(positions, deltas, all_classes):
            axes[1].text(xpos, delta + (.10 if delta >= 0 else -.10), f"{delta:+.2f}",
                         ha="center", va="bottom" if delta >= 0 else "top", fontsize=6)
            delta_rows.append({"reference": baseline, "class": cls, "ours_minus_reference_pp": delta})
    axes[1].axhline(0, color="#444444", linewidth=.7)
    axes[1].set_xticks(np.arange(4), ["BG", "Intact", "Damaged", "Destroyed"])
    axes[1].set_ylim(-.65, 5.8)
    axes[1].set_ylabel("IoU difference (pp)")
    axes[1].set_title("(b) Paired class differences", loc="left", pad=34)
    axes[1].legend(loc="lower left", bbox_to_anchor=(-.02, 1.04), ncol=1, frameon=False, handlelength=1.2)
    clean(axes[1])
    fig.text(.5, .025, "Source-reported Test scores; differences derived from displayed values.", ha="center", fontsize=6.8, color="#505960")
    write_csv("derived_class_deltas.csv", delta_rows)
    save(fig, "Fig9_class_changes")

    # Figure 10: equal-event comparisons. These are not pixel-pooled metrics.
    fig, axes = plt.subplots(1, 2, figsize=(17 / 2.54, 11.6 / 2.54), sharey=True)
    fig.subplots_adjust(left=.285, right=.98, bottom=.14, top=.88, wspace=.24)
    y = np.arange(len(event_rows))
    means = {}
    wins = {}
    for ax, baseline, key, tag in zip(axes, ["DamageFormer", "ChangeMamba"],
                                    ["ours_minus_damageformer_pp", "ours_minus_changemamba_pp"], ["a", "b"]):
        deltas = np.array([r[key] for r in event_rows])
        mean = float(deltas.mean())
        means[baseline] = mean
        wins[baseline] = int((deltas > 0).sum())
        ax.barh(y, deltas, height=.65, color=np.where(deltas >= 0, "#218b8e", "#be5b54"), zorder=3)
        ax.axvline(0, color="#3a4146", linewidth=.8)
        ax.axvline(mean, color="#324960", linewidth=.8, linestyle="--", label=f"Event mean {mean:+.2f} pp")
        for ypos, delta in zip(y, deltas):
            ax.text(delta + (.18 if delta >= 0 else -.18), ypos, f"{delta:+.2f}",
                    va="center", ha="left" if delta >= 0 else "right", fontsize=6.5)
        ax.set_xlim(-3.5, 12.3)
        ax.set_xlabel("Event mIoU difference (pp)")
        ax.set_title(f"({tag}) Ours − {baseline}", loc="left", pad=21)
        ax.legend(loc="lower left", bbox_to_anchor=(-.01, 1.0), frameon=False, handlelength=1.5, fontsize=6.7)
        ax.set_xticks([-2, 0, 4, 8, 12])
        clean(ax, grid="x")
    axes[0].set_yticks(y, [r["event"] for r in event_rows])
    axes[0].invert_yaxis()
    fig.text(.5, .035, "Equal-event comparisons derived from source-reported event scores (14 events).", ha="center", fontsize=6.8, color="#505960")
    assert round(means["DamageFormer"], 2) == 1.78
    assert round(means["ChangeMamba"], 2) == 2.65
    save(fig, "Fig10_event_deltas")

    historical_rows = []
    if history.is_file():
        content = history.read_text(encoding="utf-8")
        for line in content.splitlines():
            pieces = [v.strip() for v in line.strip("|").split("|")]
            if len(pieces) == 6 and re.fullmatch(r"\d+", pieces[0]) and all(re.fullmatch(r"\d+\.\d+", v) for v in pieces[1:]):
                historical_rows.append(dict(zip(["step", "grid_probq", "A", "B", "C", "D"], [int(pieces[0])] + list(map(float, pieces[1:])))))
        assert len(historical_rows) == 11
        assert historical_rows[-1]["D"] == 61.6356
        write_csv("historical_top64_twoslot_val_sequence.csv", historical_rows)
        fig, ax = plt.subplots(figsize=(17 / 2.54, 8.7 / 2.54))
        fig.subplots_adjust(left=.11, right=.98, bottom=.19, top=.82)
        legend = {"A": "A: spatial slots / local", "B": "B: soft slots / local",
                  "C": "C: spatial slots / Train", "D": "D: soft slots / Train"}
        for key, color in zip("ABCD", ["#87919d", "#b79460", "#348a9b", "#a94750"]):
            ax.plot([r["step"] for r in historical_rows], [r[key] for r in historical_rows],
                    color=color, marker="o", markersize=3.3, linewidth=1.2, label=legend[key])
        ax.axhline(historical_rows[0]["grid_probq"], color="#313d46", linestyle="--", linewidth=1.0, label="Frozen grid_probq = 61.9027%")
        ax.set_title("Historical Top-64 two-slot Reader experiment", loc="left", pad=10)
        ax.set_xlabel("Updates per arm")
        ax.set_ylabel("Val foreground mIoU (%)")
        ax.set_xlim(-80, 5080)
        ax.set_ylim(61.30, 62.24)
        ax.set_xticks(np.arange(0, 5001, 500))
        clean(ax)
        ax.legend(frameon=False, ncol=2, loc="upper right", fontsize=6.8)
        fig.text(.5, .035, "Historical validation records; distinct from the present Top-8 supervised-projection model.", ha="center", fontsize=6.5, color="#505960")
        save(fig, "FigS1_historical_reader_curve")

    metadata = {
        "schema": "paper2_verified_charts_v1",
        "source_docx": str(docx), "source_docx_sha256": hashlib.sha256(docx.read_bytes()).hexdigest(),
        "source_tables": [0, 1], "source_status": "source-reported numbers; not independently re-evaluated predictions",
        "main_table": main_rows, "event_table": event_rows,
        "printed_event_average": printed_event_average, "computed_equal_event_average": event_average,
        "derived_class_deltas": delta_rows,
        "derived_event_mean_deltas_pp": means, "positive_event_count_out_of_14": wins,
        "event_metric_scope": "Manuscript table 2 reports event scores without explicitly establishing foreground versus four-class averaging; chart retains generic event mIoU label.",
        "historical_curve": {"source": str(history), "rows": historical_rows,
                             "version_boundary": "Top64 cosine parent retrieval and two-slot Reader; not present Top8 supervised projection; supplementary diagnostic only"},
        "verification": {"event_means_match_printed_to_2dp": True,
                         "no_estimated_scores": True, "no_confidence_intervals_inferred": True,
                         "no_existing_figure_modified": True},
    }
    (OUT / "verified_chart_data.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"event_delta_means_pp": means, "positive_event_counts": wins,
                      "output": str(OUT)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", type=Path, default=DEFAULT_DOCX)
    parser.add_argument("--historical-source", type=Path, default=DEFAULT_HISTORY)
    args = parser.parse_args()
    main(args.docx, args.historical_source)
