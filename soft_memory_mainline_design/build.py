from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

OUT = Path(__file__).parent
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
FONT = 'Noto Sans CJK SC'
C = dict(navy='16324F', ink='243746', muted='617487', blue='3478B9', bluebg='EDF5FC', teal='147D79', tealbg='E9F6F3', gold='B47C16', goldbg='FFF5DD', coral='B6524C', coralbg='FFF0EB', gray='ECF0F4', line='C8D3DE', white='FFFFFF')

def rgb(s):
    return RGBColor.from_string(C.get(s, s))

def text(slide, x,y,w,h,content,size=17,color='ink',bold=False,align=PP_ALIGN.LEFT):
    sh=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    for i,line in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.alignment=align
        p.space_after=Pt(5)
    return sh

def box(slide,x,y,w,h,title,body='',fill='bluebg',accent='blue',size=17,body_size=14,text_color=None):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill);sh.line.color.rgb=rgb(accent);sh.line.width=Pt(1.2)
    sh.adjustments[0]=0.12
    if body:
        text(slide,x+.15,y+.12,w-.3,.38,title,size,accent,True)
        text(slide,x+.15,y+.60,w-.3,h-.68,body,body_size)
    else:
        tx=text(slide,x+.12,y+.10,w-.24,h-.20,title,size,text_color or accent,True,PP_ALIGN.CENTER)
        tx.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
    return sh

def arrow(slide,x1,y1,x2,y2,color='muted',width=1.7,dash=False):
    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb=rgb(color);sh.line.width=Pt(width)
    ln=sh._element.spPr.get_or_add_ln()
    end=OxmlElement('a:tailEnd');end.set('type','triangle');end.set('w','sm');end.set('len','sm');ln.append(end)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    return sh

def line(slide,x1,y1,x2,y2,color='line',width=1):
    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb=rgb(color);sh.line.width=Pt(width)
    return sh

def slide(title,subtitle,notes):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    text(s,.5,.35,12.35,.5,title,29,'navy',True)
    text(s,.5,.96,12.35,.4,subtitle,15,'muted')
    line(s,.5,1.43,12.83,1.43)
    text(s,.5,7.10,11.6,.22,'方案设计稿｜图表待真实实验填充 · 固定 Q 的 soft＋Memory 联合分类',9,'muted')
    text(s,12.15,7.07,.65,.25,f'{len(prs.slides):02d} / 09',10,'muted',align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text=notes
    return s

def pill(s,x,y,w,label,fill='tealbg',color='teal'):
    return box(s,x,y,w,.42,label,fill=fill,accent=fill,size=12,text_color=color)

def table(s,x,y,width,headers,rows,col_widths=None,row_h=.62,font=14):
    n=len(headers);rh=.57
    tbl=s.shapes.add_table(len(rows)+1,n,Inches(x),Inches(y),Inches(width),Inches(rh+row_h*len(rows))).table
    if col_widths:
        for c,w in zip(tbl.columns,col_widths):c.width=Inches(w)
    for i,row in enumerate([headers]+rows):
        tbl.rows[i].height=Inches(rh if i==0 else row_h)
        for j,val in enumerate(row):
            cell=tbl.cell(i,j);cell.text=str(val);cell.margin_left=Inches(.11);cell.margin_right=Inches(.08);cell.margin_top=Inches(.08);cell.margin_bottom=Inches(.03)
            cell.fill.solid();cell.fill.fore_color.rgb=rgb('navy' if i==0 else ('white' if i%2==0 else 'gray'))
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.font.name=FONT;p.font.size=Pt(font);p.font.bold=(i==0);p.font.color.rgb=rgb('white' if i==0 else 'ink')
    return tbl

# 1
s=slide('主线结构：像素、soft 区域与历史案例共同判断','固定新的 Q-baseline 前端；先验证信息组织与读取，暂不叠加 SAR→Q 改造。',
'''所有形状可编辑。主线是假设而非已证实结论。A1 原已融合灾前 RGB 与灾后 SAR；新增 SAR→Q 是不同实验。Q 使用本轮新 Q-baseline，不能用旧 E0 或旧三组成绩充当同源消融。P0 是同次 grid_probq。联合输出 softmax(log(P0)+delta)；不存在是否采纳 Memory 的拒绝阈值。父区域 T128 用于检索，子槽只改变内容组织。Train GT 只进入支持标签与训练损失，查询推理无 GT。''')
box(s,.55,2.05,2.15,2.20,'Frozen front','灾前 RGB\n灾后 SAR\n冻结 A1＋Q\nF256 · B · Q',size=16,body_size=16)
box(s,3.13,1.68,2.9,.62,'Pixel F256 · 细节直连',size=15)
box(s,3.13,2.52,2.9,.69,'P0 = grid_probq',size=18)
box(s,3.13,3.43,2.9,.69,'Parent：F256＋T128',size=15)
box(s,3.13,4.34,2.9,.85,'Soft slots：QB / Q(1−B)',fill='tealbg',accent='teal',size=16)
for yy in [1.99,2.86,3.78,4.75]:
    line(s,2.9,1.99,2.9,4.75,'blue');arrow(s,2.9,yy,3.13,yy,'blue')
arrow(s,2.7,3.15,2.9,3.15,'blue')
box(s,6.5,3.30,2.65,1.22,'Memory Reader','两个查询槽读取全部合法子槽\n父 T128 检索 Top64',fill='goldbg',accent='gold',body_size=13)
arrow(s,6.03,3.78,6.50,3.78,'gold')
arrow(s,6.03,4.75,6.50,4.25,'teal')
box(s,6.5,5.28,2.65,1.16,'Train-only Memory','父 key／子槽特征／软标签\n排除同图、同事件',fill='goldbg',accent='gold',body_size=13)
arrow(s,7.82,5.28,7.82,4.52,'gold')
box(s,3.13,5.65,2.9,.70,'Train 图像＋Train GT',fill='gray',accent='muted',size=16)
arrow(s,6.03,6.0,6.5,6.0,'gold')
text(s,3.15,6.45,5.9,.27,'GT → 子槽软标签；不进入 Val／Test 查询输入',12,'muted')
box(s,10.05,2.45,2.70,2.30,'Joint pixel classifier','像素＋父区域＋双槽\n＋Memory 回投＋P0\n\nsoftmax(log P0＋Δ)',fill='tealbg',accent='teal',body_size=16)
line(s,9.60,1.99,9.60,4.95,'teal',1.6)
for x,yy in [(6.03,1.99),(6.03,2.86),(6.03,4.95)]:
    if yy==2.86:
        line(s,x,yy,6.23,yy,'teal');line(s,6.23,yy,6.23,2.08,'teal');line(s,6.23,2.08,9.60,2.08,'teal')
    else:line(s,x,yy,9.60,yy,'teal')
line(s,6.03,3.78,6.16,3.78,'blue',1.7)
line(s,6.16,3.78,6.16,3.10,'blue',1.7)
arrow(s,6.16,3.10,10.05,3.10,'blue',1.7)
text(s,6.55,2.67,2.65,.30,'父区域按 Q 回投',13,'blue',True)
arrow(s,9.15,3.90,9.60,3.90,'gold')
arrow(s,9.60,3.50,10.05,3.50,'teal')
box(s,10.05,5.35,2.70,.85,'四类像素预测',fill='bluebg',accent='blue',size=21)
arrow(s,11.4,4.75,11.4,5.35,'teal')
s.notes_slide.notes_text_frame.text += '\nTrain-only Memory 使用与当前查询相同的冻结前端和 Q 建库；不存在单独的第三套前端。'
text(s,9.35,6.35,3.45,.54,'父区域按 Q；双槽与 Memory\n按 Q0／Q1 回投',12,'muted')

# 2
s=slide('Soft 分流：改变信息组织，不删除背景证据','B 是冻结前端预测的建筑概率；Q0、Q1 是同一父区域的两个角色。',
'''Q_{ik} 表示位置 i 对父区域 k 的分配。子槽权重 w_{ik0}=Q_{ik}(1-B_i)，w_{ik1}=Q_{ik}B_i。特征均值分母是该槽总权重；软标签分母只包含有效标签位置，忽略255。支持特征和支持标签必须同一权重，不能复制父标签。空槽 mask 需一致。两槽都可以输出任意损毁类；偏背景槽可能包含被前端漏检的真建筑。Q0+Q1=Q，但每槽单独归一化的汇总后再回投不等于原父均值。''')
box(s,.55,1.72,3.45,1.20,'保留父区域','同一 Q、同一 B、同一像素特征',body_size=16)
box(s,4.40,1.72,3.75,1.20,'偏建筑槽','w1 = Q × B',fill='tealbg',accent='teal',body_size=24)
box(s,8.55,1.72,4.22,1.20,'偏背景槽','w0 = Q × (1 − B)',fill='goldbg',accent='gold',body_size=24)
arrow(s,4.00,2.32,4.40,2.32)
line(s,2.27,2.92,2.27,3.07);line(s,2.27,3.07,10.66,3.07);arrow(s,10.66,3.07,10.66,2.92)
box(s,.55,3.25,5.95,1.65,'特征与标签使用同一槽权重','特征：Fs = Σi wis Fi / Σi wis（Σi wis > 0）\n标签：ys = Σi validi wis onehot(yi) / Σi validi wis',fill='tealbg',accent='teal',body_size=16)
box(s,6.85,3.25,5.92,1.65,'对应回投，保留原像素路径','槽输出 → 对应 Q0／Q1 → 当前像素\nQ0 + Q1 = Q；空槽使用显式有效标记',body_size=17)
box(s,.55,5.30,3.85,1.28,'不强制类别','两个角色都预测四类；\n偏背景槽也可包含受损线索。',fill='gray',accent='muted',body_size=15)
box(s,4.75,5.30,3.85,1.28,'不宣称建筑实例约束','没有逐栋身份，不保证连通、\n不保证不跨相邻建筑。',fill='gray',accent='muted',body_size=15)
box(s,8.95,5.30,3.82,1.28,'父 key 与子内容分开','本轮保留父检索；重新生成\n两槽特征及对应软标签。',fill='gray',accent='muted',body_size=15)

# 3
s=slide('Memory 读取：按父区域找案例，按子槽读内容','64 个父案例最多展开为 128 个槽；不是 128 个独立历史案例。',
'''两个查询槽并非只读同名支持槽，它们均可注意全部选中的合法子槽。检索父 key 固定；查询和库的子槽遵循相同规则。候选字段：子槽 F256、父 T128、槽质量或比例、有效标记、四类软标签。注意力 K 由特征生成，标签作为 V 的信息；事件和图像元数据只用于排除，不作为事件 embedding。无支持时上下文为零，不能隐藏学习到的替代 Memory。R0/A/B 图内参考标签字段是模型预测分布，不是查询GT。''')
steps=[('Query parent T128','冻结父区域 key'),('Top64 parent retrieval','同图／同事件排除'),('Expand child values','每个父案例 0–2 个有效槽')]
for i,(a,b) in enumerate(steps):
    x=.55+i*4.18;box(s,x,1.75,3.85,1.08,a,b,fill='goldbg',accent='gold',size=17,body_size=14)
    if i<2:arrow(s,x+3.85,2.30,x+4.18,2.30,'gold')
table(s,.55,3.10,5.00,['每个支持子槽字段','作用'],[
    ['Child F256＋Parent T128','局部证据＋父上下文'],['槽占比＋valid','质量与空槽处理'],['Train 四类软标签','案例的标注证据'],['图像／事件 ID','仅用于合法性排除']
],col_widths=[2.70,2.30],row_h=.51,font=14)
box(s,6.02,3.10,3.02,2.10,'两个 Query slots','均读取全部合法候选槽\n\nK ← 特征\nV ← 特征＋软标签',fill='tealbg',accent='teal',body_size=17)
arrow(s,5.55,4.10,6.02,4.10,'gold')
box(s,9.50,3.10,3.27,2.10,'Slot context → Pixel','对应 Q0／Q1 回投\n与像素、父区域共同分类\n\n无支持：上下文置零',body_size=16)
arrow(s,9.04,4.10,9.50,4.10,'teal')
text(s,.65,6.05,12.0,.57,'需要记录：合法支持数、独立来源数、类别软质量、空槽率；“找到邻居”不等于“找到有用的受损案例”。',16,'muted')

# 4
s=slide('四格主实验：同时检验 soft 与历史参考的增量','所有组使用同源 Q、像素路径、父区域、P0、相同 Reader 结构与训练预算。',
'''A/B参考图内空间近邻，含查询自身，标签字段来自模型预测。C/D历史案例使用Train标签；因此C-A和D-B是整套历史参考的增量，不单独识别排序或标注来源的贡献。固定棋盘格式空间硬分组是已有capacity control，不完全控制soft权重连续性；后续如做纯建筑语义解释需概率分布匹配的B错位/置换对照。旧R0/RT/RF来自旧Q及不同后端时不能作为本轮单槽因果对照。''')
table(s,.55,1.75,12.23,['区域组织方式','图内参考：模型预测分布','Train Memory：标注软分布'],[
    ['固定空间双槽（capacity）','A · 空间分组＋图内参考','C · 空间分组＋历史案例'],
    ['建筑概率双槽（soft）','B · 建筑分流＋图内参考','D · 建筑分流＋历史案例']
],col_widths=[3.40,4.42,4.41],row_h=.93,font=18)
box(s,.55,4.45,5.95,1.60,'Soft 的增量','B − A：无 Memory 时的分流作用\nD − C：有 Memory 时的分流作用',fill='tealbg',accent='teal',body_size=18)
box(s,6.85,4.45,5.93,1.60,'历史参考的增量','C − A：空间双槽条件下\nD − B：soft 双槽条件下',fill='goldbg',accent='gold',body_size=18)
text(s,.65,6.35,12,.37,'交互量 = D − C − B + A　｜　D 稳定优于 B、C 和同轮基线，才支持保留完整组合。',17,'navy',True)

# 5
s=slide('消融分两轮做：先确认收益，再解释来源','每一项都回答一个问题；候选与对应对照使用匹配种子、初始化和训练图序。',
'''先固定短训练规则，再用3个种子重做四格。不得根据Test调步数。单槽对照要匹配Reader/分类器容量和候选数尽可能接近，明确仍存在token数和计算量差异。随机支持/错配标签需训练与推理都采用同一规则；推理干预只能说明依赖或敏感性。删除直接像素分支只表示去掉直连；父区域和支持特征本来仍由像素汇总而来。B错位不使用Val真值，最好保持边际分布和尺度/平滑性，另记录mask相关性。''')
pill(s,.55,1.68,2.05,'第一轮 · 主结果')
text(s,2.90,1.70,9.80,.50,'A／B／C／D × 3 个种子；固定一个新 Q、一个案例库版本、一个训练规则。',17,'navy',True)
table(s,.55,2.47,12.23,['第二轮对照','保留什么／只改变什么','能支持的结论'],[
    ['同源单槽 Local／Memory','重新训练，匹配结构容量与预算','双槽组织的增量'],
    ['D：随机合法父案例','同一库、Top64、排除规则不变','检索排序是否有效'],
    ['D：无标签／错配标签','同容量编码；匹配训练','标注证据是否被有效使用'],
    ['D：错位建筑概率 B','保留 B 的分布，破坏位置对应','建筑概率空间对应是否有用'],
    ['可选：删除像素直连','保留所有区域路径，重新训练','直接像素细节是否必要']
],col_widths=[3.15,5.10,3.98],row_h=.62,font=15)
box(s,.55,6.32,12.23,.48,'容量对照仍有“硬空间分组 vs 连续 soft 权重”的差异；不要将 B−A 全部解释为建筑语义。',fill='coralbg',accent='coral',size=14)

# 6
s=slide('案例图一：展示同一位置的收益与代价','固定同一图、同一裁剪、同一放大倍率；每列由对应的真实模型输出生成。',
'''本页灰色框均为待填占位，不是实验结果。必须用同源同终点的真实模型预测。选择规则提前固定：按Val受损类、低B建筑、边界、跨事件等分层；每层给典型/中位案例及失败案例。不能将旧A1输出重命名为D。误差图建议将grid_probq错误而D正确标为绿色、grid_probq正确而D错误标为红色，两者都错为紫色；像素量化必须分别统计三个前景类别。若没有真实实例标注，不写逐栋准确率。''')
cols=['灾前 RGB','灾后 SAR','GT','grid_probq','A','B','C','D','D 对基线误差']
cw=1.29;xx=.55
for j,lbl in enumerate(cols):text(s,xx+j*1.37,1.65,cw,.35,lbl,12,'navy',True,PP_ALIGN.CENTER)
for row,yy in enumerate([2.15,4.25]):
    for j in range(9):
        box(s,xx+j*1.37,yy,cw,1.43,'待填真实图',fill='gray',accent='line',size=12,text_color='muted')
    text(s,.55,yy+1.55,12.23,.32,'正向案例｜建筑与背景混合减少，受损证据是否被保留' if row==0 else '失败案例｜低 B 漏楼、相邻楼混合或 SAR 位置偏差，明确发生了什么错误',15,'teal' if row==0 else 'coral',True)
legend=[('FFFFFF','背景'),('4DBD74','完好'),('F0B34D','受损'),('DA5A5A','毁坏')]
for i,(color,label) in enumerate(legend):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.65+i*1.35),Inches(6.48),Inches(.19),Inches(.19));sh.fill.solid();sh.fill.fore_color.rgb=rgb(color);sh.line.color.rgb=rgb('line');text(s,.9+i*1.35,6.43,1.05,.3,label,12,'muted')
text(s,6.1,6.40,6.6,.38,'误差图：救回／误伤／仍错分开显示；同时报告像素数。',13,'muted')

# 7
s=slide('案例图二：Memory 找到了谁，改变了什么判断','每个查询都显示原始预测与邻居证据；注意力权重只能作为辅助线索。',
'''灰色框均为待填占位。展示查询RGB/SAR、Q父区域轮廓、建筑概率B、Q0/Q1权重、P0及最终预测；邻居展示父案例真实RGB/SAR和区域掩膜、来源事件、相似度、两槽软标签分布、有效槽比例。本文选4个邻居只是视觉展示，模型依旧读取Top64。禁止只选最支持叙述的邻居；预先规定Top1/中位/相反标签/高权重等展示策略，并展示失败检索。Query GT可作离线解释但不能参与检索或模型输入。''')
query=['Query RGB / SAR','父 Q ＋ B','Q0 / Q1 权重','P0 → 最终预测']
for i,lbl in enumerate(query):
    x=.55+i*3.13;box(s,x,1.75,2.86,1.23,'待填查询图',fill='gray',accent='line',size=15,text_color='muted');text(s,x,3.06,2.86,.35,lbl,14,'navy',True,PP_ALIGN.CENTER)
arrow(s,6.5,3.42,6.5,3.73,'gold')
for i in range(4):
    x=.55+i*3.13
    box(s,x,3.92,2.86,1.23,f'真实邻居 {i+1} · RGB / SAR',fill='goldbg',accent='gold',size=13)
    text(s,x+.08,5.26,2.70,.68,'事件／相似度：待填\n两槽四类标签分布：待填',12,'muted')
text(s,.65,6.25,12,.55,'按预先规定的 Val 分层展示典型、中位与失败案例；用随机案例／错配标签消融验证因果贡献。',16,'navy',True)

# 8
s=slide('实验协议：把“短期信号”变成可复核的证据','数据划分、训练终点、学习率与报告指标先固定，再进行重复。',
'''固定Train450是从早期已见训练图像选取的诊断子集，不代表完整Train2918；需要报告抽样规则/各事件与类别覆盖，不能用差异自动断言过拟合。短训练建议固定500步截断，但保留原5000步学习率horizon，避免更换训练方案。三个种子3407、3408、3409是设计建议，不代表已执行。Test500已经探索，不进一步反复调参；可报告一次锁定配置的测试但必须披露已使用测试探索，若要求严格独立验证须额外未使用数据或外部集，不要偷偷改变标准benchmark。按事件配对bootstrap属于不确定性分析，不把高度相关像素当独立样本。''')
items=[('01','诊断已有检查点','500／1500／5000：固定 Train450＋Val450\n同一推理配置，loss／类别 IoU／混淆矩阵'),('02','锁定短训练规则','若选择 500 步截断，保留原 5000 步 LR 日程\n不能简单把总步数参数改为 500'),('03','配对重复并报告','建议种子 3407／3408／3409；同源 Q、库、Top64\n主指标由全 split 混淆矩阵汇总'),('04','控制结论与数据使用','报告逐事件差值、类别取舍、显存和时延\n500 步 Test878 已探索；后续选择只依据 Train／Val')]
for i,(n,a,b) in enumerate(items):
    yy=1.72+i*1.22
    box(s,.55,yy,.65,.65,n,fill='navy',accent='navy',size=18,text_color='white')
    text(s,1.47,yy-.01,3.25,.44,a,19,'navy',True)
    text(s,4.75,yy-.01,8.0,.80,b,16,'ink')
    if i<3:line(s,1.48,yy+1.00,12.76,yy+1.00)

# 9
s=slide('原 PPT 的改图清单：让每页对应一个可验证问题','现有 PPT 的数值不沿用；架构、协议和图例先定稿，再填入真实新实验结果。',
'''原附件含12页。本页是页面改造建议，原文件保持不变。1-3重复架构可合并为一张总图和一张soft/Memory机制图。第4页地图只有在符合本次split时保留，标注Train/Val/Test以及事件排除策略；不能混用不同BRIGHT赛道。5-7旧比较图放附录或重生成当前A-D预测。8-10替换机制和失败案例、固定Train-Val曲线，不以最好中间结果代表正式终点。11重做同源全量结果与四格消融；12展示逐事件配对变化和种子差异，不能将像素独立性夸大。全部表格先留空或写待运行，严禁虚构预期指标。''')
table(s,.55,1.74,12.23,['原页码','建议保留／更换内容','新页必须回答的问题'],[
    ['1–3','合并结构图；突出固定 Q、双槽与联合分类','模型如何组织证据？新增组件在哪里？'],
    ['4','保留数据地图，核对划分与跨事件排除规则','训练、支持案例、验证如何分开？'],
    ['5–7','旧 A1 比较入附录；重新生成 A／B／C／D 图','同一位置哪些错误被救回、哪些被误伤？'],
    ['8–10','soft 机制、检索失败、Train／Val 诊断图','为什么可能有效？失败发生在哪里？'],
    ['11','整体结果＋四格消融表，数据暂留空','组合超过基线、单独模块和同容量对照吗？'],
    ['12','逐事件与多种子差值；附计算成本','收益是否稳定，代价与边界是什么？']
],col_widths=[1.10,5.65,5.48],row_h=.63,font=15)
text(s,.65,6.37,12,.43,'论文主张随证据调整：整体 mIoU 提升 ≠ 受损类改善；相对增量 ≠ 稳定协同。',17,'coral',True)

for i,s in enumerate(prs.slides,1):
    for sh in s.shapes:
        assert sh.left>=0 and sh.top>=0, (i, sh.name, 'negative')
        assert sh.left+sh.width<=prs.slide_width+1000, (i,sh.name,'right')
        assert sh.top+sh.height<=prs.slide_height+1000, (i,sh.name,'bottom')

outfile=OUT/'soft_memory_mainline_design_editable.pptx'
prs.save(outfile)
print(outfile)
