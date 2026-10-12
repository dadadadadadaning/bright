from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(r"D:\software\dowmload\BRIGHT-master\论文写作\光学SAR建筑损毁小论文中文初稿_扩展版.docx")
BLUE = "1F4E78"
PALE = "EAF2F8"
GRAY = "F4F6F7"
BLACK = RGBColor(0, 0, 0)


def font(run, cn="宋体", en="Times New Roman", size=10.5, bold=False, color=BLACK):
    run.font.name = en
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), cn)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), en)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), en)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def para(doc, text, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=0):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(after)
    if indent:
        pf.first_line_indent = Cm(0.74)
    font(p.add_run(text))
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def eq(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    font(p.add_run(text), cn="Cambria Math", en="Cambria Math", size=11)
    return p


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, v=90, h=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", v), ("start", h), ("bottom", v), ("end", h)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def borders(table):
    tbl_pr = table._tbl.tblPr
    node = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "D9D9D9")
        node.append(e)
    tbl_pr.append(node)


def table(doc, headers, rows, widths=None, note=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders(t)
    hdr = t.rows[0]
    trpr = hdr._tr.get_or_add_trPr()
    rep = OxmlElement("w:tblHeader")
    rep.set(qn("w:val"), "true")
    trpr.append(rep)
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        shade(c, BLUE)
        margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        font(p.add_run(h), cn="黑体", size=9, bold=True, color=RGBColor(255, 255, 255))
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            c = cells[i]
            shade(c, PALE if ridx % 2 else "FFFFFF")
            margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            font(p.add_run(str(value)), size=8.7)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    if note:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        font(p.add_run(note), size=8.5)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    font(p.add_run(text), cn="宋体", size=9.5, bold=True)


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2.3), Cm(2.2)
sec.left_margin, sec.right_margin = Cm(2.6), Cm(2.4)

styles = doc.styles
styles["Normal"].font.name = "Times New Roman"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
styles["Normal"].font.size = Pt(10.5)
for name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
    s = styles[name]
    s.font.name = "Times New Roman"
    s._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = BLACK
    if name != "Title":
        s.paragraph_format.space_before = Pt(10)
        s.paragraph_format.space_after = Pt(4)

# Footer page number
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run()
font(run, size=9)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
run._r.addnext(fld)

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(80)
p.paragraph_format.space_after = Pt(20)
p.add_run("基于跨事件记忆与效用验证的\n光学–SAR建筑损毁语义分割方法")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("中文小论文初稿（扩展版）"), cn="黑体", size=15, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(30)
font(p.add_run("研究主线：局部候选构建—Direct/Memory互补证据—可靠效用验证—选择性写回"), size=11)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(175)
font(p.add_run("2026年9月"), size=11)
doc.add_page_break()

heading(doc, "摘  要", 1)
para(doc, "灾害发生后，快速获取建筑损毁的空间分布与严重程度，是应急救援力量调配、受灾人口研判和灾后恢复重建的重要信息基础。灾前高分辨率光学影像能够提供清晰的建筑轮廓与地物纹理，灾后合成孔径雷达影像则具有全天时、全天候观测能力，两者结合为复杂天气条件下的建筑损毁评估提供了可行的数据路径。然而，光学与SAR在成像机理、观测视角和空间响应方面存在显著差异，基础语义分割模型在跨事件应用时仍容易出现局部漏检、损毁等级混淆以及边界不稳定等问题。更重要的是，现有局部细化方法往往依据单一预测头或固定置信度阈值直接修改基础结果，缺少对“本次修改是否真正优于原预测”的显式判断，因而可能在修正错误的同时引入新的误判。")
para(doc, "针对上述问题，本文构建一种基于跨事件记忆与效用验证的光学–SAR建筑损毁语义分割框架。该框架首先利用现有基础模型获得四类像素级损毁预测及多模态特征，再结合建筑概率图与SLIC超像素生成局部候选区域；随后，由Direct Head根据当前区域特征给出主要修改建议，并从排除同图、排除同事件的Memory Bank中检索相似历史区域，形成跨事件补充证据；最后，Residual Utility Verifier联合基础预测不确定性、区域质量、Direct证据、Memory证据及二者关系，在“保持基础预测”“采用Direct建议”和“采用Memory建议”之间进行选择，仅将预测为正效用的局部修改写回原语义图。该设计不改变原有语义分割任务，也不要求重新构建实例级标签，其核心目标是在保留基础模型稳定性的同时，提高局部错误修正的可靠性。")
para(doc, "现阶段实验表明，基础模型P3的mIoU为60.5020%。在使用真实标签进行事后选择的Oracle分析中，Memory、Direct及二者联合的mIoU上限分别为63.5212%、63.8589%和64.4420%，联合候选相对基础模型具有3.9401个百分点的潜在改进空间，并比Direct单独Oracle高0.5831个百分点。候选关系统计进一步显示，Memory能够提供Direct未覆盖的补充区域，其中Memory-only候选的Oracle增益为0.5739个百分点。上述结果证明了跨事件历史证据具有候选互补性，但仍属于上限分析，不能替代无真实标签条件下的最终推理结果。本文据此给出M0—M5完整实验闭环，用于检验效用验证器能否将候选空间转化为稳定、可复现的实际性能增益。")
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(6)
font(p.add_run("关键词："), cn="黑体", bold=True)
font(p.add_run("建筑损毁评估；光学–SAR遥感；语义分割；跨事件记忆；局部细化；效用验证"))

heading(doc, "1 引言", 1)
para(doc, "地震、洪水、台风、山火和爆炸等突发事件往往在短时间内造成大范围建筑损毁。灾后早期阶段的道路阻断、通信中断和次生灾害会显著增加现场调查的难度，因此，能够快速覆盖大区域的遥感观测成为建筑损毁评估的重要技术手段。xBD数据集通过大规模灾前、灾后光学影像及建筑级损毁标注推动了深度学习方法的发展[1]；BRIGHT数据集进一步提供灾前光学和灾后甚高分辨率SAR配对影像，使全天候、多灾种和跨区域的建筑损毁研究具备统一的数据基础[2]。从应用角度看，建筑损毁图不仅要在总体指标上取得较高精度，还应尽可能减少受损建筑漏检，因为局部漏检可能直接影响救援优先级和资源配置。")
para(doc, "现有建筑损毁评估方法经历了从人工特征、面向对象分析到深度语义分割和多时相变化检测的发展。ChangeOS将面向对象思想与深度语义变化检测结合，实现建筑定位与损毁分类的协同建模[3]；ChangeMamba利用状态空间模型描述大范围时空关系，并将其扩展到建筑损毁评估任务[4]。BRIGHT的基准实验表明，将建筑定位和损毁分类解耦的模型通常具有较好的综合表现[2]。尽管这类方法能够输出完整的像素级损毁图，BRIGHT基准仍显示受损类别识别和跨事件迁移更具挑战性[2]，提示数量较少、形态复杂或事件差异较大的局部损毁区域仍可能被遗漏。")
para(doc, "对于灾前光学与灾后SAR组合，困难不仅来自类别不平衡，还来自模态差异与局部空间错位。光学影像反映地表的光谱与纹理信息，SAR影像则记录微波散射响应，并受到斑点噪声、侧视几何和叠掩等因素影响。即使经过全局配准，同一建筑在两种影像中的局部对应关系也可能不完全一致。传统研究已经尝试联合光学与SAR信息进行震后损毁评估[5]。BRIGHT实验表明多模态融合总体具有互补价值，但固定的对象级后处理在集合级与事件级指标上的收益并不一致[2]，说明局部预测修改仍需显式验证其可靠性。")
para(doc, "进一步地，xBD和BRIGHT均以建筑多边形及其损毁等级组织监督信息，因此同一建筑区域内的像素通常继承统一的建筑级损毁标签[1–2]。这并不意味着建筑每个局部都呈现相同物理损伤，而是说明当前任务使用语义图表达建筑级评估结论。基于这一标注特点，本文不将研究方向改为实例分割，而是使用超像素作为局部证据聚合和预测修改单元。SLIC能够在颜色与空间邻近性约束下形成相对紧凑、贴合边界的区域[6]，适合将像素级噪声整理为较稳定的候选单元；但超像素本身不提供实例编号，也不会改变最终输出的语义分割形式。")
para(doc, "本文关注的问题可以概括为：如何在保持基础语义分割结果稳定性的前提下，利用当前区域证据与跨事件历史证据识别可修正错误，并只接受具有正效用的局部修改。为此，本文不把Memory设计成覆盖全图的第二个分类器，而是把Direct Head作为当前区域的主要建议器，把跨事件Memory作为补充证据来源，再由效用验证器决定是否采纳修改。该思想与选择性预测中的“接受或拒绝”机制具有一致性：模型不仅需要给出建议，还需要识别哪些建议值得信任[7–8]。区别在于，本文的拒绝并非不输出结果，而是回退到基础预测，从而保证任何局部模块失效时仍保留一条稳定的推理路径。")
para(doc, "本文的主要贡献可归纳为三点。第一，提出一种不改变语义分割任务的建筑引导局部细化框架，利用建筑概率与超像素构造候选区域，将像素级预测、区域质量和多模态特征统一为可验证的局部描述。第二，构建Direct–Memory互补证据机制：Direct Head负责当前区域判断，排除同事件的Memory Bank负责检索跨事件相似历史证据，并显式描述二者的一致、单边建议和冲突关系。第三，引入Residual Utility Verifier，把历史区域中“修改是否改善基础预测”的残差信息作为监督，在保持基础结果、采用Direct或采用Memory之间进行选择，以降低无条件局部修正造成性能退化的风险。需要强调的是，第二、第三项能否作为最终有效创新，必须由真实无标签推理条件下的M2—M4实验确认；当前Oracle结果仅支持其研究价值和候选互补性。")

heading(doc, "2 相关工作", 1)
heading(doc, "2.1 建筑损毁遥感评估", 2)
para(doc, "大规模公开数据集是建筑损毁智能评估发展的重要基础。xBD提供多灾种、多区域的灾前与灾后光学影像、建筑多边形及损毁等级标签，促进了建筑定位与损毁分类模型的统一比较[1]。ChangeOS从语义变化检测角度处理建筑损毁，强调对象级空间一致性与多任务学习[3]。随着Transformer和状态空间模型的发展，建筑损毁模型开始加强长距离上下文和跨时相交互。ChangeMamba将视觉状态空间模型用于二值变化检测、语义变化检测和建筑损毁评估，在保持全局建模能力的同时探索更高效的时空交互[4]。这些研究主要提升主干网络和全图预测能力，而本文关注基础模型完成推理之后，如何对局部错误进行可控修正。")
para(doc, "除卫星双时相方法外，Tennant等针对火山灰降落灾害，利用灾后UAV光学影像构建了建筑定位与损毁分类流水线[9]。该研究说明不同观测平台和灾害类型对损毁表征提出了不同要求，也进一步体现了将定位、分类与后处理环节分别分析的必要性。")
heading(doc, "2.2 光学–SAR多模态建筑损毁识别", 2)
para(doc, "光学与SAR在灾害应急中具有明显互补性。灾前光学影像容易获得，并能提供直观的建筑边界、屋顶纹理及周边环境信息；灾后SAR不依赖太阳照明，并具备全天候观测能力，可在夜间或光学影像受云层影响时提供观测[2]。Wang和Jin较早研究了灾前光学与灾后SAR信息的联合使用[5]。BRIGHT将这一组合扩展为全球分布、多灾种的甚高分辨率基准，并系统评估多种深度模型与跨事件设置[2]。然而，SAR侧视成像带来的几何形变与散射差异，使同一地物在两种模态中不具备简单的像素同质性。因此，多模态特征融合之外，还需要显式考虑对齐可靠性和局部证据质量。")
heading(doc, "2.3 区域级证据聚合与预测细化", 2)
para(doc, "像素独立分类容易在建筑内部产生碎片化结果，区域级聚合可以利用局部连续性抑制孤立噪声。SLIC通过在颜色和空间坐标构成的联合空间内聚类，生成大小相对均匀且边界贴合度较高的超像素[6]。对于本文任务，超像素的作用不是替代建筑实例，而是把具有相近光学外观和空间位置的像素组织为候选区域，使局部特征、类别分布和修改效用能够在统一尺度上计算。与直接多数投票不同，本文允许候选区域保持基础预测，只有当验证器预测修改具有正效用时才执行写回。")
heading(doc, "2.4 记忆增强与可靠性估计", 2)
para(doc, "在本文中，Memory Bank用于保存训练阶段获得的历史区域特征、类别信息及修改效用，并通过近邻检索为当前区域提供补充证据。对于跨事件建筑损毁，这一机制旨在复用不同灾害中相似的局部损毁模式，但也可能受到同事件泄漏、邻居质量不稳定和背景偏置影响。本文因此在检索时排除同图，并默认排除同一灾害事件，同时记录邻居来源事件数量与事件熵，以判断证据是否过度集中于单一场景。另一方面，选择性分类通过学习选择函数，在风险与覆盖率之间进行权衡[7]；SelectiveNet进一步将预测与拒绝机制联合建模[8]。本文借鉴其可靠性思想，但把“拒绝”具体化为保持基础预测，并把验证对象从整幅图像下沉到局部候选修改。")

heading(doc, "3 所提方法", 1)
heading(doc, "3.1 问题定义与总体框架", 2)
para(doc, "设灾前光学影像为X_rgb，灾后SAR影像为X_sar，像素级标签Y包含背景、完好、受损和毁坏四类。基础模型f_base输出类别概率图P_base及基础预测Y_base。本文将已有P3检查点视为f_base，不在局部模块训练阶段更新其参数。最终输出仍为与Y_base尺寸一致的四类语义图，不额外预测独立建筑ID。")
eq(doc, "P_base = f_base(X_rgb, X_sar),       Y_base = arg max_c P_base(c)")
para(doc, "整体流程由候选生成、Direct证据、跨事件Memory证据、效用验证和选择性写回五个环节构成。基础模型负责提供稳定的全图结果；局部模块只处理建筑概率较高且具有潜在修正价值的区域；若验证器无法确认修改有效，则直接保留Y_base。因此，该框架本质上是“基础预测加受控残差”，而不是重新训练一个与基础模型竞争的完整分割网络。")
add_caption(doc, "表1  方法中各组成部分的功能与边界")
table(doc, ["组成部分", "主要输入", "主要输出", "研究定位"], [
    ["P3基础模型", "灾前RGB、灾后SAR", "全图概率与语义预测", "稳定主干，不由局部模块替代"],
    ["SLIC候选", "RGB、建筑概率", "局部区域掩膜", "证据聚合和修改单元，不是实例标签"],
    ["Direct Head", "当前区域描述", "当前区域类别建议", "主要建议器"],
    ["跨事件Memory", "区域查询向量、训练记忆", "邻居后验与可靠性", "补充历史证据"],
    ["Utility Verifier", "基础、Direct、Memory及质量特征", "H0/HD/HM选择", "控制修改并允许回退"],
], widths=[3.0, 4.2, 4.2, 4.4])

heading(doc, "3.2 建筑引导的局部候选构建", 2)
para(doc, "首先由基础模型的非背景类别概率求得建筑概率P_bld。通过阈值筛选排除显然属于背景的区域，再在灾前RGB影像上生成SLIC超像素。使用灾前RGB进行分组的原因在于其建筑边界和地物纹理通常更清晰，可以减少SAR斑点噪声直接决定区域形状的影响。每个超像素与建筑候选相交后形成局部掩膜S_j；面积过小、建筑覆盖率过低或跨越明显非建筑边界的区域可在候选筛选阶段剔除。")
para(doc, "这种设计兼顾了建筑整体监督与局部证据分析。训练标签虽然常把整栋建筑赋为统一损毁等级，但模型错误并不一定覆盖整个建筑：某些屋顶局部、建筑边缘或SAR强散射区可能先出现错误。超像素使算法能够观察这些局部差异，又不要求改变标注格式。最终若一个区域被接受修改，只在S_j内部更新语义类别；相邻区域仍由各自证据独立决定。")

heading(doc, "3.3 局部区域特征表示", 2)
para(doc, "对于候选区域S_j，本文从三类信息构造区域描述向量。第一类是多模态表征，包括区域内RGB特征、SAR特征及融合特征的平均池化与最大池化结果；第二类是基础预测状态，包括四类平均概率、主类别、预测熵、前两类概率差值margin以及区域内类别一致性；第三类是区域质量，包括建筑覆盖率、超像素面积、边界强度和可用的对齐可靠性。将三类信息拼接后得到r_j：")
eq(doc, "r_j = [ f_j^(rgb-sar), p_j^base, u_j^base, q_j^region, q_j^align ].")
para(doc, "其中，预测熵反映类别分布的不确定性，margin反映第一候选与第二候选之间的可分性；区域质量用于区分真实建筑内部候选与边界混合区域；对齐可靠性则描述RGB查询区域与SAR局部响应的一致程度。若局部对齐模块尚未获得独立实验增益，可仅将相关量作为辅助特征或消融项，而不把它单独列为主要创新。")

heading(doc, "3.4 Direct与跨事件Memory互补证据", 2)
heading(doc, "3.4.1 Direct Head当前证据", 3)
para(doc, "Direct Head读取r_j并输出区域类别分布p_j^D。它不依赖历史检索，直接回答“依据当前影像和基础特征，该区域更可能属于哪一类”。在整个框架中，Direct是主要修改建议器，因为其证据与当前样本最直接，且推理路径短、误差来源相对明确。Direct建议类别与Y_base一致时，该区域无需产生残差；二者不一致时，形成候选假设H_D。")
heading(doc, "3.4.2 排除同事件的Memory检索", 3)
para(doc, "训练阶段将区域描述、真实类别、基础预测、建议方向和修改效用写入Memory Bank。对于当前查询r_j，使用归一化特征相似度检索Top-k近邻。为了检验真正的跨事件迁移能力，候选邻居必须排除同一影像，主实验进一步排除同一灾害事件。根据近邻真实类别或历史建议加权汇总得到Memory后验p_j^M，同时计算最大相似度、平均相似度、Top-k类别一致性、邻居来源事件数和事件熵。若近邻高度分散或过度集中于单一事件，验证器可以降低Memory证据权重。")
heading(doc, "3.4.3 Direct–Memory关系", 3)
para(doc, "Direct与Memory并非简单平均。根据二者是否建议改变基础预测，可将候选划分为四类：二者建议一致、仅Direct建议修改、仅Memory建议修改以及二者建议冲突。关系类型本身是重要特征：一致候选通常具有更强证据，Direct-only反映当前样本中特有的局部信息，Memory-only体现历史证据的补充空间，冲突候选则需要更严格的效用判断。Oracle统计中，Memory-only候选带来0.5739个百分点的潜在增益，说明将Memory完全并入Direct概率可能掩盖其独立价值。")

heading(doc, "3.5 Residual Utility Memory与效用验证器", 2)
para(doc, "传统Memory主要回答“这个区域像哪一类”，但局部细化真正需要回答的是“采用这个建议是否比保持基础预测更好”。因此，本文对每个候选构造三个假设：H_0保持基础预测，H_D采用Direct建议，H_M采用Memory建议。利用训练集真实标签计算区域质量函数Q后，可得到两种建议相对基础预测的残差效用：")
eq(doc, "U_D = Q(H_D, Y) - Q(H_0, Y),       U_M = Q(H_M, Y) - Q(H_0, Y).")
para(doc, "Q可采用区域像素正确率、区域IoU变化或兼顾类别不平衡的加权得分。为避免小区域的偶然变化主导学习，训练时应固定Q的定义，并同时记录有效修改像素数和有害修改像素数。Residual Utility Memory除保存区域特征与类别信息外，还保存建议来源、修改方向、真实效用、接受或拒绝标签以及来源事件。这样，检索结果不仅告诉模型历史区域的类别，还表明相似修改在历史上是否真正改善过基础预测。")
para(doc, "Utility Verifier采用低容量模型读取基础不确定性、Direct概率与margin、Memory后验与近邻质量、Direct–Memory关系、区域质量及对齐可靠性，分别估计两种修改的效用。第一版优先使用Logistic Regression或浅层多层感知机，以减少复杂网络掩盖机制问题的风险。推理时的决策为：")
eq(doc, "a* = arg max { 0, Û_D, Û_M }.")
para(doc, "当两种预测效用均不大于零或可靠性不足时选择H_0；当Direct效用最高时选择H_D；当Memory效用最高时选择H_M。该机制把局部细化转换为三选一的决策问题，使“保持原结果”成为与两种修改平等竞争的显式选项。")

heading(doc, "3.6 选择性写回与训练推理流程", 2)
para(doc, "对被验证器接受的候选区域，将相应建议转化为局部残差并写回基础语义图；未接受区域保持不变。若多个候选区域重叠，可按照验证器效用分数由高到低写回，或使用区域分区保证每个像素仅属于一个候选。最终结果可表示为：")
eq(doc, "Y_final = Y_base + G_D ⊙ R_D + G_M ⊙ R_M,")
para(doc, "其中R_D和R_M表示Direct与Memory相对基础结果的局部类别修改，G_D和G_M为互斥门控。训练阶段先冻结基础模型，在完整训练集上生成候选区域；随后通过交叉拟合或严格样本隔离生成Direct与Memory特征，使用训练标签计算效用并训练Verifier。验证阶段只读训练Memory，禁止把验证特征或标签写回Memory，也禁止使用验证GT决定局部选择。所有阈值和模型参数应在最终测试前冻结。")
add_caption(doc, "表2  训练与验证阶段的数据隔离规则")
table(doc, ["环节", "训练阶段", "验证/测试阶段"], [
    ["基础模型", "冻结参数并生成P_base", "仅前向推理"],
    ["Memory构建", "仅使用Train区域及标签", "只读Train Memory"],
    ["邻居过滤", "排除同图，主实验排除同事件", "排除查询事件"],
    ["Utility标签", "由Train真实标签计算", "不得使用GT参与决策"],
    ["阈值与Verifier", "在训练流程内确定", "全部冻结"],
], widths=[3.5, 6.1, 6.1])

heading(doc, "4 实验设计与初步结果", 1)
heading(doc, "4.1 数据集与评价协议", 2)
para(doc, "实验采用BRIGHT建筑损毁数据集。该数据集由灾前甚高分辨率光学影像、灾后甚高分辨率SAR影像及像素级建筑损毁标注组成，覆盖多类自然和人为灾害事件[2]。本文沿用既定训练、验证和测试划分，保持基础模型的输入、预处理和评价协议不变。Memory与Verifier只使用训练集构建；跨事件实验按照灾害事件隔离，避免同一事件的背景、建筑风格或传感器特征泄漏到查询结果。")
para(doc, "主要指标为总体mIoU及背景、完好、受损和毁坏四类IoU，并补充各类F1与事件级mIoU。由于本文研究局部修正，还应报告候选覆盖率、有效修改率、有害修改率、Affected漏检恢复率、回退比例和风险–覆盖关系。效率方面记录Memory容量、Top-k检索时间、Verifier参数量、整图推理时间和峰值显存。所有最终性能必须来自不读取验证或测试真实标签的正常推理；Oracle只用于诊断候选上限。")

heading(doc, "4.2 对比设置与M0—M5闭环", 2)
para(doc, "为了区分候选生成能力、Memory互补性和效用学习的真实贡献，本文设置从M0到M5的递进实验。M0为不含局部模块的P3基础模型；M1使用固定规则接受Direct建议，用于观察无学习门控的风险；M2使用Direct证据训练Logistic Verifier，是判断Memory价值的关键对照；M3在M2基础上加入普通Memory类别与近邻可靠性证据；M4进一步加入Residual Utility Memory，使历史修改效用参与决策；M5打乱Utility标签，检验提升是否来自有效效用监督而非模型容量或偶然相关。")
add_caption(doc, "表3  M0—M5主要实验矩阵")
table(doc, ["编号", "Direct", "普通Memory", "效用Memory", "Verifier", "主要问题"], [
    ["M0", "否", "否", "否", "否", "基础性能是多少"],
    ["M1", "是", "否", "否", "固定规则", "无可靠性控制是否退化"],
    ["M2", "是", "否", "否", "Logistic", "Direct能否形成真实增益"],
    ["M3", "是", "是", "否", "Logistic/浅层", "类别Memory是否超过M2"],
    ["M4", "是", "是", "是", "Logistic/浅层", "残差效用是否超过M3"],
    ["M5", "是", "是", "标签打乱", "同M4", "增益是否依赖有效监督"],
], widths=[1.4, 2.2, 2.5, 2.5, 2.8, 5.1], note="说明：M2—M5的真实推理结果完成后，才能确定Memory和Residual Utility是否构成最终方法贡献。")

heading(doc, "4.3 Oracle上限与互补性分析", 2)
para(doc, "当前已完成的Oracle实验用于回答“候选集合中是否存在可被选择器利用的改进空间”。Oracle在每个候选区域事后读取真实标签，选择能够获得最好结果的方案，因此代表理论上限而非可部署模型。P3基础结果为60.5020% mIoU；Memory单独Oracle、Direct单独Oracle和联合Oracle分别达到63.5212%、63.8589%和64.4420%。联合Oracle相对P3提升3.9401个百分点，相对Direct Oracle仍增加0.5831个百分点，超过预先设置的0.20个百分点互补性判断阈值。由此可以得出较稳妥的结论：Memory虽然不是最强的单独建议器，但确实提供了Direct没有覆盖的有效候选。")
add_caption(doc, "表4  已完成的Oracle候选上限实验")
table(doc, ["设置", "mIoU/%", "相对P3/百分点", "可支持的结论"], [
    ["P3基础模型", "60.5020", "0.0000", "当前稳定基线"],
    ["Memory Oracle", "63.5212", "+3.0193", "Memory候选中存在有效修改"],
    ["Direct Oracle", "63.8589", "+3.3570", "Direct是更强的主要建议器"],
    ["Direct∪Memory Oracle", "64.4420", "+3.9401", "二者具有候选互补性"],
], widths=[4.1, 3.0, 3.4, 6.0], note="Oracle使用真实标签进行事后最优选择，仅用于分析候选上限，不等同于最终模型性能。")
para(doc, "候选关系统计进一步解释了联合上限的来源。在现有统计中，Direct与Memory建议相同的候选为6018个，Memory-only为345个，Direct-only为1022个，二者冲突为529个；所有Memory与Direct不同的候选共874个，其中Oracle接受436个，占49.89%。这说明Memory差异候选并非普遍正确，但其中约一半具有正效用，恰好需要Verifier区分。Memory-only候选能够恢复69,625个Affected像素，并带来0.5739个百分点的Oracle增益；在14个事件中有11个事件得到改善，14个事件均未下降。该结果支持跨事件Memory作为补充证据，而不支持在缺少验证器时直接覆盖基础预测。")
add_caption(doc, "表5  Direct与Memory候选关系统计")
table(doc, ["关系类型", "候选数", "解释"], [
    ["Direct与Memory相同", "6018", "二者提供一致局部判断"],
    ["Memory-only", "345", "Memory提供Direct未提出的修改"],
    ["Direct-only", "1022", "当前区域证据独立提出修改"],
    ["Direct–Memory冲突", "529", "二者提出不同修改方向"],
    ["Memory差异候选合计", "874", "其中Oracle接受436个，占49.89%"],
], widths=[4.3, 3.0, 9.0])

heading(doc, "4.4 Affected漏检恢复漏斗", 2)
para(doc, "漏检恢复漏斗用于定位局部修正的主要瓶颈。P3共存在4,069,077个Affected漏检像素，其中3,107,860个落在预测建筑区域，占76.38%；SLIC候选覆盖2,410,825个，占全部漏检的59.25%。Memory提出Affected建议的像素为873,987个，占21.48%；在Oracle条件下Memory正确恢复709,192个，占17.43%。Direct Oracle可恢复916,963个，占22.53%，联合Oracle可恢复984,637个，占24.20%。因此，联合候选的主要价值是扩大可恢复集合，但仍有大量漏检在建筑候选或SLIC阶段即被排除。若后续Verifier性能受限，应先区分是候选覆盖不足，还是候选已覆盖但效用判断失败。")
add_caption(doc, "表6  Affected漏检恢复过程")
table(doc, ["阶段", "像素数", "占P3漏检比例/%", "含义"], [
    ["P3 Affected漏检", "4,069,077", "100.00", "全部待恢复错误"],
    ["预测建筑区域覆盖", "3,107,860", "76.38", "受建筑候选限制"],
    ["SLIC候选覆盖", "2,410,825", "59.25", "可进入局部推理的范围"],
    ["Memory建议Affected", "873,987", "21.48", "Memory主动提出的候选"],
    ["Memory Oracle恢复", "709,192", "17.43", "Memory候选理论正确恢复"],
    ["Direct Oracle恢复", "916,963", "22.53", "Direct候选理论正确恢复"],
    ["联合Oracle恢复", "984,637", "24.20", "当前候选集合的联合上限"],
], widths=[4.0, 3.1, 3.5, 5.7])

heading(doc, "4.5 消融与机制验证", 2)
para(doc, "候选区域消融需要比较不使用SLIC、单尺度SLIC和多尺度SLIC，并改变建筑概率阈值与最小区域面积，以确定性能变化来自合理的空间聚合而非候选数量简单增加。Memory消融应依次移除类别后验、相似度、Top-k一致性、事件数量和事件熵；数据隔离消融则比较普通检索、排除同图检索和排除同事件检索。若普通检索显著优于跨事件检索，只能说明方法依赖事件内相似性，不能宣称跨事件泛化。")
para(doc, "Verifier消融应从仅使用基础预测特征开始，依次加入Direct证据、Memory证据、Direct–Memory关系、区域质量和对齐可靠性。除总体mIoU外，还需要同时报告接受率、有效修改率和有害修改率：接受率过低可能得到小幅稳定增益但缺乏实际覆盖，接受率过高则可能放大有害修改。M5打乱Utility标签是关键反事实实验；若M5仍接近M4，说明所谓效用学习可能只是在利用类别或事件偏置，需要重新检查标签构造和数据泄漏。")
add_caption(doc, "表7  建议完成的消融实验与验收问题")
table(doc, ["实验组", "比较内容", "需要回答的问题"], [
    ["候选区域", "无SLIC/单尺度/多尺度、建筑阈值、面积阈值", "区域化是否提高有效覆盖并减少噪声"],
    ["Memory内容", "posterior、相似度、一致性、事件熵逐项移除", "哪些历史证据真正有效"],
    ["数据隔离", "普通/排除同图/排除同事件", "增益是否具有跨事件意义"],
    ["Verifier特征", "Base→Direct→Memory→关系→质量", "性能来自哪类决策信息"],
    ["Utility监督", "普通Memory/RUM/打乱标签", "残差效用是否可学习且非偶然"],
    ["可靠性", "风险–覆盖曲线、接受/有害修改率", "门控能否稳定控制修改风险"],
], widths=[3.0, 6.0, 7.2])

heading(doc, "4.6 定性分析与效率评价", 2)
para(doc, "定性结果应按照统一版式展示灾前RGB、灾后SAR、P3基础预测、Direct建议、Memory建议、Verifier选择、最终预测和真实标签。样例至少覆盖四种情形：Memory补回Direct遗漏、Direct正确而Memory有害、二者一致并被接受、二者冲突后回退P3。每个样例需同时给出候选区域边界和Verifier分数，避免只展示最终成功案例。事件级分析应选取增益较大与无增益事件，讨论建筑风格、灾害类型、SAR成像差异和候选覆盖率之间的关系。")
para(doc, "效率评价不应只报告基础网络参数量，还要单独统计局部模块的额外开销。Memory检索复杂度与库容量和特征维度相关，可通过离线建库、特征归一化和近似近邻索引降低开销；Verifier本身采用低容量模型，理论上只产生较小额外计算。论文最终应报告单幅影像候选数、平均检索时间、局部模块推理时间、峰值显存和Memory存储量，以判断方法是否适合灾后快速制图。")

heading(doc, "5 讨论", 1)
heading(doc, "5.1 研究贡献的合理边界", 2)
para(doc, "当前结果足以支持“跨事件Memory提供了Direct未覆盖的候选空间”，但尚不足以支持“Memory已经提高最终模型性能”。原因在于Oracle使用了真实标签进行选择，而实际推理只能依赖Verifier。论文的核心证据链应当是：M2证明Direct建议能够在无GT条件下被可靠筛选；M3证明普通Memory证据在M2之上仍有增益；M4证明历史残差效用比单纯类别检索更有价值；M5排除标签和容量造成的伪提升。只有当这一链条成立时，Residual Utility Memory才能被表述为主要模型创新。")
heading(doc, "5.2 语义分割、建筑整体标签与局部Token的关系", 2)
para(doc, "将建筑内部特征构造成Token集合，并不等同于把任务改成实例分割。Token只是内部计算单位，SLIC也只是区域分组工具；监督标签和最终输出仍是四类像素语义图。建筑整体被标为同一损毁等级，是数据集将建筑级判定栅格化后的结果。局部Token的作用是分析模型在建筑内部和边界附近的证据差异，再通过区域聚合生成更稳健的整体建议。因此，本文可以在不重新标注实例ID的情况下开展研究，也无需推翻此前的语义分割实验。")
heading(doc, "5.3 局限性", 2)
para(doc, "方法首先受到基础建筑覆盖范围的约束：若P3将受损建筑完全预测为背景，建筑概率阈值可能在候选生成阶段将其排除。其次，SLIC依赖RGB颜色与空间连续性，阴影、屋顶材质变化和紧邻建筑可能造成过分割或欠分割。再次，Memory的可靠性取决于特征空间是否具有跨事件可比性，近邻相似并不必然意味着修改方向相同。最后，Utility标签由训练集真实标签构造，若样本不平衡或同一事件区域高度相关，Verifier可能学习事件偏置而非一般规律。因此，严格的事件隔离、反事实消融与逐事件报告是论文可信度的重要组成部分。")

heading(doc, "6 结论", 1)
para(doc, "本文围绕光学–SAR建筑损毁语义分割中的局部漏检、跨事件模式差异和无条件修正风险，提出一套基于跨事件记忆与效用验证的局部细化思路。该方法以P3基础预测为稳定起点，通过建筑引导的SLIC候选组织局部证据，利用Direct Head和排除同事件的Memory分别提供当前与历史建议，再由Residual Utility Verifier选择保持基础结果、采用Direct或采用Memory。整个流程保持像素级语义分割任务不变，并通过回退机制限制局部模块造成的负面影响。")
para(doc, "已完成的Oracle分析显示，Direct与Memory联合候选相对P3具有3.9401个百分点的理论改进空间，且比Direct单独Oracle高0.5831个百分点；Memory-only候选及事件级统计进一步证明历史证据具有互补性。但这些数值仅代表候选上限。下一阶段应优先完成M2—M5正常推理实验、跨事件隔离消融和风险–覆盖分析，再根据真实结果决定论文题目与创新表述：若M4稳定超过M3和M2，可突出Residual Utility Memory；若Memory未形成最终增益，则应将其降为互补性分析，并把论文主线收敛到效用验证与选择性回退。")

heading(doc, "参考文献", 1)
refs = [
    "[1] Gupta R, Goodman B, Patel N, et al. Creating xBD: A Dataset for Assessing Building Damage from Satellite Imagery[C]//Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition Workshops. 2019: 10–17.",
    "[2] Chen H, Song J, Dietrich O, et al. Bright: A globally distributed multimodal building damage assessment dataset with very-high-resolution for all-weather disaster response[J]. Earth System Science Data, 2025, 17(11): 6217–6253. DOI: 10.5194/essd-17-6217-2025.",
    "[3] Zheng Z, Zhong Y, Wang J, Ma A, Zhang L. Building damage assessment for rapid disaster response with a deep object-based semantic change detection framework: From natural disasters to man-made disasters[J]. Remote Sensing of Environment, 2021, 265: 112636. DOI: 10.1016/j.rse.2021.112636.",
    "[4] Chen H, Song J, Han C, Xia J, Yokoya N. ChangeMamba: Remote Sensing Change Detection with Spatiotemporal State Space Model[J]. IEEE Transactions on Geoscience and Remote Sensing, 2024, 62: 1–20, Art. no. 4409720. DOI: 10.1109/TGRS.2024.3417253.",
    "[5] Wang T L, Jin Y Q. Postearthquake Building Damage Assessment Using Multi-Mutual Information from Pre-Event Optical Image and Postevent SAR Image[J]. IEEE Geoscience and Remote Sensing Letters, 2012, 9(3): 452–456. DOI: 10.1109/LGRS.2011.2170657.",
    "[6] Achanta R, Shaji A, Smith K, et al. SLIC Superpixels Compared to State-of-the-Art Superpixel Methods[J]. IEEE Transactions on Pattern Analysis and Machine Intelligence, 2012, 34(11): 2274–2282. DOI: 10.1109/TPAMI.2012.120.",
    "[7] Geifman Y, El-Yaniv R. Selective Classification for Deep Neural Networks[C]//Advances in Neural Information Processing Systems. 2017, 30: 4878–4887.",
    "[8] Geifman Y, El-Yaniv R. SelectiveNet: A Deep Neural Network with an Integrated Reject Option[C]//Proceedings of the 36th International Conference on Machine Learning. PMLR, 2019, 97: 2151–2159.",
    "[9] Tennant E, Jenkins S F, Miller V, et al. Automating tephra fall building damage assessment using deep learning[J]. Natural Hazards and Earth System Sciences, 2024, 24: 4585–4608. DOI: 10.5194/nhess-24-4585-2024.",
]
for item in refs:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.first_line_indent = Cm(-0.74)
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(3)
    font(p.add_run(item), size=9.2)

doc.add_page_break()
heading(doc, "附录A  论文完成前的实验与写作检查表", 1)
table(doc, ["优先级", "任务", "完成标准"], [
    ["P0", "完成M2、M3、M4真实推理", "同一划分、同一协议，报告总体/类别/事件指标"],
    ["P0", "完成M5打乱Utility反事实", "M4应明显优于随机效用监督"],
    ["P0", "核查Train/Val/事件隔离", "Val不建库、不写库、不使用GT选择"],
    ["P1", "完成Verifier可靠性分析", "报告接受率、有效/有害修改率和风险–覆盖曲线"],
    ["P1", "完成候选覆盖与恢复漏斗", "区分候选不足和选择失败"],
    ["P1", "完成事件级和定性分析", "包含成功、失败、冲突和回退案例"],
    ["P2", "完成效率统计", "检索时间、总推理时间、显存、存储量"],
    ["P2", "根据结果冻结题目与贡献", "不以Oracle替代最终性能，不夸大Memory作用"],
], widths=[2.0, 6.0, 8.2])

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
